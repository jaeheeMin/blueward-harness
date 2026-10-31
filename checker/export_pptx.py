"""고객 양식 pptx 산출(#561). `checker.export` 가 지도의 `템플릿` 이 .pptx 이면 이 모듈로 쓴다.

고객 양식 파일은 손대지 않는다. 지도(`templates/forms/<양식>.yaml`, 형식 정본은
`plugins/harness/skills/spec/form-map.md` 의 "pptx 지도")가 쪽(슬라이드)마다 도형을 **예시 글자의 앞부분**으로 가리킨다.
양식에는 자리 표시자가 없고 일반 도형뿐이라서다. 양식이 바뀌면 지도의 글자만 고친다 — 이 모듈에는 특정 양식의
글자·슬라이드 번호·좌표가 없다. 가리키는 도형이 0개거나 2개 이상이면 항목 이름을 알리고 산출을 거부한다(종료코드 1).

쓰기는 python-pptx 로 한다(왕복해도 도형·그림·글꼴이 보존되는 것을 시험했다). 견본 쪽을 복제할 때 관계(그림 등)를 새 쪽에
다시 이어 주고, 쓰지 않는 그림 관계는 끊어 예시 그림이 파일에 남지 않게 한다.

글 넘침·겹침·그림 위치는 화면 없이 확인할 수 없다. 번호 설명의 높이는 글자 수·상자 폭·글꼴 크기로 **추정**한 값이라서
산출은 항상 `layout-unchecked` 경고(화면 확인 필요, 검사 불능)를 낸다 — 통과가 아니다.

    python -m checker.export_pptx --dump <양식.pptx>    쪽마다 도형 목록(종류·글자·위치·찾기로 유일한지)을 낸다. 쓰기 없음.
"""
from __future__ import annotations

import argparse
import copy
import io
import re
import sys
import unicodedata
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

from checker.export import ExportError, Section, _clean, _inside, norm

PAGE_KEYS = {"이름", "설명", "견본", "칸", "지우기", "목록", "반복", "순번", "그림", "번호_설명", "유지"}
FIELD_KEYS = {"항목", "찾기", "틀", "설명"}
TOP_KEYS = {"양식", "설명", "템플릿", "검사_규칙", "산출", "쪽"}
LIST_KEYS = {"글", "번호", "제목"}
NUM_KEYS = {"번호", "글", "쪽당_최대", "아래"}
NUM_BADGE_KEYS = {"글자", "왼쪽_이상"}
EMU_PT = 12700
GAP = 60000  # 번호 설명 사이 간격(EMU)
SECTION_EXT_URI = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"  # PowerPoint 구역 목록. 쪽을 빼거나 늘리면 어긋나므로 걷어 낸다
LAYOUT_NOTE = (
    "글 넘침·겹침·그림 위치는 화면으로 확인하지 못했다(검사 불능, 통과 아님). PowerPoint 로 열어 쪽마다 눈으로 확인한다. "
    "번호 설명의 높이는 글자 수·상자 폭·글꼴 크기로 추정한 값이다"
)


# ---------------------------------------------------------------- 지도 검증

def is_pptx_map(fmap: dict) -> bool:
    return str(fmap.get("템플릿", "")).lower().endswith(".pptx")


def validate_map(data: dict, path: str, form: str) -> None:
    """지도의 모양(키·형)을 본다. 틀리면 항목을 말하고 읽지 못함(ExportError)으로 알린다. 양식 안의 글자는 여기서 보지 않는다."""
    def bad(page: str, item: str, reason: str):
        raise ExportError(f"양식 지도 {path}: {reason}", form=form, sheet=page, item=item)

    extra = set(data) - TOP_KEYS
    if extra:
        bad("", sorted(extra)[0], f"pptx 지도가 지원하지 않는 키 {sorted(extra)}(허용: {sorted(TOP_KEYS)})")
    pages = data.get("쪽")
    if not isinstance(pages, list) or not pages:
        bad("", "쪽", "`쪽` 목록이 없다")
    seen: set[str] = set()
    samples: dict[int, str] = {}
    for p in pages:
        if not isinstance(p, dict) or not p.get("이름"):
            bad("", "쪽", "쪽에 `이름` 이 없다")
        name = norm(p["이름"])
        if name in seen:
            bad(name, "이름", f"쪽 이름 {name!r} 이 두 번 나온다")
        seen.add(name)
        extra = set(p) - PAGE_KEYS
        if extra:
            bad(name, sorted(extra)[0], f"쪽 {name!r} 이 지원하지 않는 키 {sorted(extra)}(허용: {sorted(PAGE_KEYS)})")
        if not isinstance(p.get("견본"), int) or isinstance(p.get("견본"), bool):
            bad(name, "견본", f"쪽 {name!r} 의 `견본` 은 양식의 슬라이드 번호(정수)여야 한다")
        if p["견본"] in samples:  # 한 견본 쪽을 두 쪽이 쓰면 첫 쪽으로 쓰는 견본 슬라이드가 겹쳐 채워진다(반복 쪽도 예외 없음)
            bad(name, "견본", f"쪽 {name!r} 과 {samples[p['견본']]!r} 이 같은 견본 {p['견본']} 을 쓴다(쪽마다 견본이 달라야 한다)")
        samples[p["견본"]] = name
        for f in p.get("칸") or []:
            if not isinstance(f, dict) or not f.get("항목") or not f.get("찾기"):
                bad(name, "칸", f"쪽 {name!r} 의 칸에 `항목`·`찾기` 가 없다")
            if set(f) - FIELD_KEYS:
                bad(name, str(f["항목"]), f"칸 {f['항목']!r} 이 지원하지 않는 키 {sorted(set(f) - FIELD_KEYS)}")
            if f.get("틀") is not None and "{값}" not in str(f["틀"]):
                bad(name, str(f["항목"]), f"칸 {f['항목']!r} 의 `틀` 에 `{{값}}` 이 없다")
        if p.get("칸") is not None and not isinstance(p["칸"], list):
            bad(name, "칸", f"쪽 {name!r} 의 `칸` 은 목록이어야 한다")
        if not all(isinstance(x, str) for x in (p.get("지우기") or [])):
            bad(name, "지우기", f"쪽 {name!r} 의 `지우기` 는 글자 목록이어야 한다")
        if p.get("유지") is not None and not all(isinstance(x, str) for x in p["유지"]):
            bad(name, "유지", f"쪽 {name!r} 의 `유지` 는 글자 목록이어야 한다")
        for key, keys in (("목록", LIST_KEYS), ("순번", {"찾기"}), ("그림", {"자리"}), ("번호_설명", NUM_KEYS)):
            v = p.get(key)
            if v is None:
                continue
            if not isinstance(v, dict) or set(v) - keys:
                bad(name, key, f"쪽 {name!r} 의 `{key}` 는 {sorted(keys)} 만 가진 사전이어야 한다")
        if p.get("목록") and not all(p["목록"].get(k) for k in LIST_KEYS):
            bad(name, "목록", f"쪽 {name!r} 의 `목록` 에는 `글`·`번호`·`제목`(본문 쪽의 칸 항목 이름)이 모두 있어야 한다")
        if p.get("순번") and not p["순번"].get("찾기"):
            bad(name, "순번", f"쪽 {name!r} 의 `순번` 에 `찾기` 가 없다")
        if p.get("그림") and (not isinstance(p["그림"].get("자리"), int) or p["그림"]["자리"] < 1):
            bad(name, "그림", f"쪽 {name!r} 의 `그림.자리` 는 1 이상의 정수여야 한다")
        n = p.get("번호_설명")
        if n:
            b = n.get("번호")
            if not isinstance(b, dict) or set(b) - NUM_BADGE_KEYS or not b.get("글자") or not n.get("글"):
                bad(name, "번호_설명", f"쪽 {name!r} 의 `번호_설명` 에는 `번호`({{글자, 왼쪽_이상}})와 `글` 이 있어야 한다")
            if not isinstance(n.get("쪽당_최대", 1), int) or n.get("쪽당_최대", 1) < 1:
                bad(name, "번호_설명", f"쪽 {name!r} 의 `쪽당_최대` 는 1 이상의 정수여야 한다")
    for p in pages:
        t = (p.get("목록") or {}).get("제목")
        if t and not any(norm(f["항목"]) == norm(t) for q in pages if q.get("반복") for f in q.get("칸") or []):
            bad(norm(p["이름"]), "목록", f"`목록.제목` {t!r} 이 반복 쪽의 칸 항목에 없다")


# ---------------------------------------------------------------- 도형 찾기

def _is_text(shape) -> bool:
    return shape.has_text_frame and shape._element.tag == qn("p:sp")


def _shape_text(shape) -> str:
    return norm(shape.text_frame.text) if _is_text(shape) else ""


def find_all(slide, prefix: str) -> list:
    want = norm(prefix)
    return [s for s in slide.shapes if want and _shape_text(s).startswith(want)]


def _has_fld(shape) -> bool:
    return shape._element.find(".//" + qn("a:fld")) is not None


def _is_pic(shape) -> bool:
    return shape._element.tag == qn("p:pic")


def _is_table(shape) -> bool:
    return shape._element.tag == qn("p:graphicFrame") and shape._element.find(".//" + qn("a:tbl")) is not None


def _kind(shape) -> str:
    tag = etree.QName(shape._element).localname
    if tag == "pic":
        return "그림"
    if tag == "graphicFrame":
        return "표" if _is_table(shape) else "프레임"
    if tag == "grpSp":
        return "묶음"
    if tag == "cxnSp":
        return "선"
    return "글상자" if _is_text(shape) and _shape_text(shape) else "도형"


# ---------------------------------------------------------------- 글 쓰기·높이 추정

def set_text(shape, text: str) -> None:
    """첫 문단의 글꼴·서식을 그대로 두고 글만 바꾼다. 줄바꿈(`\\n`)마다 첫 문단을 복제해 문단을 늘린다."""
    txb = shape._element.find(qn("p:txBody"))
    paras = txb.findall(qn("a:p"))
    base = paras[0]
    for p in paras[1:]:
        txb.remove(p)
    run0 = base.find(qn("a:r"))
    rpr = run0.find(qn("a:rPr")) if run0 is not None else None
    if rpr is None:
        end = base.find(qn("a:endParaRPr"))
        rpr = etree.Element(qn("a:rPr"))
        if end is not None:
            for k, v in end.attrib.items():
                rpr.set(k, v)
            for c in end:
                rpr.append(copy.deepcopy(c))
    for tag in ("a:r", "a:br", "a:fld"):
        for e in base.findall(qn(tag)):
            base.remove(e)
    proto = copy.deepcopy(base)
    last = base
    for i, line in enumerate(text.split("\n")):
        p = base if i == 0 else copy.deepcopy(proto)
        r = etree.Element(qn("a:r"))
        r.append(copy.deepcopy(rpr))
        t = etree.SubElement(r, qn("a:t"))
        t.text = line
        end = p.find(qn("a:endParaRPr"))
        if end is not None:
            end.addprevious(r)
        else:
            p.append(r)
        if i:
            last.addnext(p)
            last = p


def has_mixed_runs(shape) -> bool:
    """도형 안 run 들의 서식(rPr)이 서로 다르면 True. 글을 바꾸면 첫 run 의 서식 하나로 합쳐진다."""
    fmts = {etree.tostring(r.find(qn("a:rPr"))) if r.find(qn("a:rPr")) is not None else b""
            for r in shape._element.iter(qn("a:r"))}
    return len(fmts) > 1


def _font_pt(shape) -> float:
    for e in shape._element.iter(qn("a:rPr")):
        if e.get("sz"):
            return int(e.get("sz")) / 100
    return 18.0


def _line_pct(shape) -> float:
    e = shape._element.find(".//" + qn("a:lnSpc") + "/" + qn("a:spcPct"))
    return int(e.get("val")) / 100000 if e is not None else 1.0


def _wide(ch: str) -> bool:
    return unicodedata.east_asian_width(ch) in ("W", "F")


def estimate_height(shape, text: str) -> int:
    """글이 들어갈 상자 높이(EMU)의 **추정**. 한글·전각은 글꼴 크기 1칸, 그 밖은 0.6칸으로 세어 상자 폭에서 줄바꿈한다.
    줄 높이는 글꼴 크기 × 1.2 × 줄간격 비율. 실제 렌더링과 다를 수 있어 화면 확인이 필요하다."""
    body = shape._element.find(".//" + qn("a:bodyPr"))
    li = int(body.get("lIns", 91440)) if body is not None else 91440
    ri = int(body.get("rIns", 91440)) if body is not None else 91440
    ti = int(body.get("tIns", 45720)) if body is not None else 45720
    bi = int(body.get("bIns", 45720)) if body is not None else 45720
    em = _font_pt(shape) * EMU_PT
    avail = max(shape.width - li - ri, em)
    lines = 0
    for para in text.split("\n"):
        x, n = 0.0, 1
        for ch in para:
            w = em * (1.0 if _wide(ch) else 0.6)
            if x + w > avail:
                n += 1
                x = 0.0
            x += w
        lines += n
    return int(lines * em * 1.2 * _line_pct(shape) + ti + bi)


# ---------------------------------------------------------------- 쪽 복제·정리

def _dup(prs, tpl):
    """견본 쪽을 복제한다. 그림 등 관계는 새 쪽에 다시 잇고 도형 안의 rId 를 바꿔 단다."""
    new = prs.slides.add_slide(tpl.slide_layout)
    tree = new.shapes._spTree
    for e in list(tree)[2:]:
        tree.remove(e)
    remap: dict[str, str] = {}
    for rid, rel in tpl.part.rels.items():
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
            continue
        remap[rid] = (
            new.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref) if rel.is_external
            else new.part.rels.get_or_add(rel.reltype, rel.target_part)
        )
    bg = tpl._element.find(qn("p:cSld") + "/" + qn("p:bg"))
    if bg is not None:
        new._element.find(qn("p:cSld")).insert(0, copy.deepcopy(bg))
    for e in list(tpl.shapes._spTree)[2:]:
        tree.append(copy.deepcopy(e))
    rns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    for e in new._element.iter():
        for k, v in list(e.attrib.items()):
            if k.startswith(rns) and v in remap:
                e.set(k, remap[v])
    return new


def _drop_unused_rels(slide) -> None:
    """쪽 안에서 쓰지 않는 관계(지워진 예시 그림 등)를 끊는다. 끊긴 그림 파트는 저장되지 않는다."""
    rns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    used = {v for e in slide._element.iter() for k, v in e.attrib.items() if k.startswith(rns)}
    for rid, rel in list(slide.part.rels.items()):
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE) or rid in used:
            continue
        slide.part.rels.pop(rid)


def _remove(shape) -> None:
    el = shape._element
    el.getparent().remove(el)


def _next_id(slide) -> int:
    return max((int(e.get("id", 0)) for e in slide._element.iter(qn("p:cNvPr"))), default=0) + 1


# ---------------------------------------------------------------- 도형 해석(변경 없음)

class _PageRejected(Exception):
    """채우는 중 도형 찾기가 어긋났다. args[0] 은 [(항목, 이유)]."""


class _Missing(Exception):
    """찾기 실패·중복. (항목, 이유)."""


def _one(slide, prefix: str, item: str, taken: dict):
    hits = find_all(slide, prefix)
    if len(hits) != 1:
        raise _Missing(item, (
            f"양식에서 {prefix!r} 로 시작하는 도형을 찾지 못했다(0개). 양식 글자가 바뀌었으면 지도의 `찾기` 를 고친다. "
            "`python -m checker.export_pptx --dump <양식.pptx>` 로 도형을 본다" if not hits else
            f"양식에서 {prefix!r} 로 시작하는 도형이 {len(hits)}개라 하나로 정할 수 없다. 더 긴 앞부분으로 `찾기` 를 고친다"
        ))
    shape = hits[0]
    other = taken.get(shape.shape_id)
    if other and other != item:
        raise _Missing(item, f"{shape.shape_id}번 도형을 {other!r} 항목도 가리킨다(한 도형을 두 항목이 가리킬 수 없다)")
    taken[shape.shape_id] = item
    return shape


class _Locator:
    """한 쪽 지도가 가리키는 도형들을 견본 쪽에서 찾는다(수정 없음). 못 찾은 항목은 problems 에 모은다."""

    def __init__(self, slide, pmap: dict, prs_height: int):
        self.slide, self.pmap = slide, pmap
        self.problems: list[tuple[str, str]] = []
        self.fields: list[tuple[dict, object]] = []
        self.erase: list = []
        self.seq = self.list_text = self.list_num = self.pic = self.badge = self.note = None
        self.keep: list = []
        taken: dict = {}

        def get(prefix: str, item: str):
            try:
                return _one(slide, prefix, item, taken)
            except _Missing as e:
                self.problems.append((e.args[0], e.args[1]))
                return None

        for f in pmap.get("칸") or []:
            s = get(str(f["찾기"]), str(f["항목"]))
            if s is not None:
                self.fields.append((f, s))
        for t in pmap.get("지우기") or []:
            s = get(t, f"지우기 {t!r}")
            if s is not None:
                self.erase.append(s)
        if pmap.get("순번"):
            self.seq = get(str(pmap["순번"]["찾기"]), "순번")
        if pmap.get("목록"):
            self.list_text = get(str(pmap["목록"]["글"]), "목록.글")
            self.list_num = get(str(pmap["목록"]["번호"]), "목록.번호")
        if pmap.get("그림"):
            pics = [s for s in slide.shapes if _is_pic(s)]
            n = pmap["그림"]["자리"]
            if n > len(pics):
                self.problems.append(("그림", f"견본 쪽에 그림이 {len(pics)}개뿐이라 `그림.자리` {n} 번째를 찾지 못했다"))
            else:
                self.pic = pics[n - 1]
        nm = pmap.get("번호_설명")
        if nm:
            b = nm["번호"]
            want, left = norm(str(b["글자"])), int(b.get("왼쪽_이상", 0))
            hits = [s for s in slide.shapes if _shape_text(s) == want and s.left is not None and s.left >= left]
            if len(hits) != 1:
                self.problems.append(("번호_설명.번호", (
                    f"글자 {want!r}, 왼쪽 {left} 이상인 도형이 {len(hits)}개라 하나로 정할 수 없다. `번호.왼쪽_이상` 을 조정한다")))
            else:
                self.badge = hits[0]
                taken[self.badge.shape_id] = "번호_설명.번호"
            self.note = get(str(nm["글"]), "번호_설명.글")
        for k in pmap.get("유지") or []:
            if norm(k) == "표":
                self.keep += [s for s in slide.shapes if _is_table(s)]
            else:
                s = get(k, f"유지 {k!r}")
                if s is not None:
                    self.keep.append(s)
        self.height = prs_height


# ---------------------------------------------------------------- 원본(md) 읽기

def split_sheet(sheet: str) -> tuple[str, str]:
    head, _, rest = sheet.partition("·")
    return norm(head), norm(rest)


def normalize_sections(fmap: dict, sections: dict, form: str) -> dict:
    """작업본 소절을 지도에 맞춘다. `#### <항목>` 본문으로 쓴 칸 값은 `| 항목 | 값 |` 표처럼 fields 로 옮긴다.
    지도에 없는 쪽·칸·소제목은 읽지 못함(ExportError)이다."""
    pages = {norm(p["이름"]): p for p in fmap["쪽"]}
    out: dict = {}
    for sheet, (sec, subs) in sections.items():
        head, rest = split_sheet(sheet)
        p = pages.get(head)
        if p is None:
            raise ExportError(f"작업본의 시트 {sheet!r} 이 양식 지도에 없다(쪽: {', '.join(pages)})", form=form, sheet=sheet)
        if p.get("반복") and not rest:
            raise ExportError(
                f"쪽 {head!r} 은 반복 쪽이다 — `### {form} · {head} · <{p['반복']} 이름>` 소절로 쓴다", form=form, sheet=sheet)
        if not p.get("반복") and rest:
            raise ExportError(f"쪽 {head!r} 은 반복 쪽이 아니다 — 소절 제목은 `### {form} · {head}` 로 쓴다", form=form, sheet=sheet)
        if sec.issues:
            raise ExportError(f"{form} · {sheet}: {sec.issues[0][0]}: {sec.issues[0][1]}", form=form, sheet=sheet)
        names = {norm(f["항목"]) for f in p.get("칸") or []}
        allowed_subs = names | ({"그림"} if p.get("그림") else set()) | ({"번호 설명"} if p.get("번호_설명") else set())
        for n in sec.fields:
            if n not in names:
                raise ExportError(f"{form} · {sheet}: 칸 {n!r} 이 지도에 없다(허용: {', '.join(sorted(names)) or '없음'})",
                                  form=form, sheet=sheet, item=n)
        for n in subs:
            if n not in allowed_subs:
                raise ExportError(f"{form} · {sheet}: `#### {n}` 이 지도에 없다(허용: {', '.join(sorted(allowed_subs)) or '없음'})",
                                  form=form, sheet=sheet, item=n)
        fields = dict(sec.fields)
        for n, lines in subs.items():
            if n in names and n not in fields:
                fields[n] = "\n".join(l.rstrip() for l in lines).strip()
        new = Section(na=sec.na, fields=fields, tables=sec.tables, texts=sec.texts, issues=sec.issues)
        out[sheet] = (new, {k: v for k, v in subs.items() if k not in names})
    return out


def _picture_path(lines: list[str]) -> str:
    for l in lines:
        s = l.strip()
        if not s:
            continue
        m = re.match(r"!\[[^\]]*\]\(([^)]+)\)", s)
        s = m.group(1) if m else s
        return s.strip().strip("`'\"<>").strip()
    return ""


_NUM_LINE = re.compile(r"^\s*(\d+)[.)]\s+(.*\S)\s*$")


def _numbered(lines: list[str]) -> list[tuple[str, str]] | str:
    """`1. 글` 목록 → [(번호 글자, 글)]. 들여쓴 줄은 앞 항목의 다음 줄. 모양이 틀리면 사유 글자를 낸다."""
    items: list[list[str]] = []
    for l in lines:
        if not l.strip():
            continue
        m = _NUM_LINE.match(l)
        if m:
            items.append([m.group(1), m.group(2)])
        elif l[:1] in (" ", "\t") and items:
            items[-1][1] += "\n" + l.strip()
        else:
            return f"번호 목록 모양이 아닌 줄이 있다: {l.strip()!r}(`1. 글` 처럼 쓰고 이어지는 줄은 들여쓴다)"
    seen = [i[0] for i in items]
    if len(set(seen)) != len(seen):
        return "번호가 겹친다: " + ", ".join(sorted({n for n in seen if seen.count(n) > 1}))
    return [(a, b.strip()) for a, b in items if _clean(b)]


# ---------------------------------------------------------------- 쪽 채우기

def _fit(img_w: int, img_h: int, box: tuple[int, int, int, int]) -> tuple[int, int, int, int]:
    x, y, w, h = box
    k = min(w / img_w, h / img_h)
    nw, nh = int(img_w * k), int(img_h * k)
    return x + (w - nw) // 2, y + (h - nh) // 2, nw, nh


def _plan_chunks(loc: _Locator, items: list[tuple[str, str]], bottom: int, warn) -> list[list[tuple[str, str, int]]]:
    """번호 설명을 쪽마다 나눈다. 글 높이는 추정이다. 쪽당_최대나 영역 아래를 넘으면 다음 쪽으로 넘긴다."""
    nm = loc.pmap["번호_설명"]
    cap = nm.get("쪽당_최대") or 10**6
    badge_h = loc.badge.height
    chunks: list[list[tuple[str, str, int]]] = [[]]
    y = loc.note.top
    for num, text in items:
        h = max(estimate_height(loc.note, text), badge_h)
        cur = chunks[-1]
        if cur and (len(cur) >= cap or y + h > bottom):
            chunks.append([])
            cur = chunks[-1]
            y = loc.note.top
        if not cur and y + h > bottom:
            warn(f"번호 {num} 설명 하나가 영역({bottom - loc.note.top} EMU)보다 높을 것으로 추정된다(추정 {h} EMU) — 글을 줄이거나 지도 `아래` 를 늘린다")
        cur.append((num, text, h))
        y += h + GAP
    return chunks


def _fill_page(slide, loc_pmap: dict, prs, ctx: dict) -> None:
    """복제된(또는 견본 그대로의) 쪽 하나를 채운다. 도형은 이 쪽에서 다시 찾는다(견본에서 유일함을 이미 확인했다)."""
    loc = _Locator(slide, loc_pmap, prs.slide_height)
    if loc.problems:  # 견본에서는 찾았는데 이 쪽에서 못 찾는 경우 — 조용히 넘기지 않고 거부로 올린다
        raise _PageRejected(loc.problems)
    values, seq, titles = ctx["values"], ctx.get("seq"), ctx.get("titles")
    warn, page = ctx["warn"], norm(loc_pmap["이름"])

    def put(shape, text: str, item: str) -> None:
        if has_mixed_runs(shape):
            warn(page, item, "format-merged", f"{item} 도형의 글에 서식이 다른 run 이 여럿이라 첫 run 의 서식 하나로 합쳐졌다(글꼴·굵기·색 확인)")
        set_text(shape, text)

    for f, shape in loc.fields:
        v = values[norm(f["항목"])]
        put(shape, str(f["틀"]).replace("{값}", v) if f.get("틀") else v, str(f["항목"]))
    if loc.seq is not None:
        put(loc.seq, f"{seq:02d}", "순번")
    if loc.list_text is not None:
        put(loc.list_text, "\n".join(titles), "목록.글")
        put(loc.list_num, "\n".join(str(i) for i in range(1, len(titles) + 1)), "목록.번호")
    for s in loc.erase:
        _remove(s)
    cloned: list = []
    new_pic = None
    if loc.pic is not None:
        data, (iw, ih) = ctx["image"]
        box = (loc.pic.left, loc.pic.top, loc.pic.width, loc.pic.height)
        x, y, w, h = _fit(iw, ih, box)
        new_pic = slide.shapes.add_picture(io.BytesIO(data), x, y, w, h)
        loc.pic._element.addprevious(new_pic._element)
        _remove(loc.pic)
    if loc.badge is not None:
        b0, t0 = loc.badge._element, loc.note._element
        y = loc.note.top
        for num, text, h in ctx["chunk"]:
            b, t = copy.deepcopy(b0), copy.deepcopy(t0)
            t0.addprevious(b)
            t0.addprevious(t)
            for el in (b, t):
                el.find(".//" + qn("p:cNvPr")).set("id", str(_next_id(slide)))
            sb, st = [s for s in slide.shapes if s._element is b][0], [s for s in slide.shapes if s._element is t][0]
            put(sb, num, "번호_설명.번호")
            put(st, text, "번호_설명.글")
            st.top, st.height = y, h
            sb.top = y + (loc.badge.top - loc.note.top)
            cloned += [sb, st]
            y += h + GAP
        b0.getparent().remove(b0)
        t0.getparent().remove(t0)
    if "유지" in loc_pmap:
        keep = {id(s._element) for s in loc.keep} | {id(s._element) for _, s in loc.fields}
        keep |= {id(s._element) for s in (loc.seq, loc.list_text, loc.list_num, new_pic) if s is not None}
        keep |= {id(s._element) for s in cloned}
        for s in list(slide.shapes):
            if id(s._element) not in keep and not _has_fld(s):
                _remove(s)
    deleted = bool(loc.erase) or loc.pic is not None or loc.badge is not None or "유지" in loc_pmap
    timing = slide._element.find(qn("p:timing"))
    if deleted and timing is not None:  # 애니메이션은 도형 id 로 도형을 가리킨다. 지운 도형을 가리키면 PowerPoint 가 복구를 묻는다
        slide._element.remove(timing)
        warn(page, "애니메이션", "animation-removed", f"쪽 {page!r} 은 도형을 지우므로 애니메이션(p:timing)을 지웠다")
    _drop_unused_rels(slide)


def _order_slides(prs, ordered: list) -> list[int]:
    """sldIdLst 를 지도 순서로 다시 놓는다. 지도가 가리키지 않은 견본 쪽은 빼고 관계를 끊는다. 뺀 쪽의 원래 번호를 낸다."""
    lst = prs.slides._sldIdLst
    by_id = {int(e.get("id")): e for e in lst}
    keep_ids = [s.slide_id for s in ordered]
    dropped = []
    for i, s in enumerate(prs.slides, 1):
        if s.slide_id not in keep_ids:
            dropped.append(i)
    gone = [by_id[i] for i in by_id if i not in keep_ids]
    for e in list(lst):
        lst.remove(e)
    for sid in keep_ids:
        lst.append(by_id[sid])
    for e in gone:
        prs.part.drop_rel(e.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"))
    ext = prs.part._element.find(qn("p:extLst"))
    if ext is not None:
        for x in list(ext):
            if x.get("uri") == SECTION_EXT_URI:
                ext.remove(x)
        if not len(ext):
            prs.part._element.remove(ext)
    return dropped


def render(form: str, fmap: dict, tpl: Path, sections: dict, root: Path) -> tuple[bytes | None, list[dict], list[dict]]:
    """견본 pptx 에 작업본 값을 채워 pptx 바이트를 낸다. (바이트 또는 None, 거부 사유들, 경고들).
    거부 사유가 하나라도 있으면 바이트는 None 이다. 양식 파일을 못 읽으면 ExportError."""
    problems: list[dict] = []
    warnings: list[dict] = []

    def problem(page: str, item: str, reason: str) -> None:
        problems.append({"form": form, "sheet": page, "item": item, "reason": reason})

    def warn(page: str, item: str, kind: str, reason: str) -> None:
        warnings.append({"form": form, "sheet": page, "item": item, "type": kind, "reason": reason})

    try:
        prs = Presentation(str(tpl))
    except Exception as e:  # 깨진 zip·pptx 가 아닌 파일 등 종류가 다양하다
        raise ExportError(f"템플릿 {tpl} 를 읽지 못했다: {type(e).__name__}: {e}", form=form) from None
    slides = list(prs.slides)
    pages = fmap["쪽"]

    # 1) 견본에서 도형을 모두 찾아 본다(수정 없음). 0개·2개 이상이면 거부.
    locs: dict[str, _Locator] = {}
    for p in pages:
        name = norm(p["이름"])
        n = p["견본"]
        if not 1 <= n <= len(slides):
            problem(name, "견본", f"견본 {n} 이 양식의 슬라이드 수({len(slides)}) 밖이다")
            continue
        locs[name] = _Locator(slides[n - 1], p, prs.slide_height)
        for item, reason in locs[name].problems:
            problem(name, item, reason)

    # 2) 작업본 값 모으기
    by_page: dict[str, list[tuple[str, Section, dict]]] = {norm(p["이름"]): [] for p in pages}
    for sheet, (sec, subs) in sections.items():
        head, rest = split_sheet(sheet)
        by_page[head].append((rest, sec, subs))
    shared: dict[str, str] = {}  # 반복하지 않는 쪽의 칸 값(다른 쪽이 같은 항목을 쓸 때 빌려 간다)
    for p in pages:
        if not p.get("반복"):
            for _, sec, _ in by_page[norm(p["이름"])]:
                for k, v in sec.fields.items():
                    if _clean(v):
                        shared.setdefault(k, _clean(v))

    plan: dict[str, list[dict]] = {}  # 쪽 이름 → 쪽 하나씩의 채울 내용
    for p in pages:
        name = norm(p["이름"])
        if name not in locs:
            continue
        loc = locs[name]
        entries = by_page[name]
        if p.get("반복") and not entries:
            problem(name, p["반복"], f"반복 쪽 {name!r} 의 소절(`### {form} · {name} · <{p['반복']} 이름>`)이 작업본에 없다")
        if not p.get("반복") and not entries:
            entries = [("", Section(), {})]
        out: list[dict] = []
        seq = 0
        for rest, sec, subs in entries:
            seq += 1
            label = f"{name} · {rest}" if rest else name
            values: dict[str, str] = {}
            for f in p.get("칸") or []:
                item = norm(f["항목"])
                v = _clean(sec.fields.get(item)) or (shared.get(item, "") if not p.get("반복") else "")
                if not v:
                    problem(label, item, f"{item} 값이 없다(비었거나 미정이다)")
                values[item] = v
            ctx: dict = {"values": values, "seq": seq, "chunk": [], "label": label}
            if p.get("그림") and loc.pic is not None:
                rel = _picture_path(subs.get("그림", []))
                if not rel:
                    problem(label, "그림", "`#### 그림` 이 없거나 경로가 비어 있다(저장소 루트 기준 PNG 경로를 쓴다)")
                else:
                    try:
                        path = _inside(root, root / rel, "그림", form)
                    except ExportError as e:
                        problem(label, "그림", e.reason)
                        path = None
                    if path is not None:
                        try:
                            from PIL import Image

                            data = path.read_bytes()
                            with Image.open(io.BytesIO(data)) as im:
                                size = im.size
                            ctx["image"] = (data, size)
                        except OSError as e:
                            problem(label, "그림", f"그림 파일 {rel} 을 읽지 못했다(없거나 이미지가 아니다): {e}")
            nm = p.get("번호_설명")
            if nm and loc.badge is not None and loc.note is not None:
                parsed = _numbered(subs.get("번호 설명", []))
                if isinstance(parsed, str):
                    problem(label, "번호 설명", parsed)
                    parsed = []
                elif not parsed:
                    problem(label, "번호 설명", "`#### 번호 설명` 이 없거나 비어 있다")
                bottom = int(nm["아래"]) if nm.get("아래") else (
                    loc.pic.top + loc.pic.height if loc.pic is not None else prs.slide_height - 400000)
                chunks = _plan_chunks(loc, parsed, bottom, lambda m, label=label: warn(label, "번호 설명", "est-overflow", m))
                if len(chunks) > 1:
                    warn(label, "번호 설명", "pages-split",
                         f"번호 설명 {len(parsed)}개를 {len(chunks)}쪽으로 나눴다(글 높이 추정). 그림·제목은 쪽마다 반복한다")
                for c in chunks:
                    out.append({**ctx, "chunk": c})
            else:
                out.append(ctx)
        plan[name] = out

    # 목차의 목록 제목 — 반복 쪽의 칸 값(화면마다 한 줄)
    for p in pages:
        lm = p.get("목록")
        if not lm or norm(p["이름"]) not in plan:
            continue
        key = norm(lm["제목"])
        titles: list[str] = []
        for q in pages:
            if q.get("반복"):
                seen: set[str] = set()
                for c in plan.get(norm(q["이름"]), []):
                    if c["label"] not in seen:
                        seen.add(c["label"])
                        titles.append(c["values"].get(key, ""))
        for c in plan[norm(p["이름"])]:
            c["titles"] = [t for t in titles if t]

    if problems:
        return None, problems, warnings

    # 3) 쪽 만들기: 견본 쪽을 첫 쪽으로 쓰고 나머지는 변경 전에 복제한다
    made: dict[str, list] = {}
    for p in pages:
        name = norm(p["이름"])
        sample = slides[p["견본"] - 1]
        count = max(len(plan.get(name, [])), 1)
        made[name] = [sample] + [_dup(prs, sample) for _ in range(count - 1)]
    for p in pages:
        name = norm(p["이름"])
        contexts = plan.get(name) or [{"values": {}, "seq": 1, "chunk": []}]
        for slide, ctx in zip(made[name], contexts):
            try:
                _fill_page(slide, p, prs, {**ctx, "warn": lambda pg, it, kind, why: warn(pg, it, kind, why)})
            except _PageRejected as e:
                for item, reason in e.args[0]:
                    problem(name, item, reason)
    if problems:
        return None, problems, warnings
    ordered = [s for p in pages for s in made[norm(p["이름"])]]
    for i in _order_slides(prs, ordered):
        warn("", f"슬라이드 {i}", "slide-dropped", f"양식의 슬라이드 {i} 은 지도가 가리키지 않아 산출물에서 뺐다")
    warn("", "전체", "layout-unchecked", LAYOUT_NOTE)
    buf = io.BytesIO()
    prs.save(buf)
    unique = [w for i, w in enumerate(warnings) if w not in warnings[:i]]  # 이어진 쪽마다 같은 경고가 반복된다
    return buf.getvalue(), [], unique


# ---------------------------------------------------------------- 도형 목록

def dump(path: Path) -> str:
    """양식 pptx 의 도형 목록. 쪽마다 도형 종류·글자·위치(EMU)와, 그 글자 앞부분이 `찾기` 로 유일한지를 보인다."""
    prs = Presentation(str(path))
    lines = [f"{path.name}: 슬라이드 {len(prs.slides)}장, 크기 {prs.slide_width}x{prs.slide_height} EMU"]
    for i, slide in enumerate(prs.slides, 1):
        lines.append(f"-- 슬라이드 {i} (배치: {slide.slide_layout.name})")
        pics = 0
        for s in slide.shapes:
            kind = _kind(s)
            if kind == "그림":
                pics += 1
                kind = f"그림 {pics}번째"
            where = f"위치 {s.left},{s.top} 크기 {s.width}x{s.height}" if s.left is not None else "위치 없음"
            text = _shape_text(s)
            note = ""
            if text:
                cand = text[:30]
                n = len(find_all(slide, cand))
                verdict = "유일" if n == 1 else f"중복 {n}개 — 더 긴 앞부분이 필요"
                note = f"  글자 {text[:40]!r}  찾기 {cand!r}: {verdict}"
            if _has_fld(s):
                note += "  (자동 쪽 번호 — 산출이 그대로 둔다)"
            lines.append(f"  {s.shape_id:>4} {kind:<8} {where}{note}")
    lines.append("글자는 앞부분 일치로 찾는다(공백은 하나로 본다). 위 `찾기` 는 그 도형 글자의 앞 30자다 — 지도에는 필요한 만큼 줄여 쓴다.")
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> int:
    for stream in (sys.stdout, sys.stderr):
        reconfigure = getattr(stream, "reconfigure", None)
        if reconfigure is not None:
            reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(prog="python -m checker.export_pptx", description="양식 pptx 의 도형 목록을 낸다(쓰기 없음).")
    parser.add_argument("--dump", type=Path, required=True, metavar="양식.pptx")
    args = parser.parse_args(argv)
    try:
        sys.stdout.write(dump(args.dump) + "\n")
    except Exception as e:
        sys.stdout.write(f"양식을 읽지 못했다(검사 불능): {type(e).__name__}: {e}\n")
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
