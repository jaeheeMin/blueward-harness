"""작업본(md) → 양식 xlsx 산출(`checker.export`)을 검증한다(이슈 #192).

픽스처 xlsx 는 커밋하지 않고 테스트 안에서 만든다. openpyxl 로 기본 통합문서를 만든 뒤 zip 수준으로
가짜 drawing 파트를 끼워 넣어, 산출이 손대지 않은 zip 항목을 바이트·순서 그대로 두는지 본다.
"""
from __future__ import annotations

import datetime as dt
import io
import json
import os
import re
import subprocess
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import PatternFill

from checker import export

MAP = """\
양식: 시험양식
설명: 시험용
템플릿: ../시험.xlsx
검사_규칙: ../../rules/test.yaml
산출:
  폴더: docs/out
  파일명: "T_{ID}_{제목}_v{버전}.xlsx"
  자리:
    버전: {표: 이력, 열: Ver, 행: 마지막}
    ID: {칸: 코드, 떼기: "X-"}
시트:
  - 이름: Info
    칸:
      - 항목: 코드
        위치: B2
        필수: true
        예시: SAMPLE
      - 항목: 제목
        위치: B3
        필수: true
        형식: "제목 : {값}"
      - 항목: 자동값
        위치: B4
        자동: true
      - 항목: 메모
        위치: B5
        예시: 옛 메모
      - 항목: 수량
        위치: B6
    표:
      - 이름: 이력
        머리글_행: 8
        시작_행: 9
        열:
          Ver: A
          Date: B
          Desc: C
        최대_행: 3
        필수: true
    서술:
      - 이름: 설명
        위치: A14:A16
  - 이름: Log
    표:
      - 이름: 로그
        머리글_행: 1
        시작_행: 2
        열:
          "No": A
          Text: B
        여러_줄: 행
  - 이름: Pair
    표:
      - 이름: 왼쪽
        머리글_행: 1
        시작_행: 2
        열: {L1: A, L2: B}
        짝: 오른쪽
        최대_행: 2
      - 이름: 오른쪽
        머리글_행: 1
        시작_행: 2
        열: {R1: D, R2: E}
        짝: 왼쪽
        최대_행: 2
"""

RULES = """\
관할: "docs/out/**"
규칙:
  - 종류: filename
    패턴: '^T_[0-9A-Za-z]+_.+_v\\d+\\.\\d+\\.xlsx$'
"""

SPEC = """\
# DEV-001 시험

## 근거 요구사항

REQ-001

## 양식 항목

### 시험양식 · Info

| 항목 | 값 |
|---|---|
| 코드 | X-01 |
| 제목 | 시험 건 1 |
| 자동값 | 자동 |
| 메모 | 미정 |
| 수량 | 12 |

#### 이력

| Ver | Date | Desc |
|---|---|---|
| 1.0 | 2026-10-06 | 최초 & <시작> |
| 1.1 | 2026-10-08 | 수정<br>두 줄 |

#### 설명

1. 첫째
2. 둘째

### 시험양식 · Log

#### 로그

| No | Text |
|---|---|
| 1 | a<br>b<br>c |
| 2 | d |
| 3 | e |
| 4 | f |

### 시험양식 · Pair

#### 왼쪽

| L1 | L2 |
|---|---|
| a | b |
| c | d |

#### 오른쪽

| R1 | R2 |
|---|---|
| e | f |
| g | h |

## 변경 이력
"""

DRAWING = b'<?xml version="1.0"?><xdr:wsDr xmlns:xdr="x"><xdr:sp/><xdr:pic></xdr:pic></xdr:wsDr>'


def _make_template(path: Path) -> None:
    wb = openpyxl.Workbook()
    info = wb.active
    info.title = "Info"
    info["A2"], info["B2"] = "코드", "SAMPLE"
    info["B4"] = '=B2&"x"'
    info["B5"] = "옛 메모"
    for r, v in enumerate(["0.1", "0.2", "0.3"], start=9):
        info.cell(r, 1, v)
        info.cell(r, 2, dt.date(2020, 1, r)).number_format = "yyyy-mm-dd"
        info.cell(r, 3, "샘플 내용")
    info["A8"], info["B8"], info["C8"] = "Ver", "Date", "Desc"
    info["A14"], info["A15"] = "샘플1", "샘플2"
    info["A20"] = "손대지 않을 칸"
    log = wb.create_sheet("Log")
    log["A1"], log["B1"] = "No", "Text"
    fill = PatternFill("solid", fgColor="FFFF00")
    for r in (2, 3, 4):
        log.cell(r, 1, f"s{r}")
        log.cell(r, 2, f"샘플{r}")
    log["A4"].fill = fill
    log["B4"].fill = fill
    pair = wb.create_sheet("Pair")
    for c, h in zip("ABDE", ["L1", "L2", "R1", "R2"]):
        pair[f"{c}1"] = h
    pair["A2"], pair["D2"] = "old", "old"
    buf = io.BytesIO()
    wb.save(buf)
    # 가짜 drawing 파트를 끼운다(내용은 산출이 건드리지 않아야 한다).
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as zi, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename == "[Content_Types].xml":
                data = data.replace(
                    b"</Types>",
                    b'<Override PartName="/xl/drawings/drawing1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>',
                )
            zo.writestr(it.filename, data)
        zo.writestr("xl/drawings/drawing1.xml", DRAWING)
        zo.writestr("xl/drawings/_rels/drawing1.xml.rels", b'<?xml version="1.0"?><Relationships/>')
        # Pair 시트(sheet3)가 그 drawing 을 쓴다.
        zo.writestr(
            "xl/worksheets/_rels/sheet3.xml.rels",
            '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" '
            'Target="../drawings/drawing1.xml"/></Relationships>'.encode("utf-8"),
        )


@pytest.fixture
def proj(tmp_path: Path) -> Path:
    (tmp_path / "templates" / "forms").mkdir(parents=True)
    (tmp_path / "rules").mkdir()
    (tmp_path / "templates" / "forms" / "시험양식.yaml").write_text(MAP, encoding="utf-8")
    (tmp_path / "rules" / "test.yaml").write_text(RULES, encoding="utf-8")
    _make_template(tmp_path / "templates" / "시험.xlsx")
    return tmp_path


def _spec(proj: Path, text: str = SPEC) -> Path:
    p = proj / "spec.md"
    p.write_text(text, encoding="utf-8")
    return p


def _run(proj: Path, capsys, text: str = SPEC, *extra: str):
    code = export.main(["--spec", str(_spec(proj, text)), "--root", str(proj), *extra])
    out = capsys.readouterr().out
    assert out.count("\n") == 1, "stdout 은 JSON 한 줄이어야 한다"
    return code, json.loads(out)


def _assert_sheets_sane(path: Path) -> None:
    """well-formed 에 더해 행·셀이 오름차순이고 셀 r 이 행 번호와 맞는지 본다."""
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with zipfile.ZipFile(path) as zf:
        for n in zf.namelist():
            if n.startswith("xl/worksheets/sheet") and n.endswith(".xml"):
                last_row = 0
                for row in ET.fromstring(zf.read(n)).iter(ns + "row"):
                    r = int(row.get("r"))
                    assert r > last_row, f"{n}: 행 순서"
                    last_row, last_col = r, 0
                    for c in row.findall(ns + "c"):
                        m = re.fullmatch(r"([A-Z]+)(\d+)", c.get("r"))
                        assert m and int(m.group(2)) == r, f"{n}: 셀 {c.get('r')} 가 {r}행에 있다"
                        col = export.col2n(m.group(1))
                        assert col > last_col, f"{n}: {r}행 셀 순서"
                        last_col = col


def _produced(proj: Path, report: dict) -> Path:
    assert len(report["files"]) == 1
    path = proj / report["files"][0]
    _assert_sheets_sane(path)
    return path


def test_산출이_성공하고_JSON_모양이_맞다(proj, capsys):
    code, report = _run(proj, capsys)
    assert code == 0
    assert report["status"] == "ok"
    assert report["problems"] == []
    assert report["warnings"] == []
    assert report["next"].startswith("다음:")
    assert report["files"] == ["docs/out/T_01_시험 건 1_v1.1.xlsx"]
    assert _produced(proj, report).exists()


def test_파일명_자리_표_마지막_행과_떼기(proj, capsys):
    _, report = _run(proj, capsys)
    # ID: 코드 X-01 에서 `X-` 를 뗀 01, 버전: 이력 표 Ver 의 마지막 행 1.1
    assert Path(report["files"][0]).name.startswith("T_01_")
    assert Path(report["files"][0]).name.endswith("_v1.1.xlsx")


def test_칸_형식_자동_샘플_지우기(proj, capsys):
    _, report = _run(proj, capsys)
    ws = openpyxl.load_workbook(_produced(proj, report))["Info"]
    assert ws["B2"].value == "X-01"
    assert ws["B3"].value == "제목 : 시험 건 1"
    assert ws["B4"].value == '=B2&"x"'  # 자동 칸은 쓰지 않는다
    assert ws["B5"].value is None  # 예시 칸이 미정 → 샘플 지움
    assert ws["B6"].value == 12  # 정수는 숫자로
    assert ws["A20"].value == "손대지 않을 칸"  # 지도에 없는 칸은 그대로


def test_표_날짜_최대행_샘플_지우기(proj, capsys):
    _, report = _run(proj, capsys)
    ws = openpyxl.load_workbook(_produced(proj, report))["Info"]
    assert [ws.cell(9, c).value for c in (1, 2, 3)] == ["1.0", dt.datetime(2026, 10, 6), "최초 & <시작>"]
    assert ws["C10"].value == "수정\n두 줄"  # 여러_줄 셀 기본: 셀 안 줄바꿈
    assert [ws.cell(11, c).value for c in (1, 2, 3)] == [None, None, None]  # 채우지 않은 샘플 행


def test_서술_범위는_줄마다_첫_열의_한_행씩_쓰고_나머지_샘플을_지운다(proj, capsys):
    _, report = _run(proj, capsys)
    ws = openpyxl.load_workbook(_produced(proj, report))["Info"]
    assert [ws[f"A{r}"].value for r in (14, 15, 16)] == ["1. 첫째", "2. 둘째", None]


def test_서술_줄은_빈_줄을_빼고_br도_줄로_나눈다(proj, capsys):
    text = SPEC.replace("1. 첫째\n2. 둘째", "1. a<br>2. b\n\n3. c")
    _, report = _run(proj, capsys, text)
    ws = openpyxl.load_workbook(_produced(proj, report))["Info"]
    assert [ws[f"A{r}"].value for r in (14, 15, 16)] == ["1. a", "2. b", "3. c"]


def test_서술이_범위_행_수를_넘으면_거부(proj, capsys):
    code, report = _run(proj, capsys, SPEC.replace("1. 첫째\n2. 둘째", "1\n2\n3\n4"))
    assert code == 1
    assert any(p["item"] == "설명" and "넘는다" in p["reason"] for p in report["problems"])


def test_서술_위치가_한_셀이면_줄바꿈_포함_한_셀(proj, capsys):
    (proj / "templates" / "forms" / "시험양식.yaml").write_text(MAP.replace("A14:A16", "A14"), encoding="utf-8")
    _, report = _run(proj, capsys)
    ws = openpyxl.load_workbook(_produced(proj, report))["Info"]
    assert ws["A14"].value == "1. 첫째\n2. 둘째"
    assert ws["A15"].value == "샘플2"  # 지도 위치(A14) 밖이라 건드리지 않는다


def test_여러_줄_행_표와_새_행_스타일_복제(proj, capsys):
    _, report = _run(proj, capsys)
    wb = openpyxl.load_workbook(_produced(proj, report))
    ws = wb["Log"]
    got = [(ws.cell(r, 1).value, ws.cell(r, 2).value) for r in range(2, 9)]
    assert got == [(1, "a"), (None, "b"), (None, "c"), (2, "d"), (3, "e"), (4, "f"), (None, None)]
    # 템플릿 마지막 데이터 행(4행)의 서식이 새 행(5~7행)에 복제된다.
    for r in (5, 6, 7):
        assert ws.cell(r, 1).fill.fgColor.rgb == ws["A4"].fill.fgColor.rgb == "00FFFF00"
    assert ws.cell(8, 1).fill.fgColor.rgb != "00FFFF00"


def test_짝_표_채우기(proj, capsys):
    _, report = _run(proj, capsys)
    ws = openpyxl.load_workbook(_produced(proj, report))["Pair"]
    assert [ws.cell(2, c).value for c in (1, 2, 4, 5)] == ["a", "b", "e", "f"]
    assert [ws.cell(3, c).value for c in (1, 2, 4, 5)] == ["c", "d", "g", "h"]


def test_손대지_않은_zip_항목은_바이트와_순서가_그대로(proj, capsys):
    _, report = _run(proj, capsys)
    tpl = zipfile.ZipFile(proj / "templates" / "시험.xlsx")
    out = zipfile.ZipFile(_produced(proj, report))
    assert [i.filename for i in out.infolist()] == [i.filename for i in tpl.infolist()]
    changed = {n for n in tpl.namelist() if tpl.read(n) != out.read(n)}
    sheets = {n for n in tpl.namelist() if n.startswith("xl/worksheets/sheet")}
    assert changed <= sheets | {"xl/workbook.xml"}
    assert out.read("xl/drawings/drawing1.xml") == DRAWING
    assert out.read("xl/styles.xml") == tpl.read("xl/styles.xml")


def test_시트_XML이_well_formed이고_수식_통합문서는_전체_재계산(proj, capsys):
    _, report = _run(proj, capsys)
    out = zipfile.ZipFile(_produced(proj, report))
    for n in out.namelist():
        if n.startswith("xl/worksheets/sheet") and n.endswith(".xml"):
            ET.fromstring(out.read(n))
    wb = out.read("xl/workbook.xml").decode("utf-8")
    ET.fromstring(wb)
    assert 'fullCalcOnLoad="1"' in wb


def test_산출물이_openpyxl로_다시_읽히고_파일명_규칙을_통과한다(proj, capsys):
    _, report = _run(proj, capsys)
    wb = openpyxl.load_workbook(_produced(proj, report))
    assert wb.sheetnames == ["Info", "Log", "Pair"]


def test_양식을_지정하면_그_양식만(proj, capsys):
    code, report = _run(proj, capsys, SPEC, "--form", "시험양식")
    assert code == 0 and len(report["files"]) == 1


def test_필수_칸이_미정이면_거부하고_아무것도_쓰지_않는다(proj, capsys):
    code, report = _run(proj, capsys, SPEC.replace("| 제목 | 시험 건 1 |", "| 제목 | 미정 |"))
    assert code == 1
    assert report["status"] == "rejected" and report["files"] == []
    assert {"form": "시험양식", "sheet": "Info", "item": "제목", "reason": "필수 칸이 미정이거나 비어 있다"} in report["problems"]
    assert not (proj / "docs").exists()


def test_필수_표가_비면_거부(proj, capsys):
    text = SPEC.replace("| 1.0 | 2026-10-06 | 최초 & <시작> |\n| 1.1 | 2026-10-08 | 수정<br>두 줄 |\n", "|  |  |  |\n")
    code, report = _run(proj, capsys, text)
    assert code == 1
    assert any(p["item"] == "이력" for p in report["problems"])


def test_최대_행_초과는_거부(proj, capsys):
    rows = "".join(f"| 1.{i} | 2026-10-0{i} | x |\n" for i in range(1, 5))
    text = SPEC.replace("| 1.0 | 2026-10-06 | 최초 & <시작> |\n| 1.1 | 2026-10-08 | 수정<br>두 줄 |\n", rows)
    code, report = _run(proj, capsys, text)
    assert code == 1
    assert any("최대 3행" in p["reason"] for p in report["problems"])


def test_짝_표_행_수가_다르면_거부(proj, capsys):
    code, report = _run(proj, capsys, SPEC.replace("| g | h |\n", ""))
    assert code == 1
    assert any(p["item"] in ("왼쪽", "오른쪽") and "행 수" in p["reason"] for p in report["problems"])


def test_금지_문자가_파일명에_들어가면_거부(proj, capsys):
    code, report = _run(proj, capsys, SPEC.replace("시험 건 1", "a/b"))
    assert code == 1
    assert any(p["item"] == "파일명" and "쓸 수 없는 문자" in p["reason"] for p in report["problems"])


def test_파일명_자리_값이_없으면_거부(proj, capsys):
    text = SPEC.replace("| 1.0 | 2026-10-06 | 최초 & <시작> |\n| 1.1 | 2026-10-08 | 수정<br>두 줄 |\n", "|  |  |  |\n")
    code, report = _run(proj, capsys, text)
    assert code == 1
    assert any("{버전}" in p["item"] for p in report["problems"])


def test_이미_있는_파일은_덮지_않고_거부(proj, capsys):
    code, report = _run(proj, capsys)
    assert code == 0
    produced = _produced(proj, report)
    before = produced.read_bytes()
    code, report = _run(proj, capsys)
    assert code == 1
    # git 저장소가 아니라 main 에 있는 판인지 판단하지 못한다 — 덮지 않고 사유를 밝힌다
    assert "판단하지 못" in report["problems"][0]["reason"] and "통과 아님" in report["problems"][0]["reason"]
    assert produced.read_bytes() == before


def test_파일명이_검사_규칙_패턴과_다르면_거부(proj, capsys):
    (proj / "rules" / "test.yaml").write_text(RULES.replace("^T_", "^Z_"), encoding="utf-8")
    code, report = _run(proj, capsys)
    assert code == 1
    assert any("패턴" in p["reason"] for p in report["problems"])


def test_지도에_없는_항목은_읽지_못함_종료코드_2(proj, capsys):
    code, report = _run(proj, capsys, SPEC.replace("| 수량 | 12 |", "| 없는칸 | 1 |"))
    assert code == 2
    assert report["status"] == "error" and report["files"] == []
    assert report["problems"][0]["item"] == "없는칸"


def test_지도에_없는_시트_표_열은_종료코드_2(proj, capsys):
    assert _run(proj, capsys, SPEC.replace("시험양식 · Pair", "시험양식 · 없는시트"))[0] == 2
    assert _run(proj, capsys, SPEC.replace("#### 로그", "#### 없는표"))[0] == 2
    assert _run(proj, capsys, SPEC.replace("| No | Text |", "| No | 없는열 |"))[0] == 2


def test_지도나_템플릿을_못_읽으면_종료코드_2(proj, capsys):
    (proj / "templates" / "시험.xlsx").unlink()
    code, report = _run(proj, capsys)
    assert code == 2 and report["status"] == "error"
    (proj / "templates" / "forms" / "시험양식.yaml").unlink()
    assert _run(proj, capsys)[0] == 2


def test_작업본에_양식_항목_절이_없으면_종료코드_2(proj, capsys):
    assert _run(proj, capsys, "# 제목\n\n## 개요\n")[0] == 2


def test_해당_없음_시트는_예시와_표_샘플을_비운다(proj, capsys):
    text = SPEC.replace("시험양식 · Pair\n", "시험양식 · Pair\n\n해당 없음\n").split("#### 왼쪽")[0] + "\n## 변경 이력\n"
    code, report = _run(proj, capsys, text)
    assert code == 0
    ws = openpyxl.load_workbook(_produced(proj, report))["Pair"]
    assert ws["A2"].value is None and ws["D2"].value is None
    # 지도 밖 샘플(머리글 셀, 그림·도형 2개)은 남기고 알린다.
    warns = [w for w in report["warnings"] if w["sheet"] == "Pair"]
    assert len(warns) == 1 and "그림·도형 2개" in warns[0]["reason"] and "엑셀에서 확인" in warns[0]["reason"]


def test_새_행이_아래_다른_항목과_겹치면_거부(proj, capsys):
    # Info 시트의 이력(최대_행 있음)은 해당 없고, Log 표 아래에 서술을 둔 지도로 바꿔 본다.
    map_text = MAP.replace(
        "        여러_줄: 행\n",
        "        여러_줄: 행\n    서술:\n      - 이름: 아래\n        위치: A30\n",
    )
    (proj / "templates" / "forms" / "시험양식.yaml").write_text(map_text, encoding="utf-8")
    code, report = _run(proj, capsys)
    assert code == 1
    assert any(p["item"] == "로그" and "아래에 다른 항목" in p["reason"] for p in report["problems"])


def _set_map(proj: Path, text: str) -> None:
    (proj / "templates" / "forms" / "시험양식.yaml").write_text(text, encoding="utf-8")


def _nothing_written(proj: Path) -> bool:
    return not any(p.is_file() for p in (proj / "docs").rglob("*")) if (proj / "docs").exists() else True


def test_산출_폴더가_root_밖이면_종료코드_2(proj, capsys):
    outside = proj.parent / "outside-out"
    _set_map(proj, MAP.replace("폴더: docs/out", "폴더: ../outside-out"))
    code, report = _run(proj, capsys)
    assert code == 2 and report["status"] == "error" and "밖" in report["problems"][0]["reason"]
    assert not outside.exists()
    _set_map(proj, MAP.replace("폴더: docs/out", f'폴더: "{outside.as_posix()}"'))
    code, report = _run(proj, capsys)
    assert code == 2 and not outside.exists()


def test_템플릿_검사_규칙_경로가_root_밖이면_종료코드_2(proj, capsys):
    _set_map(proj, MAP.replace("템플릿: ../시험.xlsx", "템플릿: ../../../시험.xlsx"))
    assert _run(proj, capsys)[0] == 2
    _set_map(proj, MAP.replace("검사_규칙: ../../rules/test.yaml", "검사_규칙: ../../../rules/test.yaml"))
    code, report = _run(proj, capsys)
    assert code == 2 and "밖" in report["problems"][0]["reason"]


def test_양식_이름에_경로_구분자가_있으면_종료코드_2(proj, capsys):
    for bad in ("../x", "a/b", "a\\b", "C:x"):
        assert _run(proj, capsys, SPEC.replace("### 시험양식 · Info", f"### {bad} · Info"))[0] == 2
    assert _run(proj, capsys, SPEC, "--form", "../시험양식")[0] == 2


def test_예상하지_못한_예외는_종료코드_2와_JSON(proj, capsys, monkeypatch):
    def boom(*a, **k):
        raise RuntimeError("boom")

    monkeypatch.setattr(export, "build_form", boom)
    code, report = _run(proj, capsys)
    assert code == 2 and report["status"] == "error" and report["files"] == []
    assert "boom" in report["problems"][0]["reason"]


def _two_form_spec() -> str:
    second = SPEC.split("### 시험양식 · Info", 1)[1].split("## 변경 이력")[0].replace("시험양식", "시험양식2")
    return SPEC.replace("## 변경 이력", "### 시험양식2 · Info" + second + "## 변경 이력")


def _add_second_form(proj: Path) -> None:
    second = MAP.replace("양식: 시험양식", "양식: 시험양식2").replace("docs/out", "docs/out2").replace(
        "T_{ID}_{제목}_v{버전}", "T_{ID}_{제목}_B_v{버전}"
    )
    (proj / "templates" / "forms" / "시험양식2.yaml").write_text(second, encoding="utf-8")


def test_두_양식을_함께_산출한다(proj, capsys):
    _add_second_form(proj)
    code, report = _run(proj, capsys, _two_form_spec())
    assert code == 0 and len(report["files"]) == 2


def test_둘째_양식_쓰기가_실패하면_아무_파일도_남기지_않는다(proj, capsys, monkeypatch):
    _add_second_form(proj)
    real = os.rename
    calls = []

    def flaky(src, dst):
        calls.append(dst)
        if len(calls) == 2:
            raise OSError("디스크 오류")
        real(src, dst)

    monkeypatch.setattr(os, "rename", flaky)
    code, report = _run(proj, capsys, _two_form_spec())
    assert code == 2 and report["status"] == "error" and "디스크 오류" in report["problems"][0]["reason"]
    assert len(calls) == 2  # 첫째는 최종 이름까지 갔다가 지워졌다
    assert _nothing_written(proj)
    assert not (proj / "docs").exists()


def test_셀_값이_32767자를_넘으면_거부(proj, capsys):
    code, report = _run(proj, capsys, SPEC.replace("| 메모 | 미정 |", "| 메모 | " + "가" * 32768 + " |"))
    assert code == 1 and report["files"] == []
    assert any(p["item"] == "메모" and "32767" in p["reason"] for p in report["problems"])
    code, _ = _run(proj, capsys, SPEC.replace("| 메모 | 미정 |", "| 메모 | " + "가" * 32767 + " |"))
    assert code == 0


def test_XML에_못_넣는_문자는_지우고_warnings로_알린다(proj, capsys):
    code, report = _run(proj, capsys, SPEC.replace("| 메모 | 미정 |", "| 메모 | a\x0bb\x01c" + chr(0xFFFE) + " |"))
    assert code == 0
    ws = openpyxl.load_workbook(_produced(proj, report))["Info"]
    assert ws["B5"].value == "abc"
    assert [w["item"] for w in report["warnings"]] == ["메모"] and "3개" in report["warnings"][0]["reason"]


def test_표가_끊긴_뒤_다시_나오면_종료코드_2(proj, capsys):
    text = SPEC.replace("| 1.0 | 2026-10-06 | 최초 & <시작> |\n| 1.1", "| 1.0 | 2026-10-06 | 최초 & <시작> |\n\n| 1.1")
    code, report = _run(proj, capsys, text)
    assert code == 2 and "끊긴" in report["problems"][0]["reason"]


def test_머리글_다음_줄이_구분선이_아니면_종료코드_2(proj, capsys):
    code, report = _run(proj, capsys, SPEC.replace("| Ver | Date | Desc |\n|---|---|---|\n", "| Ver | Date | Desc |\n"))
    assert code == 2 and "구분선" in report["problems"][0]["reason"]


def test_행의_셀_수가_머리글과_다르면_거부(proj, capsys):
    code, report = _run(proj, capsys, SPEC.replace("| 1.1 | 2026-10-08 | 수정<br>두 줄 |", "| 1.1 | 2026-10-08 |"))
    assert code == 1
    assert any(p["item"] == "이력" and "셀 수" in p["reason"] for p in report["problems"])


def test_같은_소절이_두_번이면_종료코드_2(proj, capsys):
    text = SPEC.replace("## 변경 이력", "### 시험양식 · Pair\n\n해당 없음\n\n## 변경 이력")
    code, report = _run(proj, capsys, text)
    assert code == 2 and "두 번" in report["problems"][0]["reason"]


def test_칸_표_머리글이_항목_값이_아니면_종료코드_2(proj, capsys):
    code, report = _run(proj, capsys, SPEC.replace("| 항목 | 값 |", "| 이름 | 값 |"))
    assert code == 2 and "항목" in report["problems"][0]["reason"]


def test_calcPr_삽입은_스키마_순서를_지킨다():
    base = '<workbook><sheets><sheet name="a" sheetId="1" r:id="rId1"/></sheets>{mid}<oleSize ref="A1"/></workbook>'
    with_names = export._with_full_calc(
        base.format(mid="<definedNames><definedName name=\"x\">a!A1</definedName></definedNames>")
    )
    assert with_names.index("</definedNames>") < with_names.index("<calcPr") < with_names.index("<oleSize")
    without = export._with_full_calc(base.format(mid=""))
    assert without.index("</sheets>") < without.index("<calcPr") < without.index("<oleSize")
    again = export._with_full_calc('<workbook><sheets/><calcPr calcId="1"/></workbook>')
    assert 'calcId="1" fullCalcOnLoad="1"/>' in again and again.count("calcPr") == 1


def test_자동_표시_없는_수식_칸에_쓰려_하면_종료코드_2(proj, capsys):
    _set_map(proj, MAP.replace("        위치: B4\n        자동: true\n", "        위치: B4\n"))
    code, report = _run(proj, capsys)
    assert code == 2 and "수식" in report["problems"][0]["reason"]
    assert _nothing_written(proj)


def test_검사_규칙의_pattern_별칭도_읽는다(proj, capsys):
    (proj / "rules" / "test.yaml").write_text(
        RULES.replace("패턴:", "pattern:").replace("^T_", "^Z_"), encoding="utf-8"
    )
    code, report = _run(proj, capsys)
    assert code == 1 and any("패턴" in p["reason"] for p in report["problems"])


# ---------------------------------------------------------------- SRS 원본(#228)
#
# 고객 양식은 저장소에 넣지 않으므로 Backlog 모양의 합성 xlsx 를 XML 로 직접 만든다. openpyxl 은 x14 확장
# (extLst) 유효성을 저장 때 잃기 때문에 이 픽스처를 만드는 데도 쓰지 않는다.

_NS = (
    'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
)

BACKLOG_MAP = """\
양식: Backlog
설명: SRS 요구사항 표를 Backlog 시트로
템플릿: ../Backlog.xlsx
산출:
  폴더: docs/out
  파일명: "BW_Backlog List_v{버전}.xlsx"
  자리:
    버전: {메타: 버전}
시트:
  - 이름: Backlog
    표:
      - 이름: Backlog
        원본: SRS
        머리글_행: 3
        시작_행: 4
        번호_열: Backlog ID
        md_전용: [출처]
        열:
          Backlog ID: A
          LoB: B
          Title: C
          Short Description: D
          Status: E
          Priority (PO): F
          Development RICEFW ID: G
          Flag: H
  - 이름: Values Drop Down
    설명: 드롭다운 목록 시트 — 건드리지 않는다
"""

SRS_MD = """\
# 시험 SRS

버전: 0.2

## 배경과 목표

x

## 요구사항

| ID | Title | LoB | Status | Priority (PO) | Development RICEFW ID | 출처 |
|---|---|---|---|---|---|---|
| REQ-002 | 결재 | Finance | New | 1. Must-Have | DEV-002 | 회의록 |
| REQ-003 | 보고서 | Sales | Planned | 2. Should-Have |  | 메일 |

### REQ-002 결재

#### Short Description

첫 줄
둘째 줄

### REQ-003 보고서

#### Short Description

보고서 한 줄

## 참조

없음
"""


def _backlog_template(
    path: Path,
    *,
    lob_formula: str = "'Values Drop Down'!$D$4:$D$7",
    hidden_row: bool = True,
    lob_sqref: str = "B4:B20 B30:B40",
    flag_attrs: str = 'showErrorMessage="1"',
    local_cells: bool = False,
) -> None:
    """Backlog 모양의 합성 xlsx. 표준 유효성(정의 이름·직접 목록)과 x14 유효성(다른 시트 범위)을 둘 다 가진다."""
    heads = ["Backlog ID", "LoB", "Title", "Short Description\n", "Status", "Priority\n(PO)", "Development RICEFW ID", "Flag"]
    head_cells = "".join(
        f'<c r="{c}3" s="1" t="inlineStr"><is><t xml:space="preserve">{h}</t></is></c>' for c, h in zip("ABCDEFGH", heads)
    )
    head_cells = head_cells.replace(">Status<", "> Status<")
    rows = [
        '<row r="1"><c r="A1" t="inlineStr"><is><t>Backlog</t></is></c>'
        + ('<c r="H1" t="inlineStr"><is><t>Alpha</t></is></c>' if local_cells else "")
        + "</row>",
        '<row r="2">' + ('<c r="H2" t="inlineStr"><is><t>Beta</t></is></c>' if local_cells else "") + "</row>"
        if local_cells else '<row r="2"/>',
        f'<row r="3" ht="40" customHeight="1">{head_cells}</row>',
    ]
    for r in range(4, 9):
        hid = ' hidden="1"' if (hidden_row and r == 5) else ""
        rows.append(f'<row r="{r}" spans="1:8" ht="30" customHeight="1"{hid}>' + "".join(f'<c r="{c}{r}" s="1"/>' for c in "ABCDEFGH") + "</row>")
    x14 = (
        '<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}" xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
        '<x14:dataValidations count="2" xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
        '<x14:dataValidation type="list" allowBlank="1" showErrorMessage="1"><x14:formula1>'
        f"<xm:f>{lob_formula}</xm:f></x14:formula1><xm:sqref>{lob_sqref}</xm:sqref></x14:dataValidation>"
        '<x14:dataValidation type="list" allowBlank="1" showErrorMessage="1"><x14:formula1>'
        "<xm:f>'Values Drop Down'!$B$4:$B$8</xm:f></x14:formula1><xm:sqref>E4:E20</xm:sqref></x14:dataValidation>"
        "</x14:dataValidations></ext></extLst>"
    )
    sheet1 = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><worksheet {_NS}>'
        '<dimension ref="A1:H8"/><sheetViews><sheetView workbookViewId="0"/></sheetViews>'
        f'<sheetData>{"".join(rows)}</sheetData>'
        '<autoFilter ref="A3:H6"></autoFilter>'
        '<dataValidations count="2">'
        '<dataValidation type="list" allowBlank="1" showErrorMessage="1" sqref="F4:F20"><formula1>Priorities</formula1></dataValidation>'
        f'<dataValidation type="list" allowBlank="1" {flag_attrs} sqref="H4:H20"><formula1>"Y,N"</formula1></dataValidation>'
        "</dataValidations>"
        '<pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" footer="0.3"/>'
        f"{x14}</worksheet>"
    )

    def inl(ref: str, text: str) -> str:
        return f'<c r="{ref}" t="inlineStr"><is><t>{text}</t></is></c>'

    vdd_rows = [
        f'<row r="4"><c r="A4" t="s"><v>0</v></c>{inl("B4", "New")}<c r="D4" t="s"><v>3</v></c></row>',
        f'<row r="5"><c r="A5" t="s"><v>1</v></c>{inl("B5", "Planned")}{inl("D5", "Sales")}</row>',
        f'<row r="6"><c r="A6" t="s"><v>2</v></c></row>',  # B6, D6 는 빈 칸(무시돼야 한다)
        f'<row r="7">{inl("B7", "In Process")}{inl("D7", "R&amp;D")}</row>',
        f'<row r="8">{inl("B8", "Done")}</row>',
    ]
    sheet2 = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><worksheet {_NS}>'
        f'<dimension ref="A4:D8"/><sheetData>{"".join(vdd_rows)}</sheetData></worksheet>'
    )
    sst = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><sst {_NS} count="4" uniqueCount="4">'
        "<si><t>1. Must-Have</t></si><si><t>2. Should-Have</t></si><si><t>3. Could-Have</t></si>"
        '<si><r><t>Fin</t></r><r><t>ance</t></r><rPh sb="0" eb="1"><t>무시</t></rPh></si></sst>'
    )
    workbook = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><workbook {_NS}><bookViews><workbookView/></bookViews>'
        '<sheets><sheet name="Backlog" sheetId="1" r:id="rId1"/><sheet name="Values Drop Down" sheetId="2" r:id="rId2"/></sheets>'
        '<definedNames><definedName name="_xlnm._FilterDatabase" localSheetId="0" hidden="1">Backlog!$A$3:$H$6</definedName>'
        "<definedName name=\"Priorities\">'Values Drop Down'!$A$4:$A$6</definedName></definedNames>"
        '<calcPr calcId="191029"/></workbook>'
    )
    rels_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    ofd = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    wb_rels = (
        f'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="{rels_ns}">'
        f'<Relationship Id="rId1" Type="{ofd}/worksheet" Target="worksheets/sheet1.xml"/>'
        f'<Relationship Id="rId2" Type="{ofd}/worksheet" Target="worksheets/sheet2.xml"/>'
        f'<Relationship Id="rId3" Type="{ofd}/styles" Target="styles.xml"/>'
        f'<Relationship Id="rId4" Type="{ofd}/sharedStrings" Target="sharedStrings.xml"/></Relationships>'
    )
    styles = (
        f'<?xml version="1.0" encoding="UTF-8"?><styleSheet {_NS}>'
        '<fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts>'
        '<fills count="2"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill></fills>'
        '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
        '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
        '<cellXfs count="2"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
        '<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0" applyAlignment="1"><alignment wrapText="1"/></xf></cellXfs></styleSheet>'
    )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        '<Override PartName="/xl/worksheets/sheet2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
        '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>'
    )
    root_rels = (
        f'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="{rels_ns}">'
        f'<Relationship Id="rId1" Type="{ofd}/officeDocument" Target="xl/workbook.xml"/></Relationships>'
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zo:
        zo.writestr("[Content_Types].xml", content_types)
        zo.writestr("_rels/.rels", root_rels)
        zo.writestr("xl/workbook.xml", workbook)
        zo.writestr("xl/_rels/workbook.xml.rels", wb_rels)
        zo.writestr("xl/styles.xml", styles)
        zo.writestr("xl/sharedStrings.xml", sst)
        zo.writestr("xl/worksheets/sheet1.xml", sheet1)
        zo.writestr("xl/worksheets/sheet2.xml", sheet2)


@pytest.fixture
def bproj(tmp_path: Path) -> Path:
    (tmp_path / "templates" / "forms").mkdir(parents=True)
    (tmp_path / "templates" / "forms" / "Backlog.yaml").write_text(BACKLOG_MAP, encoding="utf-8")
    _backlog_template(tmp_path / "templates" / "Backlog.xlsx")
    return tmp_path


def _srs_run(proj: Path, capsys, text: str = SRS_MD, *extra: str):
    srs = proj / "SRS.md"
    srs.write_text(text, encoding="utf-8")
    code = export.main(["--spec", str(srs), "--root", str(proj), "--form", "Backlog", *extra])
    out = capsys.readouterr().out
    assert out.count("\n") == 1
    return code, json.loads(out)


def _read_parts(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


def _sheet1(path: Path) -> str:
    return zipfile.ZipFile(path).read("xl/worksheets/sheet1.xml").decode("utf-8")


def _cell_values(path: Path) -> dict[str, object]:
    ws = openpyxl.load_workbook(path)["Backlog"]
    return {c.coordinate: c.value for row in ws.iter_rows(min_row=4, max_row=8) for c in row if c.value is not None}


def test_SRS_원본은_요구사항_표를_여러_행에_쓰고_REQ_번호를_숫자로_쓴다(bproj, capsys):
    code, report = _srs_run(bproj, capsys)
    assert code == 0 and report["status"] == "ok" and report["warnings"] == []
    assert report["files"] == ["docs/out/BW_Backlog List_v0.2.xlsx"]
    got = _cell_values(bproj / report["files"][0])
    assert got["A4"] == 2 and got["A5"] == 3  # REQ-002 → 2
    assert [got[f"{c}4"] for c in "BCEF"] == ["Finance", "결재", "New", "1. Must-Have"]
    assert got["G4"] == "DEV-002" and "G5" not in got
    assert got["D4"] == "첫 줄\n둘째 줄"  # 상세 절 → 셀, 여러 줄은 셀 안 줄바꿈
    assert got["D5"] == "보고서 한 줄"
    assert "회의록" not in _sheet1(bproj / report["files"][0])  # md 전용 열은 엑셀로 가지 않는다


def test_다음_단계_열은_지도_md_전용에_없어도_오류_없이_산출되고_엑셀에_안_간다(bproj, capsys):
    text = SRS_MD.replace("| 출처 |\n|---|---|---|---|---|---|---|", "| 출처 | 다음 단계 |\n|---|---|---|---|---|---|---|---|").replace(
        "| 회의록 |", "| 회의록 | 개발 |"
    ).replace("| 메일 |", "| 메일 | 설정 |")
    assert "md_전용: [출처]" in BACKLOG_MAP  # 지도는 다음 단계를 적지 않았다
    code, report = _srs_run(bproj, capsys, text)
    assert code == 0 and report["status"] == "ok" and report["warnings"] == []
    got = _cell_values(bproj / report["files"][0])
    assert got["A4"] == 2 and got["G4"] == "DEV-002" and "G5" not in got
    sheet = _sheet1(bproj / report["files"][0])
    assert "개발" not in sheet and "설정" not in sheet


def test_SRS_원본_산출물의_시트_XML이_well_formed(bproj, capsys):
    _, report = _srs_run(bproj, capsys)
    _assert_sheets_sane(bproj / report["files"][0])


def test_지도에도_md_전용에도_없는_열은_종료코드_2(bproj, capsys):
    code, report = _srs_run(bproj, capsys, SRS_MD.replace("출처", "엉뚱"))
    assert code == 2 and "엉뚱" in report["problems"][0]["reason"]


def test_상세_절_소제목이_지도_열에_없으면_종료코드_2(bproj, capsys):
    code, report = _srs_run(bproj, capsys, SRS_MD.replace("#### Short Description\n\n보고서", "#### Nope\n\n보고서"))
    assert code == 2 and "Nope" in report["problems"][0]["reason"]


def test_표에_없는_REQ의_상세_절은_종료코드_2(bproj, capsys):
    code, report = _srs_run(bproj, capsys, SRS_MD.replace("### REQ-003 보고서", "### REQ-009 보고서"))
    assert code == 2 and "REQ-009" in report["problems"][0]["reason"]


def test_표_칸과_상세_절_양쪽에_있으면_거부(bproj, capsys):
    text = SRS_MD.replace("| ID | Title |", "| ID | Short Description | Title |").replace("|---|---|---|", "|---|---|---|---|", 1)
    text = text.replace("| REQ-002 | 결재 |", "| REQ-002 | 표에도 있음 | 결재 |").replace("| REQ-003 | 보고서 |", "| REQ-003 |  | 보고서 |")
    code, report = _srs_run(bproj, capsys, text)
    assert code == 1
    assert any(p["item"] == "REQ-002" and "양쪽" in p["reason"] for p in report["problems"])


_NO_DETAIL_3 = "### REQ-003 보고서\n\n#### Short Description\n\n보고서 한 줄\n"


def test_ID가_REQ_숫자가_아니거나_겹치면_거부(bproj, capsys):
    dup = SRS_MD.replace("| REQ-003 | 보고서", "| REQ-002 | 보고서").replace(_NO_DETAIL_3, "")
    code, report = _srs_run(bproj, capsys, dup)
    assert code == 1 and any("두 줄 이상" in p["reason"] for p in report["problems"])
    bad = SRS_MD.replace("| REQ-003 | 보고서", "| 3 | 보고서").replace(_NO_DETAIL_3, "")
    code, report = _srs_run(bproj, capsys, bad)
    assert code == 1 and any("REQ-숫자" in p["reason"] for p in report["problems"])


def test_버전은_SRS_머리의_버전_줄에서_온다(bproj, capsys):
    code, report = _srs_run(bproj, capsys, SRS_MD.replace("버전: 0.2", "- 버전: v1.3"))
    assert code == 0 and report["files"] == ["docs/out/BW_Backlog List_v1.3.xlsx"]


def test_버전_줄이_없으면_거부하고_파일을_쓰지_않는다(bproj, capsys):
    code, report = _srs_run(bproj, capsys, SRS_MD.replace("버전: 0.2\n", ""))
    assert code == 1 and any("버전" in p["item"] for p in report["problems"])
    assert not (bproj / "docs").exists()


def test_SRS_지도는_양식을_고르지_않으면_종료코드_2로_알려_준다(bproj, capsys):
    srs = bproj / "SRS.md"
    srs.write_text(SRS_MD, encoding="utf-8")
    code = export.main(["--spec", str(srs), "--root", str(bproj)])
    report = json.loads(capsys.readouterr().out)
    assert code == 2 and "--form <지도>" in report["problems"][0]["reason"] and "Backlog" in report["problems"][0]["reason"]


def test_작업본_지도에_SRS를_먹이면_종료코드_2(proj, capsys):
    srs = proj / "SRS.md"
    srs.write_text(SRS_MD, encoding="utf-8")
    # 원본이 SRS 가 아닌 지도에는 --form 을 줘도 작업본 경로로 간다(`## 양식 항목` 이 없으니 읽지 못함)
    code = export.main(["--spec", str(srs), "--root", str(proj), "--form", "시험양식"])
    report = json.loads(capsys.readouterr().out)
    assert code == 2 and "양식 항목" in report["problems"][0]["reason"]


# -- 드롭다운 허용 값

def test_드롭다운_허용_값을_표준_x14_정의_이름_직접_목록에서_읽는다(bproj, capsys):
    code = export.main(["--root", str(bproj), "--form", "Backlog", "--list-values"])
    report = json.loads(capsys.readouterr().out)
    assert code == 0 and report["status"] == "ok" and report["warnings"] == []
    cols = {c["column"]: c for c in report["sheets"][0]["columns"]}
    assert report["sheets"][0]["sheet"] == "Backlog"
    # x14 + 다른 시트 범위: 빈 칸(D6)은 무시, sharedStrings 서식 있는 글(Fin+ance)은 이어 붙이고 윗주는 뺀다, &amp; 는 푼다
    assert cols["LoB"]["allowed"] == ["Finance", "Sales", "R&D"]
    assert cols["LoB"]["cells"] == "B4:B20 B30:B40" and cols["LoB"]["letter"] == "B"
    assert cols["Status"]["allowed"] == ["New", "Planned", "In Process", "Done"]  # 가운데 빈 칸 무시
    # 표준 + 정의 이름(Priorities → 'Values Drop Down'!$A$4:$A$6, sharedStrings)
    assert cols["Priority (PO)"]["allowed"] == ["1. Must-Have", "2. Should-Have", "3. Could-Have"]
    # 직접 적은 목록
    assert cols["Flag"]["allowed"] == ["Y", "N"]
    assert "Title" not in cols  # 유효성이 없는 열은 내지 않는다


def test_읽지_못하는_목록_수식은_허용_값_null과_사유로_알린다(bproj, capsys):
    _backlog_template(bproj / "templates" / "Backlog.xlsx", lob_formula="INDIRECT(&quot;A1&quot;)")
    code = export.main(["--root", str(bproj), "--form", "Backlog", "--list-values"])
    report = json.loads(capsys.readouterr().out)
    lob = next(c for c in report["sheets"][0]["columns"] if c["column"] == "LoB")
    # 읽지 못한 열이 있으면 통과처럼 0 으로 두지 않고 3(경고)
    assert code == 3 and report["status"] == "warning" and lob["allowed"] is None and lob["unreadable"]
    assert any(w["type"] == "dropdown-unchecked" and w["item"] == "LoB" for w in report["warnings"])


def test_목록_밖_값은_행_열_값_허용_값과_함께_경고하고_종료코드_3(bproj, capsys):
    text = SRS_MD.replace("| Finance |", "| Marketing |").replace("| 1. Must-Have |", "| 9. Nope |")
    code, report = _srs_run(bproj, capsys, text)
    assert code == 3 and report["status"] == "warning"
    assert (bproj / report["files"][0]).exists()  # 산출은 한다
    drops = [w for w in report["warnings"] if w["type"] == "dropdown"]
    lob = next(w for w in drops if w["column"] == "LoB")
    assert (lob["item"], lob["value"], lob["cell"], lob["row"]) == ("REQ-002", "Marketing", "B4", 4)
    assert lob["allowed"] == ["Finance", "Sales", "R&D"] and "허용 값: Finance, Sales, R&D" in lob["reason"]
    assert any(w["column"] == "Priority (PO)" and w["value"] == "9. Nope" for w in drops)
    assert len(drops) == 2  # 목록 안 값(Sales, New, Planned, 2. Should-Have)은 경고하지 않는다
    # 칸 값은 그대로 들어 있다
    assert _cell_values(bproj / report["files"][0])["B4"] == "Marketing"


def test_목록_안_값은_경고_없이_종료코드_0(bproj, capsys):
    code, report = _srs_run(bproj, capsys)
    assert code == 0 and not [w for w in report["warnings"] if w["type"] == "dropdown"]


def test_check는_파일을_쓰지_않고_같은_경고를_낸다(bproj, capsys):
    code, report = _srs_run(bproj, capsys, SRS_MD.replace("| Finance |", "| Marketing |"), "--check")
    assert code == 3 and report["status"] == "warning" and report["dry_run"] is True
    assert report["files"] == ["docs/out/BW_Backlog List_v0.2.xlsx"]
    assert not (bproj / "docs").exists()
    code, report = _srs_run(bproj, capsys, SRS_MD, "--check")
    assert code == 0 and report["dry_run"] is True and not (bproj / "docs").exists()


def test_목록을_읽지_못한_칸은_검사_불능_경고와_종료코드_3(bproj, capsys):
    _backlog_template(bproj / "templates" / "Backlog.xlsx", lob_formula="INDIRECT(&quot;A1&quot;)")
    code, report = _srs_run(bproj, capsys, SRS_MD.replace("| Finance |", "| Marketing |"))
    assert code == 3 and report["status"] == "warning"
    assert [w["type"] for w in report["warnings"]] == ["dropdown-unchecked"]
    assert "검사하지 못한 열이 있다" in report["next"] and "목록 밖" not in report["next"]
    assert (bproj / report["files"][0]).exists()


# -- 행 속성·서식·확장 유효성 보존

def test_쓰는_행의_hidden은_지운다(bproj, capsys):
    tpl = zipfile.ZipFile(bproj / "templates" / "Backlog.xlsx").read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert '<row r="5" spans="1:8" ht="30" customHeight="1" hidden="1">' in tpl
    _, report = _srs_run(bproj, capsys)
    out = _sheet1(bproj / report["files"][0])
    assert "hidden" not in out
    assert re.search(r'<row r="5"[^>]*>', out)


def test_쓰지_않는_행의_hidden은_그대로(bproj, capsys):
    # 요구사항 한 줄만 쓰면 5행은 쓰지 않는다
    text = SRS_MD.replace("| REQ-003 | 보고서 | Sales | Planned | 2. Should-Have |  | 메일 |\n", "").replace(
        "### REQ-003 보고서\n\n#### Short Description\n\n보고서 한 줄\n", ""
    )
    _, report = _srs_run(bproj, capsys, text)
    assert 'r="5"' in _sheet1(bproj / report["files"][0]) and 'hidden="1"' in _sheet1(bproj / report["files"][0])


def test_산출물에_x14_유효성과_표준_유효성이_바이트_그대로_남는다(bproj, capsys):
    _, report = _srs_run(bproj, capsys)
    tpl = zipfile.ZipFile(bproj / "templates" / "Backlog.xlsx").read("xl/worksheets/sheet1.xml").decode("utf-8")
    out = _sheet1(bproj / report["files"][0])
    ext = lambda x: re.search(r"<extLst>.*</extLst>", x, re.S).group(0)
    std = lambda x: re.search(r"<dataValidations count=.*?</dataValidations>", x, re.S).group(0)
    assert ext(out) == ext(tpl) and std(out) == std(tpl)
    assert out.count("<x14:dataValidation ") == 2
    # 목록 시트와 나머지 항목도 바이트 그대로
    a, b = zipfile.ZipFile(bproj / "templates" / "Backlog.xlsx"), zipfile.ZipFile(bproj / report["files"][0])
    for n in ("xl/worksheets/sheet2.xml", "xl/styles.xml", "xl/sharedStrings.xml"):
        assert a.read(n) == b.read(n)


def test_자동_필터는_쓴_행이_범위를_넘을_때만_늘린다(bproj, capsys):
    # 2건(4~5행)은 필터 A3:H6 안 — 필터와 숨은 이름 모두 그대로
    _, report = _srs_run(bproj, capsys)
    out = _read_parts(bproj / report["files"][0])
    assert '<autoFilter ref="A3:H6"></autoFilter>' in out["xl/worksheets/sheet1.xml"].decode("utf-8")
    assert "Backlog!$A$3:$H$6" in out["xl/workbook.xml"].decode("utf-8")
    # 5건(4~8행)은 넘는다 — 둘 다 8행까지
    rows = "".join(f"| REQ-{n:03d} | t{n} | Sales | New | 1. Must-Have |  | x |\n" for n in range(4, 7))
    text = SRS_MD.replace("\n### REQ-002", rows + "\n### REQ-002", 1)
    for f in (bproj / "docs" / "out").glob("*"):
        f.unlink()
    _, report = _srs_run(bproj, capsys, text)
    out = _read_parts(bproj / report["files"][0])
    assert '<autoFilter ref="A3:H8"></autoFilter>' in out["xl/worksheets/sheet1.xml"].decode("utf-8")
    wb = out["xl/workbook.xml"].decode("utf-8")
    assert "Backlog!$A$3:$H$8" in wb and "$H$6" not in wb


def test_다른_시트_드롭다운_목록_시트는_건드리지_않아도_경고가_없다(bproj, capsys):
    _, report = _srs_run(bproj, capsys)
    assert not any(w.get("sheet") == "Values Drop Down" for w in report["warnings"])


def test_SRS_원본_지도에_칸이나_원본_없는_표가_섞이면_종료코드_2(bproj, capsys):
    path = bproj / "templates" / "forms" / "Backlog.yaml"
    path.write_text(BACKLOG_MAP + "    칸:\n      - 항목: x\n        위치: A1\n", encoding="utf-8")
    # 두 번째 시트 아래에 칸이 붙는다
    code, report = _srs_run(bproj, capsys)
    assert code == 2 and "SRS" in report["problems"][0]["reason"]


# -- 줄 높이·목록 읽기·검사 대상 (#228 검토 지적)

def test_값을_쓴_행의_고정_줄_높이는_지우고_다른_행은_그대로(bproj, capsys):
    _, report = _srs_run(bproj, capsys)
    out = _sheet1(bproj / report["files"][0])
    row = lambda n: re.search(rf'<row r="{n}"[^>]*>', out).group(0)
    for n in (4, 5):
        assert "ht=" not in row(n) and "customHeight" not in row(n)
    for n in (6, 8):
        assert 'ht="30" customHeight="1"' in row(n)
    assert 'ht="40" customHeight="1"' in row(3)


def test_시트_이름_없는_같은_시트_범위_목록은_그_시트_기준으로_읽는다(bproj, capsys):
    _backlog_template(bproj / "templates" / "Backlog.xlsx", lob_formula="$H$1:$H$2", local_cells=True)
    code = export.main(["--root", str(bproj), "--form", "Backlog", "--list-values"])
    report = json.loads(capsys.readouterr().out)
    lob = next(c for c in report["sheets"][0]["columns"] if c["column"] == "LoB")
    assert code == 0 and lob["allowed"] == ["Alpha", "Beta"]


_WITH_FLAG = SRS_MD.replace("| 출처 |\n|---|---|---|---|---|---|---|", "| 출처 | Flag |\n|---|---|---|---|---|---|---|---|").replace(
    "| 회의록 |", "| 회의록 | Q |"
).replace("| 메일 |", "| 메일 | Q |")


def test_오류_메시지가_켜진_stop_유효성은_목록_밖_값을_경고한다(bproj, capsys):
    code, report = _srs_run(bproj, capsys, _WITH_FLAG)
    assert code == 3 and [w["column"] for w in report["warnings"] if w["type"] == "dropdown"] == ["Flag", "Flag"]


def test_정보_스타일이거나_오류_메시지가_꺼진_유효성은_검사하지_않는다(bproj, capsys):
    for attrs in ('errorStyle="information" showErrorMessage="1"', 'errorStyle="warning" showErrorMessage="1"', ""):
        for f in (bproj / "docs" / "out").glob("*") if (bproj / "docs").exists() else []:
            f.unlink()
        _backlog_template(bproj / "templates" / "Backlog.xlsx", flag_attrs=attrs)
        code, report = _srs_run(bproj, capsys, _WITH_FLAG)
        assert code == 0 and not [w for w in report["warnings"] if w["type"] == "dropdown"], attrs
    code = export.main(["--root", str(bproj), "--form", "Backlog", "--list-values"])
    flag = next(c for c in json.loads(capsys.readouterr().out)["sheets"][0]["columns"] if c["column"] == "Flag")
    assert flag["enforced"] is False and flag["allowed"] == ["Y", "N"]


def test_목록_범위가_실제로_비어_있으면_알림만_내고_종료코드에_영향_없다(bproj, capsys):
    _backlog_template(bproj / "templates" / "Backlog.xlsx", lob_formula="'Values Drop Down'!$F$4:$F$6")
    code, report = _srs_run(bproj, capsys)
    assert code == 0 and report["status"] == "ok"
    assert [w["type"] for w in report["warnings"]] == ["dropdown-empty"]
    assert "검사하지 못한" not in report["next"]


def test_정의_이름을_못_찾으면_읽지_못한_것이라_종료코드_3(bproj, capsys):
    _backlog_template(bproj / "templates" / "Backlog.xlsx", lob_formula="NoSuchName")
    code, report = _srs_run(bproj, capsys)
    assert code == 3 and [w["type"] for w in report["warnings"]] == ["dropdown-unchecked"]


def test_코드_펜스_안의_제목은_소제목으로_보지_않는다(bproj, capsys):
    text = SRS_MD.replace(
        "첫 줄\n둘째 줄", "첫 줄\n```\n#### Nope\n### REQ-009 x\n## 요구사항\n```\n둘째 줄"
    )
    code, report = _srs_run(bproj, capsys, text)
    assert code == 0, report
    assert _cell_values(bproj / report["files"][0])["D4"] == "첫 줄\n```\n#### Nope\n### REQ-009 x\n## 요구사항\n```\n둘째 줄"


def test_쓰는_행이_드롭다운_범위_밖이면_그_줄만_알린다(bproj, capsys):
    _backlog_template(bproj / "templates" / "Backlog.xlsx", lob_sqref="B4:B5 B30:B40")
    rows = "| REQ-004 | t4 | Sales | New | 1. Must-Have |  | x |\n"
    code, report = _srs_run(bproj, capsys, SRS_MD.replace("\n### REQ-002", rows + "\n### REQ-002", 1))
    miss = [w for w in report["warnings"] if w["type"] == "dropdown-missing"]
    assert code == 0 and len(miss) == 1
    assert miss[0]["row"] == 6 and miss[0]["item"] == "REQ-004"
    assert "6행" in miss[0]["reason"] and "이 줄엔 드롭다운이 없다" in miss[0]["reason"]


def test_대소문자만_다른_값은_힌트를_붙여_경고한다(bproj, capsys):
    code, report = _srs_run(bproj, capsys, SRS_MD.replace("| Finance |", "| finance |"))
    assert code == 3
    w = next(w for w in report["warnings"] if w["type"] == "dropdown")
    assert "대소문자만 다르다('Finance')" in w["reason"]


# ---------------------------------------------------------------- --form 경로, 미병합 판 다시 산출

def _form_results(proj: Path, capsys, form: str):
    code, report = _run(proj, capsys, SPEC, "--check", "--form", form)
    return code, report["files"], report["status"]


def test_form_이름과_경로_세_형태가_같은_결과(proj, capsys):
    by_name = _form_results(proj, capsys, "시험양식")
    assert by_name[0] == 0 and len(by_name[1]) == 1
    absolute = str((proj / "templates" / "forms" / "시험양식.yaml").resolve())
    for form in ("templates/forms/시험양식.yaml", absolute, "templates/forms/시험양식.yaml".replace("/", os.sep)):
        assert _form_results(proj, capsys, form) == by_name
    # 구분자 없는 `이름.yaml` 은 templates/forms 아래 파일 이름으로 본다
    assert _form_results(proj, capsys, "시험양식.yaml") == by_name


def test_form_경로가_없거나_밖이면_종료코드_2(proj, capsys):
    (proj / "other.yaml").write_text(MAP, encoding="utf-8")
    for form in ("templates/forms/없는양식.yaml", "other.yaml", "templates/other.yaml", "../x/templates/forms/시험양식.yaml"):
        code, report = _run(proj, capsys, SPEC, "--check", "--form", form)
        assert code == 2 and report["status"] == "error", form
        assert report["files"] == []


def test_list_values_도_form_경로를_받는다(proj, capsys):
    code = export.main(["--root", str(proj), "--form", "templates/forms/시험양식.yaml", "--list-values"])
    out = json.loads(capsys.readouterr().out)
    assert code in (0, 3) and out["form"] == "시험양식"
    code = export.main(["--root", str(proj), "--form", "templates/forms/없음.yaml", "--list-values"])
    assert code == 2
    capsys.readouterr()


def _git(cwd: Path, *args: str) -> None:
    subprocess.run(
        ["git", "-c", "user.name=t", "-c", "user.email=t@t", "-c", "commit.gpgsign=false", *args],
        cwd=cwd, check=True, capture_output=True,
    )


@pytest.fixture
def gproj(proj: Path, tmp_path_factory) -> Path:
    """origin(bare)을 가진 git 저장소. 산출물은 아직 없고 origin/main 에 초기 커밋만 있다."""
    bare = tmp_path_factory.mktemp("origin") / "o.git"
    subprocess.run(["git", "init", "--bare", "-b", "main", str(bare)], check=True, capture_output=True)
    _git(proj, "init", "-b", "main")
    _git(proj, "add", "-A")
    _git(proj, "commit", "-m", "init")
    _git(proj, "remote", "add", "origin", str(bare))
    _git(proj, "push", "origin", "main")
    return proj


def test_main_에_없는_기존_산출_파일은_같은_버전으로_덮어쓴다(gproj, capsys):
    code, report = _run(gproj, capsys)
    assert code == 0 and "replaced" not in report
    produced = _produced(gproj, report)
    produced.write_bytes(b"stale")
    code, report = _run(gproj, capsys)
    assert code == 0
    assert report["replaced"] is True and report["replaced_files"][0]["file"] == report["files"][0]
    assert "origin/main 에 없" in report["replaced_files"][0]["reason"]
    assert produced.read_bytes() != b"stale"
    _assert_sheets_sane(produced)
    assert not list(produced.parent.glob(".*.bak")) and not list(produced.parent.glob(".*.tmp"))


def test_main_에_있는_산출_파일은_거부한다(gproj, capsys):
    code, report = _run(gproj, capsys)
    produced = _produced(gproj, report)
    _git(gproj, "add", "-A")
    _git(gproj, "commit", "-m", "산출")
    _git(gproj, "push", "origin", "main")
    before = produced.read_bytes()
    code, report = _run(gproj, capsys)
    assert code == 1 and "replaced" not in report
    reason = report["problems"][0]["reason"]
    assert "main 에 이미 있는 판" in reason and "버전을 올려" in reason
    assert produced.read_bytes() == before


def test_origin_main_을_읽지_못하면_거부하고_사유를_밝힌다(gproj, capsys):
    code, report = _run(gproj, capsys)
    produced = _produced(gproj, report)
    _git(gproj, "remote", "remove", "origin")  # origin/main 참조가 사라진다
    before = produced.read_bytes()
    code, report = _run(gproj, capsys)
    assert code == 1
    reason = report["problems"][0]["reason"]
    assert "판단하지 못" in reason and "origin/main" in reason
    assert produced.read_bytes() == before


# -- 엑셀 행 높이 상한을 넘어 잘릴 긴 셀 경고 (#240)

def _set_layout(proj: Path, cols: str = "", merges: str = "") -> None:
    """Backlog 템플릿 시트에 열 너비(`<cols>`)와 병합 셀을 끼워 넣는다."""
    path = proj / "templates" / "Backlog.xlsx"
    parts = _read_parts(path)
    with zipfile.ZipFile(path) as z:
        infos = z.infolist()
    sheet = parts["xl/worksheets/sheet1.xml"].decode("utf-8")
    sheet = sheet.replace("<sheetData>", f"<cols>{cols}</cols><sheetData>") if cols else sheet
    sheet = sheet.replace('<autoFilter', f"<mergeCells>{merges}</mergeCells><autoFilter") if merges else sheet
    parts["xl/worksheets/sheet1.xml"] = sheet.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zo:
        for info in infos:
            zo.writestr(info.filename, parts[info.filename])


def _long_srs(desc: str) -> str:
    return SRS_MD.replace("첫 줄\n둘째 줄", desc)


def _tall(report: dict) -> list[dict]:
    return [w for w in report["warnings"] if w.get("type") == "cell-too-tall"]


def test_좁은_열의_긴_셀은_행_높이_상한을_넘어_경고하고_종료코드_3(bproj, capsys):
    _set_layout(bproj, '<col min="4" max="4" width="10" customWidth="1"/>')
    code, report = _srs_run(bproj, capsys, _long_srs("x" * 300))  # 10칸씩 30줄 = 450pt
    w = _tall(report)
    assert code == 3 and report["status"] == "warning" and len(w) == 1
    assert (w[0]["row"], w[0]["column"], w[0]["cell"], w[0]["req"]) == (4, "Short Description", "D4", "REQ-002")
    assert w[0]["est_lines"] == 30 and w[0]["est_height_pt"] == 450 and w[0]["limit_pt"] == 409
    assert (bproj / report["files"][0]).exists()  # 파일은 썼다


def test_전각_글자는_두_칸으로_세어_같은_길이도_더_높다(bproj, capsys):
    _set_layout(bproj, '<col min="4" max="4" width="20" customWidth="1"/>')
    code, report = _srs_run(bproj, capsys, _long_srs("가" * 300))  # 10자씩 30줄
    assert code == 3 and _tall(report)[0]["est_lines"] == 30
    (bproj / "docs" / "out" / "BW_Backlog List_v0.2.xlsx").unlink()
    code, report = _srs_run(bproj, capsys, _long_srs("x" * 300))  # 20자씩 15줄 = 225pt
    assert code == 0 and report["warnings"] == []


def test_명시한_줄바꿈도_줄_수에_센다(bproj, capsys):
    _set_layout(bproj, '<col min="4" max="4" width="100" customWidth="1"/>')
    code, report = _srs_run(bproj, capsys, _long_srs("\n".join(f"항목 {i}" for i in range(30))))
    w = _tall(report)
    assert code == 3 and len(w) == 1 and w[0]["est_lines"] == 30


def test_짧은_셀은_좁은_열이어도_경고_없다(bproj, capsys):
    _set_layout(bproj, '<col min="4" max="4" width="10" customWidth="1"/>')
    code, report = _srs_run(bproj, capsys)
    assert code == 0 and report["warnings"] == []


def test_병합_셀은_범위_너비의_합으로_센다(bproj, capsys):
    cols = '<col min="4" max="6" width="10" customWidth="1"/>'
    _set_layout(bproj, cols)
    code, report = _srs_run(bproj, capsys, _long_srs("x" * 300))  # D 한 칸 10 → 30줄
    assert code == 3
    (bproj / "docs" / "out" / "BW_Backlog List_v0.2.xlsx").unlink()
    _set_layout(bproj, cols, '<mergeCell ref="D4:F4"/>')  # D~F 합 30칸 → 10줄
    code, report = _srs_run(bproj, capsys, _long_srs("x" * 300))
    assert code == 0 and report["warnings"] == []


def test_줄바꿈_서식이_없는_셀은_한_줄로_보이므로_경고하지_않는다(bproj, capsys):
    path = bproj / "templates" / "Backlog.xlsx"
    parts = _read_parts(path)
    parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(b'r="D4" s="1"', b'r="D4" s="0"')
    with zipfile.ZipFile(path) as z:
        order = [i.filename for i in z.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zo:
        for n in order:
            zo.writestr(n, parts[n])
    _set_layout(bproj, '<col min="4" max="4" width="10" customWidth="1"/>')
    code, report = _srs_run(bproj, capsys, _long_srs("x" * 300))
    assert code == 0 and report["warnings"] == []


def test_check도_긴_셀을_경고하고_파일은_쓰지_않는다(bproj, capsys):
    _set_layout(bproj, '<col min="4" max="4" width="10" customWidth="1"/>')
    code, report = _srs_run(bproj, capsys, _long_srs("x" * 300), "--check")
    assert code == 3 and report["dry_run"] is True and len(_tall(report)) == 1
    assert not (bproj / "docs" / "out").exists()


def test_양식_밖_열이_여럿이면_하나도_빠뜨리지_않고_모두_알리며_멈춘다_316(bproj, capsys):
    text = SRS_MD.replace("| 출처 |\n|---|---|---|---|---|---|---|", "| 출처 | 엉뚱 | Scope |\n|---|---|---|---|---|---|---|---|---|")
    text = text.replace("| 회의록 |", "| 회의록 | a | b |").replace("| 메일 |", "| 메일 | c | d |")
    code, report = _srs_run(bproj, capsys, text)
    reason = report["problems"][0]["reason"]
    assert code == 2 and report["files"] == []
    assert "'엉뚱'" in reason and "'Scope'" in reason
    assert "허용:" in reason and "Backlog ID" not in reason.split("허용:")[0] and "다음: 양식 지도" in reason
    assert not (bproj / "docs").exists()  # 건너뛰고 산출하지 않는다


def test_양식_밖_상세_소제목이_여럿이면_REQ별로_모두_알리며_멈춘다_316(bproj, capsys):
    text = SRS_MD.replace("#### Short Description\n\n첫 줄", "#### Notes\n\n첫 줄").replace(
        "#### Short Description\n\n보고서", "#### Memo\n\n보고서"
    )
    code, report = _srs_run(bproj, capsys, text)
    reason = report["problems"][0]["reason"]
    assert code == 2 and "REQ-002: `#### Notes`" in reason and "REQ-003: `#### Memo`" in reason
    assert "다음: 양식 지도" in reason
    assert not (bproj / "docs").exists()


def test_작업본_양식_밖_항목이_여럿이면_모두_알리며_멈춘다_316(proj, capsys):
    spec = SPEC.replace("| 수량 | 12 |", "| 없는칸 | 1 |").replace("| No | Text |", "| No | 없는열 |")
    code, report = _run(proj, capsys, spec)
    reason = report["problems"][0]["reason"]
    assert code == 2 and report["files"] == []
    assert "없는칸" in reason or "없는열" in reason


def test_템플릿_시트_이름이_NFD여도_지도의_NFC_이름으로_찾아_산출한다(bproj, capsys):
    import unicodedata

    tpl = bproj / "templates" / "Backlog.xlsx"
    parts = _read_parts(tpl)
    nfd = unicodedata.normalize("NFD", "백로그")
    assert nfd != "백로그"
    parts["xl/workbook.xml"] = parts["xl/workbook.xml"].decode("utf-8").replace('name="Backlog"', f'name="{nfd}"').encode("utf-8")
    with zipfile.ZipFile(tpl, "w", zipfile.ZIP_DEFLATED) as zo:
        for n, data in parts.items():
            zo.writestr(n, data)
    mp = bproj / "templates" / "forms" / "Backlog.yaml"
    mp.write_text(mp.read_text(encoding="utf-8").replace("  - 이름: Backlog\n", "  - 이름: 백로그\n"), encoding="utf-8")
    code, report = _srs_run(bproj, capsys)
    assert code == 0, report
    ws = openpyxl.load_workbook(bproj / report["files"][0])[nfd]
    assert ws["A4"].value == 2
