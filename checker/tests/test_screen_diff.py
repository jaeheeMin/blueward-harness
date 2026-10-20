"""화면 그림의 지난 판 비교와 변경 표시(`checker.screen_diff`, `checker.screen_draw`)를 검증한다(이슈 #347).

산출은 origin/main 의 같은 경로 작업본과 지금 작업본의 `#### 화면 1 그림` 을 비교해 추가·수정·삭제를 번호와 겹침 스타일로 표시하고
지도 `변경:` 의 변경 내용 표를 채운다. git 이 필요한 시험은 임시 폴더에 `git init` 하고 가짜 `refs/remotes/origin/main` 을 만들어 쓴다.
"""
from __future__ import annotations

import subprocess
from pathlib import Path

import openpyxl
import pytest

from checker import screen_diff, screen_draw
from checker.tests.test_screen_draw import (  # noqa: F401  (proj 는 fixture)
    BTN, COND, MAP, RES, SPEC_HEAD, _map, _merges, _nothing_written, _run, _sheet, _template, proj, screen,
)

CHANGE = "        변경: {시작_행: 20, 최대_행: %d, 열: {번호: B, 구분: C, 대상: E, 내용: H}}\n"


@pytest.fixture(autouse=True)
def _no_previous_spec():
    """conftest 의 "첫 판으로 둔다" 덮개를 끈다 — 이 파일은 진짜 git 으로 지난 판을 읽는다."""


def _git(root: Path, *args: str) -> None:
    subprocess.run(["git", "-C", str(root), *args], check=True, capture_output=True)


def _origin(proj: Path, old: str | None, *, ref: bool = True) -> None:
    """proj 를 git 저장소로 만들고 old(작업본 본문)를 spec.md 로 커밋해 origin/main 으로 삼는다. old 가 None 이면 spec.md 가 없는 main."""
    _git(proj, "init", "-q")
    _git(proj, "config", "user.email", "t@example.com")
    _git(proj, "config", "user.name", "t")
    _git(proj, "config", "commit.gpgsign", "false")
    (proj / "README.md").write_text("x", encoding="utf-8")
    files = ["README.md"]
    if old is not None:
        (proj / "spec.md").write_bytes((SPEC_HEAD + old).encode("utf-8"))
        files.append("spec.md")
    _git(proj, "add", *files)
    _git(proj, "commit", "-q", "-m", "x")
    if ref:
        _git(proj, "update-ref", "refs/remotes/origin/main", "HEAD")


def _at(ws, text):
    for row in ws.iter_rows():
        for c in row:
            if c.value == text:
                return c
    raise AssertionError(f"{text!r} 칸이 없다")


def _fill(c) -> str:
    return c.fill.fgColor.rgb


def _table(ws, first: int = 20, n: int = 10) -> list[tuple]:
    rows = [tuple(ws[f"{c}{r}"].value for c in "BCEH") for r in range(first, first + n)]
    return [r for r in rows if any(r)]


def _tbl(name: str, rows: list[str], req: list[str] | None = None) -> str:
    head = "| 이름 | 필수 |\n|---|---|\n" if req is not None else f"| {name or '이름'} |\n|---|\n"
    if req is not None:
        return head + "".join(f"| {r} | {q} |\n" for r, q in zip(rows, req))
    return head + "".join(f"| {r} |\n" for r in rows)


OLD = screen()
YELLOW, GREY = "FFFFFF00", "FFD9D9D9"


# ---------------------------------------------------------------- 판정(순수 함수)

def test_삭제_항목은_옛_판에서_바로_앞에_있던_남은_항목_뒤에_둔다():
    merged = screen_diff.merge(
        ([(n, False) for n in "abcd"], [], [], False), ([(n, False) for n in "acx"], [], [], False),
    )
    assert [n for n, _ in merged.conds] == ["a", "b", "c", "d", "x"]  # b 는 a 뒤, d 는 c 뒤
    assert {k: v for k, v in merged.marks.items()} == {
        ("필터 바", "b"): ("삭제", 1), ("필터 바", "d"): ("삭제", 2), ("필터 바", "x"): ("추가", 3),
    }
    # 맨 앞 항목이 지워지면 소절 맨 앞에 남는다.
    m2 = screen_diff.merge(([], [], list("abc"), False), ([], [], list("bc"), False))
    assert m2.cols == ["a", "b", "c"] and m2.marks == {("결과 테이블", "a"): ("삭제", 1)}


def test_순서만_바뀐_것과_필수_빈칸은_표시하지_않는다():
    old = ([("a", False), ("b", True)], ["x", "y"], ["p", "q"], False)
    new = ([("b", True), ("a", False)], ["y", "x"], ["q", "p"], False)
    merged = screen_diff.merge(old, new)
    assert merged.marks == {} and merged.changes == []
    assert merged.conds == [("b", True), ("a", False)]  # 새 판의 순서를 따른다


def test_번호는_필터_바_Go_액션_결과_테이블_순이다():
    old = ([], ["a"], ["t"], False)
    new = ([("f", False)], ["a", "b"], ["t", "u"], True)
    merged = screen_diff.merge(old, new)
    assert [(c.number, c.target) for c in merged.changes] == [
        (1, "필터 바 · f"), (2, "액션 · Go"), (3, "액션 · b"), (4, "결과 테이블 · u"),
    ]


def test_번호_표기는_20까지_원문자_그_뒤는_괄호():
    assert screen_diff.circled(1) == "①" and screen_diff.circled(20) == "⑳" and screen_diff.circled(21) == "(21)"


# ---------------------------------------------------------------- 추가·수정·삭제 표시

def test_추가_수정_삭제가_번호와_스타일로_표시되고_변경_표가_채워진다(proj, capsys):
    _map(proj, extra=CHANGE % 10)
    _origin(proj, OLD)
    cond = _tbl("", ["판매 조직", "기간", "고객"], ["N", "Y", ""])
    btn = _tbl("", ["취소", "조회"])
    res = _tbl("열", ["판매문서", "품목", "수량", "금액"])
    code, report = _run(proj, capsys, screen(cond=cond, btn=btn, res=res))
    assert code == 0, report
    ws = _sheet(proj, report)

    mod = _at(ws, "① 판매 조직")  # 수정: 필수 Y → N 이라 별표가 없다. 라벨만 노랗다
    assert _fill(mod) == YELLOW and _fill(ws.cell(mod.row, mod.column + 2)) == "FFFFFFFF"
    assert mod.font.b and mod.border.left.style == "thin"  # 기존 상자 스타일 위에 덮어쓴 겹침이다

    gone = _at(ws, "② 판매처")  # 삭제: 라벨·입력 둘 다 회색 취소선
    for c in (gone, ws.cell(gone.row, gone.column + 2)):
        assert _fill(c) == GREY
    assert gone.font.strike and gone.font.color.rgb == "FF808080"
    assert ws.cell(gone.row, gone.column + 2).font.strike

    new = _at(ws, "③ 고객")  # 추가: 라벨·입력 둘 다 노랗다
    assert _fill(new) == YELLOW and _fill(ws.cell(new.row, new.column + 2)) == YELLOW and not new.font.strike

    assert _fill(_at(ws, "④ 일괄 취소")) == GREY and _at(ws, "④ 일괄 취소").font.strike
    assert _fill(_at(ws, "⑤ 조회")) == YELLOW

    head = _at(ws, "⑥ 상태")  # 결과 테이블 삭제: 머리글과 빈 행 모두 회색 취소선
    assert _fill(head) == GREY and head.font.strike
    assert _fill(ws.cell(head.row + 1, head.column)) == GREY and _fill(ws.cell(head.row + 2, head.column)) == GREY
    add = _at(ws, "⑦ 금액")  # 결과 테이블 추가: 머리글만 노랗다
    assert _fill(add) == YELLOW and ws.cell(add.row + 1, add.column).fill.fill_type is None

    assert _table(ws) == [
        ("①", "수정", "필터 바 · 판매 조직", "필수 Y → N"),
        ("②", "삭제", "필터 바 · 판매처", "필터 바에서 '판매처' 삭제"),
        ("③", "추가", "필터 바 · 고객", "필터 바에 '고객' 추가"),
        ("④", "삭제", "액션 · 일괄 취소", "액션에서 '일괄 취소' 삭제"),
        ("⑤", "추가", "액션 · 조회", "액션에 '조회' 추가"),
        ("⑥", "삭제", "결과 테이블 · 상태", "결과 테이블에서 '상태' 삭제"),
        ("⑦", "추가", "결과 테이블 · 금액", "결과 테이블에 '금액' 추가"),
    ]
    # 바뀌지 않은 항목은 번호도 표시도 없다
    assert _fill(_at(ws, "기간 *")) == "FFD9E1F2" and _fill(_at(ws, "취소")) != YELLOW and not _at(ws, "취소").font.strike


def test_이름을_바꾸면_삭제와_추가다(proj, capsys):
    _map(proj, extra=CHANGE % 5)
    _origin(proj, OLD)
    cond = _tbl("", ["판매 조직", "기간(연월)", "판매처"], ["Y", "Y", ""])
    code, report = _run(proj, capsys, screen(cond=cond))
    assert code == 0, report
    ws = _sheet(proj, report)
    assert _at(ws, "① 기간 *").font.strike and _fill(_at(ws, "② 기간(연월) *")) == YELLOW
    assert [(r[0], r[1]) for r in _table(ws)] == [("①", "삭제"), ("②", "추가")]


def test_조회가_자동에서_Go로_바뀌면_Go_액션_추가다(proj, capsys):
    _map(proj, extra=CHANGE % 5)
    _origin(proj, OLD)
    code, report = _run(proj, capsys, screen(settings="| 항목 | 값 |\n|---|---|\n| 조회 | Go |\n"))
    assert code == 0, report
    ws = _sheet(proj, report)
    go = _at(ws, "① Go")
    assert _fill(go) == YELLOW and go.column == 18  # 영역 오른쪽 끝(R열)
    assert _table(ws) == [("①", "추가", "액션 · Go", "액션에 'Go' 추가")]


def test_조회가_Go에서_자동으로_바뀌면_Go_액션_삭제다(proj, capsys):
    _map(proj, extra=CHANGE % 5)
    _origin(proj, screen(settings="| 항목 | 값 |\n|---|---|\n| 조회 | Go |\n"))
    code, report = _run(proj, capsys, OLD)
    assert code == 0, report
    ws = _sheet(proj, report)
    go = _at(ws, "① Go")  # 삭제된 Go 도 그림에 남는다
    assert _fill(go) == GREY and go.font.strike
    assert _table(ws) == [("①", "삭제", "액션 · Go", "액션에서 'Go' 삭제")]


def test_순서만_바뀌면_표시도_변경_표도_없다(proj, capsys):
    _origin(proj, OLD)  # 지도에 `변경:` 이 없어도 변경이 없으면 거부하지 않는다
    cond = _tbl("", ["판매처", "기간", "판매 조직"], ["", "Y", "Y"])
    btn = _tbl("", ["일괄 취소", "취소"])
    code, report = _run(proj, capsys, screen(cond=cond, btn=btn))
    assert code == 0, report
    ws = _sheet(proj, report)
    assert not [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value[:1] in "①②③"]
    assert _fill(_at(ws, "판매처")) == "FFD9E1F2" and ws["B20"].value is None


def test_필수_빈칸은_N으로_본다(proj, capsys):
    _origin(proj, OLD)
    cond = _tbl("", ["판매 조직", "기간", "판매처"], ["Y", "Y", "N"])  # 옛 판의 판매처는 빈칸
    code, report = _run(proj, capsys, screen(cond=cond))
    assert code == 0, report


def test_캡처로_그리는_화면은_비교하지_않는다(proj, capsys, monkeypatch):
    """그림 절이 없고 캡처 절만 있는 화면은 이 비교의 대상이 아니다 — 지난 판을 읽으려 git 을 부르지도 않는다."""
    from checker import export

    def boom(root, rel):
        raise AssertionError("git 을 부르면 안 된다")

    monkeypatch.setattr(export, "_origin_main_text", boom)
    _map(proj, extra="        캡처: {설명: U7:V20}\n")
    _, report = _run(proj, capsys, "#### 화면 1 캡처\n\n![캡처](docs/spec/img/x_v1.png)\n\n| 번호 | 설명 |\n|---|---|\n| 1 | a |\n")
    assert "git 을 부르면" not in str(report)


# ---------------------------------------------------------------- 비교 기준

def test_main에_작업본이_없으면_첫_판이라_표시_없이_그린다(proj, capsys):
    _origin(proj, None)
    code, report = _run(proj, capsys, OLD)
    assert code == 0, report
    ws = _sheet(proj, report)
    assert _at(ws, "판매 조직 *") and ws["B20"].value is None
    assert not any("지난 판" in w["reason"] for w in report["warnings"])


def test_origin_main을_못_읽으면_검사_불능이다(proj, capsys):
    _origin(proj, OLD, ref=False)  # git 저장소지만 origin/main 이 없다
    code, report = _run(proj, capsys, OLD)
    assert code == 2 and report["status"] == "error", report
    assert "지난 판과 비교하지 못했다" in report["problems"][0]["reason"] and "git fetch" in report["problems"][0]["reason"]
    assert _nothing_written(proj)


def test_git_저장소가_아니어도_검사_불능이다(proj, capsys):
    code, report = _run(proj, capsys, OLD)  # proj 는 git 저장소가 아니다
    assert code == 2 and "지난 판과 비교하지 못했다" in report["problems"][0]["reason"]
    assert _nothing_written(proj)


def test_그릴_화면이_없으면_git을_부르지_않는다(proj, capsys, monkeypatch):
    from checker import export

    def boom(root, rel):
        raise AssertionError("git 을 부르면 안 된다")

    monkeypatch.setattr(export, "_origin_main_text", boom)
    code, report = _run(proj, capsys, "")  # 그림 절이 없다 → 그리지 않았다는 경고만
    assert code in (0, 3) and any("그리지 않았다" in w["reason"] for w in report["warnings"]), report


def test_main_판에_그림_절이_없으면_표시_없이_그리고_경고한다(proj, capsys):
    _origin(proj, "")  # main 의 작업본에는 화면 그림 절이 없다
    code, report = _run(proj, capsys, OLD)
    assert code in (0, 3), report
    ws = _sheet(proj, report)
    assert _at(ws, "판매 조직 *") and ws["B20"].value is None
    assert any("지난 판에 그림 절이 없어 변경 표시를 하지 않았다" in w["reason"] or "지난 판(origin/main)에 그림 절이 없어 변경 표시를 하지 않았다" in w["reason"] for w in report["warnings"])


def test_main_판_그림_절이_깨졌으면_검사_불능이다(proj, capsys):
    _origin(proj, "#### 화면 1 그림\n\n##### 필터 바\n\n| 이름 | 값 |\n|---|---|\n| a | b |\n")
    code, report = _run(proj, capsys, OLD)
    assert code == 2, report
    assert "지난 판" in report["problems"][0]["reason"] and "검사 불능" in report["problems"][0]["reason"]
    assert _nothing_written(proj)


def test_변경이_없으면_main과_같은_판이라도_아무것도_쓰지_않는다(proj, capsys):
    _map(proj, extra=CHANGE % 5)
    _origin(proj, OLD)
    code, report = _run(proj, capsys, OLD)
    assert code == 0, report
    ws = _sheet(proj, report)
    assert ws["B20"].value is None and ws["H20"].value is None


def test_check_모드도_같은_비교를_한다(proj, capsys):
    """`--check` 는 같은 build_form 을 쓰므로 지난 판을 못 읽으면 검사 불능, 변경 표 자리가 없으면 거부다."""
    _origin(proj, OLD, ref=False)
    code, report = _run(proj, capsys, OLD, "--check")
    assert code == 2
    proj2 = proj  # origin/main 을 만든 뒤에는 변경 표 자리가 없다는 거부가 같이 나온다
    _git(proj2, "update-ref", "refs/remotes/origin/main", "HEAD")
    code, report = _run(proj2, capsys, screen(btn=_tbl("", ["취소"])), "--check")
    assert code == 1 and "변경 내용 표 자리가 지도에 없다" in report["problems"][0]["reason"]


# ---------------------------------------------------------------- 변경 내용 표 거부

def test_변경_표_자리가_지도에_없으면_거부한다(proj, capsys):
    _origin(proj, OLD)
    code, report = _run(proj, capsys, screen(btn=_tbl("", ["취소"])))
    assert code == 1, report
    assert "변경 내용 표 자리가 지도에 없다 — 지도 `화면:` 에 `변경:` 을 적는다" in report["problems"][0]["reason"]
    assert _nothing_written(proj)


def test_변경이_최대_행보다_많으면_자르지_않고_거부한다(proj, capsys):
    _map(proj, extra=CHANGE % 1)
    _origin(proj, OLD)
    code, report = _run(proj, capsys, screen(btn=_tbl("", ["새 버튼"])))  # 추가 1 + 삭제 2
    assert code == 1, report
    assert "최대_행 1" in report["problems"][0]["reason"] and "자르지 않는다" in report["problems"][0]["reason"]
    assert _nothing_written(proj)


def test_변경_표를_쓸_칸에_값이_남아_있으면_거부한다(proj, capsys):
    _template(proj / "templates" / "시험.xlsx", sample={"E21": "견본"})
    _map(proj, extra=CHANGE % 10)
    _origin(proj, OLD)
    code, report = _run(proj, capsys, screen(btn=_tbl("", ["새 버튼"])))
    assert code == 1, report
    assert "이미 값이 있다" in report["problems"][0]["reason"] and "E21" in report["problems"][0]["reason"]
    assert _nothing_written(proj)


# ---------------------------------------------------------------- 지도 검증

def _problems(**kw) -> list[str]:
    item = {"이름": "화면 1", "영역": "C7:S14", **kw}
    return [why for _, why in screen_draw.map_problems({"시트": [{"이름": "S", "화면": [item]}]})]


def _spot(**over):
    base = {"시작_행": 20, "최대_행": 10, "열": {"번호": "B", "구분": "C", "대상": "E", "내용": "H"}}
    return {**base, **over}


def test_지도_변경_표가_올바르면_오류가_없다():
    assert _problems(변경=_spot()) == []
    item = screen_draw.items({"화면": [{"이름": "화면 1", "영역": "C7:S14", "변경": _spot()}]})[0]
    assert item["변경"] == _spot()


@pytest.mark.parametrize("spot,needle", [
    ("x", "목록이어야"),
    (_spot(모름=1), "모르는 키"),
    (_spot(시작_행=0), "시작_행 는 1 이상의 정수"),
    (_spot(최대_행="3"), "최대_행 는 1 이상의 정수"),
    (_spot(최대_행=True), "최대_행 는 1 이상의 정수"),
    (_spot(열={"번호": "B", "구분": "C", "대상": "E"}), "`내용` 항목이 없다"),
    (_spot(열={"번호": "B", "구분": "C", "대상": "E", "내용": "H", "비고": "I"}), "모른다"),
    (_spot(열={"번호": "B", "구분": "C", "대상": "E", "내용": "8"}), "열 문자"),
    (_spot(열="B"), "열 문자 목록이어야"),
    ({"최대_행": 5}, "시작_행"),
    (_spot(시작_행=1048570, 최대_행=20), "마지막 행"),
])
def test_지도_변경_표_오류(spot, needle):
    assert any(needle in why for why in _problems(변경=spot)), _problems(변경=spot)


def test_지도_변경_표가_다른_범위와_겹치면_오류다():
    assert any("겹친다" in w for w in _problems(변경=_spot(시작_행=7)))  # 화면 영역 C7:S14 의 C·E·H 열과 겹친다
    assert any("겹친다" in w for w in _problems(변경=_spot(열={"번호": "B", "구분": "C", "대상": "C", "내용": "H"})))  # 같은 열 두 번
    assert any("겹친다" in w for w in _problems(캡처={"설명": "B22:C30"}, 변경=_spot()))  # 같은 화면의 캡처 설명과
    # 지도가 맡은 칸·표·서술·유지와
    owned = {"시트": [{"이름": "S", "칸": [{"항목": "코드", "위치": "B22"}], "화면": [{"이름": "화면 1", "영역": "C7:S14", "변경": _spot()}]}]}
    assert any("겹친다" in why for _, why in screen_draw.map_problems(owned))
    # 다른 화면의 변경 표와
    two = {"시트": [{"이름": "S", "화면": [
        {"이름": "화면 1", "영역": "C7:S14", "변경": _spot()},
        {"이름": "화면 2", "영역": "C30:S40", "변경": _spot(시작_행=25)},
    ]}]}
    assert any("겹친다" in why for _, why in screen_draw.map_problems(two))


def test_지도_스타일이_추가_수정_삭제_이름을_받는다():
    assert _problems(스타일={"추가": {"채우기": "00FF00"}, "삭제": {"글꼴_색": "FF0000", "취소선": True}}) == []
    bad = _problems(스타일={"추가": {"채우기": "초록"}, "변경": {"굵게": True}})
    assert any("스타일 추가" in w for w in bad) and any("스타일 이름 '변경'" in w for w in bad)
    assert any("취소선" in w for w in _problems(스타일={"삭제": {"취소선": "예"}}))


def test_지도_스타일이_표시_스타일을_통째로_바꾼다(proj, capsys):
    _map(proj, extra=CHANGE % 5 + "        스타일:\n          추가: {채우기: 00FF00}\n          삭제: {글꼴_색: FF0000}\n")
    _origin(proj, OLD)
    code, report = _run(proj, capsys, screen(btn=_tbl("", ["취소", "새 버튼"])))
    assert code == 0, report
    ws = _sheet(proj, report)
    assert _fill(_at(ws, "② 새 버튼")) == "FF00FF00"  # 추가 스타일이 지도 값으로 바뀌었다
    gone = _at(ws, "① 일괄 취소")
    assert gone.font.color.rgb == "FFFF0000" and not gone.font.strike and _fill(gone) == "FFD9D9D9"  # 삭제 스타일은 통째로 바뀌어 취소선·기본 채우기 대신 액션 상자 스타일


# ---------------------------------------------------------------- 중복 이름

def test_소절_안_이름이_겹치면_양식_밖으로_거부한다():
    item = screen_draw.items({"화면": [{"이름": "화면 1", "영역": "C7:S14"}]})[0]
    body = ("#### 화면 1 그림\n\n##### 필터 바\n\n" + _tbl("", ["a", "a", "b"], ["Y", "", ""])
            + "\n##### 액션\n\n" + _tbl("", ["취소", "취소"])).splitlines()
    off = screen_draw.offenders(item, body)
    assert [(k, n) for k, n, _ in off if k == "화면 항목 중복"] == [
        ("화면 항목 중복", "화면 1 그림 · 필터 바: 'a' (2번)"), ("화면 항목 중복", "화면 1 그림 · 액션: '취소' (2번)"),
    ]


def test_조회_Go와_이름이_Go인_액션은_겹친다():
    item = screen_draw.items({"화면": [{"이름": "화면 1", "영역": "C7:S14"}]})[0]
    body = ("#### 화면 1 그림\n\n| 항목 | 값 |\n|---|---|\n| 조회 | Go |\n\n##### 액션\n\n" + _tbl("", ["Go"])).splitlines()
    assert any(k == "화면 항목 중복" for k, _, _ in screen_draw.offenders(item, body))


def test_중복_이름은_산출도_거부한다(proj, capsys):
    _origin(proj, OLD)
    code, report = _run(proj, capsys, screen(btn=_tbl("", ["취소", "취소"])))
    assert code == 2 and "화면 항목 중복" in report["problems"][0]["reason"]
    assert _nothing_written(proj)


def test_넘침은_삭제_항목을_더한_목록으로_본다(proj, capsys):
    """삭제 항목도 그림에 남으므로, 새 판만으로는 들어가도 합치면 영역을 넘을 수 있고 그러면 지금처럼 거절한다."""
    _map(proj, text=MAP.replace("영역: C7:S14", "영역: C7:S7"), extra=CHANGE % 10)
    new = screen(cond=_tbl("", ["d", "e", "f"], ["", "", ""]), btn=None, res=None)
    item = screen_draw.items({"화면": [{"이름": "화면 1", "영역": "C7:S7"}]})[0]
    assert screen_draw.layout(item, [(n, False) for n in "def"], [], [])[1] == []  # 새 판만으로는 한 줄에 들어간다
    _origin(proj, screen(cond=_tbl("", ["a", "b", "c"], ["", "", ""]), btn=None, res=None))
    code, report = _run(proj, capsys, new)
    assert code == 1 and "영역 C7:S7 를 넘는다" in report["problems"][0]["reason"], report
    assert _nothing_written(proj)


# ---------------------------------------------------------------- 지난 판은 느슨하게 읽는다

def test_지난_판의_중복_미정_표_밖_글은_무시하고_비교한다(proj, capsys):
    """main 의 그림 절이 새 판에서는 거부될 모양이어도(중복·미정·표 밖 글), 그것을 고치는 PR 이 검사 불능이 되면 안 된다."""
    _map(proj, extra=CHANGE % 5)
    old = (
        "#### 화면 1 그림\n\n비고: 손으로 적은 글\n\n"
        "| 항목 | 값 |\n|---|---|\n| 모름 | 값 |\n| 조회 | 아무거나 |\n\n"
        "##### 필터 바\n\n| 이름 | 필수 |\n|---|---|\n| 판매 조직 | Y |\n| 판매 조직 | N |\n| 기간 | Y |\n| 미정 | |\n| | |\n| 판매처 | |\n\n"
        "##### 액션\n\n| 이름 |\n|---|\n| 취소 |\n| 취소 |\n| 일괄 취소 |\n\n"
        "##### 결과 테이블\n\n" + RES
    )
    _origin(proj, old)
    # 새 판: 중복·미정을 고치고 액션 하나를 더했다 — 지난 판에서 처음 것만 쓰므로 바뀐 것은 추가 하나뿐이다.
    code, report = _run(proj, capsys, screen(btn=_tbl("", ["취소", "일괄 취소", "조회"])))
    assert code == 0, report
    ws = _sheet(proj, report)
    assert _fill(_at(ws, "① 조회")) == YELLOW
    assert _fill(_at(ws, "판매 조직 *")) == "FFD9E1F2"  # 처음 것(필수 Y)과 같다 — 수정이 아니다
    assert _table(ws) == [("①", "추가", "액션 · 조회", "액션에 '조회' 추가")]


def test_지난_판_소절_머리글이_틀리면_main을_먼저_고치라고_알린다(proj, capsys):
    _origin(proj, "#### 화면 1 그림\n\n##### 액션\n\n| 이름 | 값 |\n|---|---|\n| a | b |\n")
    code, report = _run(proj, capsys, OLD)
    assert code == 2, report
    assert "main 의 해당 그림 절을 먼저 고친다" in report["problems"][0]["reason"]


def test_Spec이_하위_폴더에_있고_이름에_한글_공백이_있어도_지난_판과_비교한다(proj, capsys):
    import json

    from checker import export

    _map(proj, extra=CHANGE % 5)
    rel = "docs/spec/DEV-001 화면 시험.md"
    spec = proj / rel
    spec.parent.mkdir(parents=True)
    _git(proj, "init", "-q")
    _git(proj, "config", "user.email", "t@example.com")
    _git(proj, "config", "user.name", "t")
    _git(proj, "config", "commit.gpgsign", "false")
    spec.write_bytes((SPEC_HEAD + OLD).encode("utf-8"))
    _git(proj, "add", "--", rel)
    _git(proj, "commit", "-q", "-m", "x")
    _git(proj, "update-ref", "refs/remotes/origin/main", "HEAD")
    new = screen(btn=_tbl("", ["취소", "새 버튼"]))  # 일괄 취소 삭제, 새 버튼 추가
    spec.write_bytes((SPEC_HEAD + new).encode("utf-8"))
    code = export.main(["--spec", str(spec), "--root", str(proj)])
    report = json.loads(capsys.readouterr().out)
    assert code == 0, report
    ws = _sheet(proj, report)
    assert _at(ws, "① 일괄 취소").font.strike and _fill(_at(ws, "② 새 버튼")) == YELLOW
    assert [r[:2] for r in _table(ws)] == [("①", "삭제"), ("②", "추가")]
