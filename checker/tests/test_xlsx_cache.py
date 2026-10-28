"""xlsx 리더의 통합문서 재사용(#515): 같은 파일은 프로세스에서 한 번만 연다."""
from __future__ import annotations

import os
from pathlib import Path

import openpyxl
import pytest

from checker.readers import xlsx
from checker.rules import table_headers


def _make(path: Path, header: list[str], sheet: str = "Title") -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(["제목"])
    ws.append(header)
    wb.save(path)
    wb.close()
    return path


@pytest.fixture
def loads(monkeypatch):
    """openpyxl.load_workbook 호출 횟수를 센다."""
    calls = []
    real = openpyxl.load_workbook

    def counting(*args, **kwargs):
        calls.append(args[0])
        return real(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", counting)
    return calls


def test_같은_파일은_여러_행을_읽어도_한_번만_연다(tmp_path, loads):
    f = _make(tmp_path / "a.xlsx", ["A", "B"])
    assert xlsx.read_row(f, "Title", 2) == ["A", "B"]
    assert xlsx.read_row(f, "Title", 1) == ["제목"]
    assert xlsx.read_row(f, "없는 시트", 1) is None
    assert xlsx.read_xlsx(f)["sheets"] == ["Title"]
    assert len(loads) == 1


def test_table_headers_규칙_셋이_같은_템플릿을_한_번만_연다(tmp_path, loads):
    tpl = _make(tmp_path / "tpl.xlsx", ["A", "B"])
    for _ in range(3):
        params = table_headers.from_template({"시트": "Title", "헤더_행": 2}, tpl)
        assert params["_expected"] == ["A", "B"]
    assert len(loads) == 1


def test_파일이_바뀌면_다시_연다(tmp_path, loads):
    f = _make(tmp_path / "a.xlsx", ["A", "B"])
    assert xlsx.read_row(f, "Title", 2) == ["A", "B"]
    _make(f, ["A", "B", "C"])
    st = f.stat()
    os.utime(f, ns=(st.st_atime_ns, st.st_mtime_ns + 5_000_000_000))  # 같은 초에 저장돼도 mtime 이 달라지게
    assert xlsx.read_row(f, "Title", 2) == ["A", "B", "C"]
    assert len(loads) == 2


def test_열기에_실패하면_캐시하지_않고_다음에_다시_시도한다(tmp_path, loads):
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"not a zip")
    for _ in range(2):
        with pytest.raises(Exception):
            xlsx.read_row(bad, "Title", 1)
    assert len(loads) == 2
    assert not xlsx._OPEN
    # 같은 경로를 올바른 파일로 바꾸면 읽힌다.
    _make(bad, ["X"])
    assert xlsx.read_row(bad, "Title", 2) == ["X"]


def test_없는_파일은_예외이고_캐시되지_않는다(tmp_path):
    with pytest.raises(FileNotFoundError):
        xlsx.read_row(tmp_path / "none.xlsx", "Title", 1)
    assert not xlsx._OPEN


def test_열어_둔_통합문서_수에_한도가_있고_밀려난_것은_닫힌다(tmp_path, loads):
    files = [_make(tmp_path / f"f{i}.xlsx", [str(i)]) for i in range(xlsx._MAX_OPEN + 2)]
    for f in files:
        xlsx.read_row(f, "Title", 2)
    assert len(xlsx._OPEN) == xlsx._MAX_OPEN
    xlsx.clear_cache()
    assert not xlsx._OPEN
