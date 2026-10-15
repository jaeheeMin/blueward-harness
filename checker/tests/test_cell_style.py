"""셀 서식·병합 쓰기 기반(`checker.cell_style`)을 검증한다(이슈 #335, PR C1).

합성 xlsx 를 openpyxl 로 만들고, 시트·styles XML 을 zip 수준에서 고친 뒤 openpyxl 로 다시 열어 서식을 확인한다.
"""
from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Font, PatternFill

from checker import cell_style
from checker.cell_style import StylesXml, merge_range, merges_of, style_range, validate_style
from checker.export import ExportError, SheetXml

SHEET = "xl/worksheets/sheet1.xml"
STYLES = "xl/styles.xml"


def _make(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "화면"
    ws["A1"] = 1234.5
    ws["A1"].number_format = "0.00"
    ws["A1"].font = Font(name="Arial", size=11, bold=True, italic=True)
    ws["C3"] = "채움"
    ws["C3"].fill = PatternFill("solid", fgColor="FFFF00")
    ws["B10"] = "기존 병합"
    ws.merge_cells("B10:C11")
    ws.sheet_properties.tabColor = "FF0000"
    wb.save(path)


@pytest.fixture
def xlsx(tmp_path: Path) -> Path:
    p = tmp_path / "base.xlsx"
    _make(p)
    return p


def _load(path: Path) -> tuple[SheetXml, StylesXml]:
    with zipfile.ZipFile(path) as z:
        return SheetXml(z.read(SHEET).decode(), set(), False), StylesXml(z.read(STYLES).decode())


def _save(src: Path, dst: Path, ed: SheetXml, st: StylesXml) -> Path:
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for info in zi.infolist():
            data = zi.read(info.filename)
            if info.filename == SHEET:
                data = ed.render().encode()
            elif info.filename == STYLES:
                data = st.render().encode()
            zo.writestr(info.filename, data)
    return dst


def test_derive_adds_and_reuses(xlsx: Path):
    ed, st = _load(xlsx)
    before = st.xf_count()
    a = st.derive(0, {"채우기": "DDEBF7", "테두리": "얇게"})
    b = st.derive(0, {"테두리": "얇게", "채우기": "ddebf7"})
    assert a == b == before
    assert st.xf_count() == before + 1
    c = st.derive(0, {"채우기": "DDEBF7", "테두리": "굵게"})
    assert c == before + 1
    # 같은 채우기는 fills 에서 재사용된다(채우기 항목 1개만 늘었다).
    out = st.render()
    assert re.search(rf'<cellXfs count="{before + 2}"', out)
    assert out.count('rgb="FFDDEBF7"') == 1
    assert re.search(rf'<fills count="{st.fills.original + 1}"', out)


def test_derive_overrides_only_given_parts(xlsx: Path):
    ed, st = _load(xlsx)
    base = int(ed.cell("A1").style)
    new = st.derive(base, {"글꼴_크기": 9, "글꼴_색": "FF0000"})
    out = _save(xlsx, xlsx.with_name("o.xlsx"), ed, st)  # 쓰기 가능한지
    ed.set_style("A1", str(new))
    _save(xlsx, out, ed, st)
    c = openpyxl.load_workbook(out)["화면"]["A1"]
    assert c.number_format == "0.00"  # 숫자 형식 유지
    assert c.font.b and c.font.i and c.font.name == "Arial"  # 글꼴의 나머지 유지
    assert c.font.sz == 9 and c.font.color.rgb == "FFFF0000"
    assert c.fill.fill_type is None


def test_font_bold_off_and_alignment(xlsx: Path):
    ed, st = _load(xlsx)
    base = int(ed.cell("A1").style)
    new = st.derive(base, {"굵게": False, "가로": "가운데", "세로": "위", "줄바꿈": True})
    ed.set_style("A1", str(new))
    c = openpyxl.load_workbook(_save(xlsx, xlsx.with_name("o.xlsx"), ed, st))["화면"]["A1"]
    assert not c.font.b and c.font.i
    assert (c.alignment.horizontal, c.alignment.vertical, c.alignment.wrap_text) == ("center", "top", True)
    # 정렬만 바꾸면 기존 정렬 속성(가로)이 남는다.
    ed2, st2 = _load(xlsx)
    n1 = st2.derive(0, {"가로": "오른쪽"})
    n2 = st2.derive(n1, {"줄바꿈": False, "세로": "아래"})
    ed2.set_style("A1", str(n2))
    a = openpyxl.load_workbook(_save(xlsx, xlsx.with_name("p.xlsx"), ed2, st2))["화면"]["A1"].alignment
    assert (a.horizontal, a.vertical, a.wrap_text) == ("right", "bottom", None)


def test_border_color_and_none_fill(xlsx: Path):
    ed, st = _load(xlsx)
    c3 = int(ed.cell("C3").style)
    new = st.derive(c3, {"채우기": "없음", "테두리": "중간", "테두리_색": "336699"})
    ed.set_style("C3", str(new))
    c = openpyxl.load_workbook(_save(xlsx, xlsx.with_name("o.xlsx"), ed, st))["화면"]["C3"]
    assert c.fill.fill_type is None
    assert c.border.left.style == c.border.bottom.style == "medium"
    assert c.border.top.color.rgb == "FF336699"
    ed2, st2 = _load(xlsx)
    n = st2.derive(c3, {"테두리": "없음"})
    ed2.set_style("C3", str(n))
    c = openpyxl.load_workbook(_save(xlsx, xlsx.with_name("p.xlsx"), ed2, st2))["화면"]["C3"]
    assert c.border.left.style is None and c.fill.fill_type == "solid"  # 채우기는 그대로


def test_untouched_parts_stay_identical(xlsx: Path):
    _, st = _load(xlsx)
    assert st.render() == st.xml  # 아무것도 안 했으면 바이트 그대로
    st.derive(0, {"굵게": True})
    out = st.render()
    for tag in ("cellStyleXfs", "cellStyles", "dxfs"):
        m = re.search(rf"<{tag}\b.*?(?:</{tag}>|/>)", st.xml, re.S)
        if m:
            assert m.group(0) in out


def test_style_cells_that_do_not_exist(xlsx: Path):
    ed, st = _load(xlsx)
    n = style_range(ed, st, "E20", "G21", {"채우기": "DDEBF7", "테두리": "얇게"}, "화면")
    assert n == 6
    out = _save(xlsx, xlsx.with_name("o.xlsx"), ed, st)
    ws = openpyxl.load_workbook(out)["화면"]
    for ref in ("E20", "F20", "G20", "E21", "F21", "G21"):
        assert ws[ref].fill.fgColor.rgb == "FFDDEBF7" and ws[ref].border.right.style == "thin"
    assert ws["A1"].value == 1234.5 and ws["C3"].value == "채움"
    assert '<dimension ref="A1:G21"/>' in zipfile.ZipFile(out).read(SHEET).decode()
    # 행 안 셀 순서가 열 순서다.
    xml = zipfile.ZipFile(out).read(SHEET).decode()
    row = re.search(r'<row r="20".*?</row>', xml, re.S).group(0)
    assert re.findall(r'<c r="([A-Z]+)20"', row) == ["E", "F", "G"]


def test_style_range_inserts_between_existing_cells(xlsx: Path):
    ed, st = _load(xlsx)
    style_range(ed, st, "B3", "B3", {"굵게": True}, "화면")  # C3 앞
    style_range(ed, st, "D3", None, {"굵게": True}, "화면")  # C3 뒤
    out = _save(xlsx, xlsx.with_name("o.xlsx"), ed, st)
    xml = zipfile.ZipFile(out).read(SHEET).decode()
    row = re.search(r'<row r="3".*?</row>', xml, re.S).group(0)
    assert re.findall(r'<c r="([A-Z]+)3"', row) == ["B", "C", "D"]
    assert openpyxl.load_workbook(out)["화면"]["C3"].fill.fgColor.rgb == "00FFFF00"  # 기존 칸 서식 유지


def test_new_cell_follows_column_default_style(xlsx: Path):
    _, st2 = _load(xlsx)
    fill_xf = st2.derive(0, {"채우기": "FFC000"})
    with zipfile.ZipFile(xlsx) as z:
        xml = z.read(SHEET).decode().replace("<sheetData", f'<cols><col min="6" max="6" width="9" style="{fill_xf}" customWidth="1"/></cols><sheetData', 1)
    ed = SheetXml(xml, set(), False)
    style_range(ed, st2, "F5", None, {"굵게": True}, "화면")
    s = ed.cell("F5").style
    attrs, _ = cell_style._xf_parts(st2.xfs.items[int(s)])
    fill_id = int(dict(attrs)["fillId"])
    assert "FFC000" in st2.fills.items[fill_id]  # 열 기본 서식 위에 굵게를 덮었다


def test_merge_add_and_count(xlsx: Path):
    ed, _ = _load(xlsx)
    assert merges_of(ed) == ["B10:C11"]
    assert merge_range(ed, "E2", "G3", "화면") == "E2:G3"
    assert merges_of(ed) == ["B10:C11", "E2:G3"]
    out = _save(xlsx, xlsx.with_name("o.xlsx"), ed, _load(xlsx)[1])
    xml = zipfile.ZipFile(out).read(SHEET).decode()
    assert '<mergeCells count="2">' in xml
    assert {str(r) for r in openpyxl.load_workbook(out)["화면"].merged_cells.ranges} == {"B10:C11", "E2:G3"}


def test_merge_creates_mergecells_in_schema_order(tmp_path: Path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    p = tmp_path / "n.xlsx"
    wb.save(p)
    ed, st = _load(p)
    assert merges_of(ed) == []
    merge_range(ed, "B2", "C2")
    xml = ed.render()
    assert re.search(r'</sheetData><mergeCells count="1"><mergeCell ref="B2:C2"/></mergeCells><pageMargins', xml)
    out = _save(p, tmp_path / "o.xlsx", ed, st)
    assert [str(r) for r in openpyxl.load_workbook(out).active.merged_cells.ranges] == ["B2:C2"]


@pytest.mark.parametrize("first,last", [("B10", "B10"), ("C11", "D12"), ("A9", "D12"), ("C10", "C10")])
def test_merge_overlap_or_single_refused(xlsx: Path, first: str, last: str):
    ed, _ = _load(xlsx)
    with pytest.raises(ExportError) as ei:
        merge_range(ed, first, last, "화면")
    assert ei.value.sheet == "화면"
    assert merges_of(ed) == ["B10:C11"]  # 쓰지 않았다


def test_merge_refuses_overlap_with_new_merge(xlsx: Path):
    ed, _ = _load(xlsx)
    merge_range(ed, "E2", "F3")
    with pytest.raises(ExportError, match="겹친다"):
        merge_range(ed, "F3", "G4")


def test_merge_refuses_hidden_values(xlsx: Path):
    ed, _ = _load(xlsx)
    with pytest.raises(ExportError, match="값이 있다"):
        merge_range(ed, "A1", "C3")  # C3 에 값
    merge_range(ed, "C3", "D3")  # 왼쪽 위 칸의 값은 괜찮다


@pytest.mark.parametrize("bad,expect", [
    ({"색": "red"}, "모른다"),
    ({"채우기": "빨강"}, "RRGGBB"),
    ({"채우기": "FFF"}, "RRGGBB"),
    ({"테두리": "점선"}, "얇게"),
    ({"테두리_색": "000000"}, "때만"),
    ({"테두리": "없음", "테두리_색": "000000"}, "때만"),
    ({"굵게": "예"}, "true 또는 false"),
    ({"줄바꿈": 1}, "true 또는 false"),
    ({"글꼴_크기": 0}, "숫자"),
    ({"글꼴_크기": "9"}, "숫자"),
    ({"글꼴_크기": True}, "숫자"),
    ({"글꼴_색": 12}, "RRGGBB"),
    ({"가로": "중앙"}, "왼쪽"),
    ({"세로": "중간"}, "위"),
    ({}, "비어"),
    ("굵게", "목록"),
    (None, "목록"),
])
def test_validate_rejects(bad, expect: str):
    problems = validate_style(bad)
    assert problems and any(expect in p for p in problems), problems


def test_validate_accepts_full_definition():
    full = {"채우기": "DDEBF7", "테두리": "얇게", "테두리_색": "000000", "굵게": True, "글꼴_크기": 9.5,
            "글꼴_색": "000000", "가로": "왼쪽", "세로": "가운데", "줄바꿈": True}
    assert validate_style(full) == []
    assert validate_style({"채우기": "없음", "테두리": "없음"}) == []


def test_invalid_definition_stops_everything(xlsx: Path):
    ed, st = _load(xlsx)
    before = st.xf_count()
    with pytest.raises(ExportError, match="모른다"):
        st.derive(0, {"굵게": True, "그림자": True})
    with pytest.raises(ExportError, match="스타일 정의가 틀렸다") as ei:
        style_range(ed, st, "A5", "B6", {"채우기": "nope"}, "화면")
    assert ei.value.sheet == "화면"
    assert st.xf_count() == before and not st.dirty and ed.cell("A5") is None
    with pytest.raises(ExportError, match="모양"):
        style_range(ed, st, "A5", "BB", {"굵게": True}, "화면")
    with pytest.raises(ExportError, match="cellXfs"):
        st.derive(9999, {"굵게": True})


def test_mutated_package_opens_and_other_parts_identical(xlsx: Path):
    ed, st = _load(xlsx)
    style_range(ed, st, "E5", "F6", {"채우기": "DDEBF7", "테두리": "얇게"}, "화면")
    merge_range(ed, "E5", "F5", "화면")
    out = _save(xlsx, xlsx.with_name("o.xlsx"), ed, st)
    with zipfile.ZipFile(xlsx) as a, zipfile.ZipFile(out) as b:
        assert a.namelist() == b.namelist()
        for n in a.namelist():
            if n not in (SHEET, STYLES):
                assert a.read(n) == b.read(n)
        import xml.etree.ElementTree as ET
        for n in (SHEET, STYLES):
            ET.fromstring(b.read(n))
    ws = openpyxl.load_workbook(out)["화면"]
    assert "E5:F5" in {str(r) for r in ws.merged_cells.ranges}
    assert ws["F6"].fill.fgColor.rgb == "FFDDEBF7"


# ---------------------------------------------------------------- 검토 반영(C1 재검토)

STYLES_XML = (
    '<?xml version="1.0"?><styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
    '<fonts count="1"><font><sz val="11"/><color theme="1"/><name val="Calibri"/><family val="2"/></font></fonts>'
    '<fills count="2"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill></fills>'
    '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
    '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
    '<cellXfs count="2"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
    '<xf numFmtId="14" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1" applyAlignment="1" applyProtection="1">'
    '<alignment indent="2" horizontal="left"/><protection locked="0"/></xf></cellXfs></styleSheet>'
)


def _sheet(body: str = "", rows: str = "", cols: str = "") -> SheetXml:
    xml = f'<worksheet><dimension ref="A1:B2"/>{cols}<sheetData>{rows}</sheetData>{body}</worksheet>'
    return SheetXml(xml, set(), False)


def test_unreadable_existing_merge_stops_overlap_check():
    for bad in ('<mergeCell ref="Zzz"/>', '<mergeCell ref="B2:C3x"/>', "<mergeCell/>"):
        ed = _sheet(f'<mergeCells count="1">{bad}</mergeCells><pageMargins/>')
        with pytest.raises(ExportError, match="검사할 수 없다"):
            merge_range(ed, "E5", "F6", "화면")
        assert bad in ed.post  # 쓰지 않았다
    ed = _sheet('<mergeCells count="1"><mergeCell ref="A1:B2"/>junk</mergeCells>')
    with pytest.raises(ExportError, match="읽지 못하는"):
        merge_range(ed, "E5", "F6")


def test_existing_cell_without_s_uses_xf0_not_column_style():
    st = StylesXml(STYLES_XML)
    col_xf = st.derive(0, {"채우기": "FFC000"})
    ed = _sheet(rows='<row r="2"><c r="A2"><v>1</v></c></row>', cols=f'<cols><col min="1" max="2" width="9" style="{col_xf}"/></cols>')
    style_range(ed, st, "A2", "B2", {"굵게": True})
    fill_of = lambda ref: dict(cell_style._xf_parts(st.xfs.items[int(ed.cell(ref).style)])[0])["fillId"]  # noqa: E731
    assert fill_of("A2") == "0"  # 셀 요소가 있고 s 가 없으면 서식 0 이다(열 서식을 따르지 않는다)
    assert fill_of("B2") != "0"  # 새로 만든 셀은 열 기본 서식 위에 쓴다


def test_new_cell_prefers_row_style_over_column_style():
    st = StylesXml(STYLES_XML)
    row_xf = st.derive(0, {"채우기": "112233"})
    col_xf = st.derive(0, {"채우기": "445566"})
    ed = _sheet(rows=f'<row r="3" s="{row_xf}" customFormat="1"/>', cols=f'<cols><col min="1" max="1" width="9" style="{col_xf}"/></cols>')
    style_range(ed, st, "A3", None, {"굵게": True})
    fill = dict(cell_style._xf_parts(st.xfs.items[int(ed.cell("A3").style)])[0])["fillId"]
    assert "112233" in st.fills.items[int(fill)]


def test_empty_selfclosing_mergecells_is_replaced_not_duplicated():
    for tag in ('<mergeCells count="0"/>', "<mergeCells/>"):
        ed = _sheet(f"{tag}<pageMargins/>")
        assert merges_of(ed) == []
        merge_range(ed, "A5", "B6")
        assert ed.post.count("<mergeCells") == 1
        assert ed.post.startswith('<mergeCells count="1"><mergeCell ref="A5:B6"/></mergeCells><pageMargins')
        assert merges_of(ed) == ["A5:B6"]


def test_unparsable_mergecells_tag_stops():
    ed = _sheet("<mergeCells count='1'")  # 닫히지 않은 태그
    with pytest.raises(ExportError, match="mergeCells"):
        merge_range(ed, "A5", "B6")


def test_merge_with_single_quoted_attrs():
    ed = _sheet("<mergeCells count='1'><mergeCell ref='A1:B2'/></mergeCells>")
    assert merges_of(ed) == ["A1:B2"]
    with pytest.raises(ExportError, match="겹친다"):
        merge_range(ed, "B2", "C3")
    merge_range(ed, "E5", "F6")
    assert merges_of(ed) == ["A1:B2", "E5:F6"] and ed.post.count("<mergeCells") == 1


def test_single_quoted_xf_attributes_are_kept():
    st = StylesXml(STYLES_XML.replace('numFmtId="14" fontId="0"', "numFmtId='14' fontId='0'"))
    n = st.derive(1, {"굵게": True})
    attrs = dict(cell_style._xf_parts(st.xfs.items[n])[0])
    assert attrs["numFmtId"] == "14" and attrs["applyNumberFormat"] == "1" and attrs["fontId"] != "0"
    # 속성을 읽지 못하면 말없이 버리지 않고 멈춘다.
    bad = StylesXml(STYLES_XML.replace('numFmtId="14"', "numFmtId=14"))
    with pytest.raises(ExportError, match="읽지 못했다"):
        bad.derive(1, {"굵게": True})


def test_merge_position_with_autofilter_cf_and_validations():
    ed = _sheet('<autoFilter ref="A1:B2"/><conditionalFormatting sqref="A1"/><dataValidations count="0"/><pageMargins/>')
    merge_range(ed, "A5", "B6")
    assert re.search(r'<autoFilter[^>]*/><mergeCells count="1">.*?</mergeCells><conditionalFormatting', ed.post)
    ed = _sheet('<sheetProtection sheet="1"/><dataValidations count="0"/>')
    merge_range(ed, "A5", "B6")
    assert re.search(r'<sheetProtection[^>]*/><mergeCells.*?</mergeCells><dataValidations', ed.post)
    ed = _sheet('<sheetProtection sheet="1"/>')
    merge_range(ed, "A5", "B6")
    assert ed.post.endswith("</mergeCells></worksheet>") or ed.render().endswith("</mergeCells></worksheet>")


def test_alignment_override_keeps_protection_and_other_alignment_attrs():
    st = StylesXml(STYLES_XML)
    n = st.derive(1, {"가로": "가운데", "줄바꿈": True})
    attrs, inner = cell_style._xf_parts(st.xfs.items[n])
    assert dict(attrs)["numFmtId"] == "14" and dict(attrs)["applyProtection"] == "1"
    assert '<protection locked="0"/>' in inner
    al = dict(cell_style._attrs(re.search(r"<alignment([^>]*?)/>", inner).group(1), "t"))
    assert al == {"indent": "2", "horizontal": "center", "wrapText": "1"}
    assert inner.count("<alignment") == 1 and inner.index("<alignment") < inner.index("<protection")


def test_style_range_is_atomic_on_failure():
    st = StylesXml(STYLES_XML)
    ed = _sheet(rows='<row r="1"><c r="A1" s="1"/></row>')
    before_xml, before_xfs = ed.render(), st.xf_count()
    real = st.derive
    calls = {"n": 0}

    def flaky(base, defn):
        calls["n"] += 1
        if calls["n"] == 3:
            raise ExportError("깨짐")
        return real(base, defn)

    st.derive = flaky  # type: ignore[method-assign]
    with pytest.raises(ExportError, match="깨짐") as ei:
        style_range(ed, st, "A1", "C2", {"채우기": "112233"}, "화면")
    assert ei.value.sheet == "화면"
    assert ed.render() == before_xml and not ed.dirty
    assert st.xf_count() == before_xfs and not st.dirty and st.render() == st.xml


def test_cell_xfs_limit(monkeypatch):
    monkeypatch.setattr(cell_style, "MAX_CELL_XFS", 3)
    st = StylesXml(STYLES_XML)
    st.derive(0, {"굵게": True})  # 3번째 xf
    ed = _sheet()
    with pytest.raises(ExportError, match="한도"):
        style_range(ed, st, "A1", None, {"굵게": False, "채우기": "123456"}, "화면")
    assert ed.cell("A1") is None and st.xf_count() == 3
    st.derive(0, {"굵게": True})  # 이미 있는 xf 는 한도에 걸리지 않고 재사용된다
