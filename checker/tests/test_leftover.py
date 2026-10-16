"""산출 결과에 템플릿 견본이 남으면 산출을 멈추는 검사(`checker.leftover`, 지도 `견본_검사`·`유지`)를 검증한다(이슈 #335, PR B).

합성 템플릿(openpyxl)에 지도 밖 견본(셀 값·그림·메모)을 심어 두고 `checker.export` 로 산출·`--check` 한다. 그림은 zip 수준으로 넣는다.
"""
from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path

import openpyxl
import pytest
import yaml
from openpyxl.comments import Comment

from checker import blank_template, export

REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
RELS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"

MAP = """\
양식: 시험양식
설명: 시험용
템플릿: ../시험.xlsx
%(top)s산출:
  폴더: docs/out
  파일명: "T_{ID}_v1.xlsx"
  자리:
    ID: {칸: 코드}
시트:
  - 이름: Main
    칸:
      - 항목: 코드
        위치: B2
        필수: true
      - 항목: 작성일
        위치: B3
        자동: true
    표:
      - 이름: 목록
        머리글_행: 5
        시작_행: 6
        최대_행: 3
        열: {번호: A, 이름: B}
%(main)s  - 이름: Other
%(other)s"""

SPEC = """\
# DEV-001 시험

## 양식 항목

### 시험양식 · Main

| 항목 | 값 |
|---|---|
| 코드 | X1 |

#### 목록

| 번호 | 이름 |
|---|---|
| 1 | 가 |

### 시험양식 · Other

해당 없음
"""


def _template(path: Path, *, other: str | None = "고정 글자", extra: dict | None = None, comment_at: str | None = None,
              drawing_at: str | None = None, picture: bool = False, dv: str | None = None, hf: str | None = None,
              empty_inline: bool = False) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Main"
    ws["A1"] = "라벨"  # 지도 밖 고정 글자
    ws["B2"] = "SAMPLE"  # 칸 코드: 산출이 덮어쓴다
    ws["B3"] = "=TODAY()"  # 자동 칸
    ws["A5"], ws["B5"] = "번호", "이름"  # 표 머리글 행
    ws["A6"], ws["B6"] = "9", "다른 개발 건"  # 표 데이터 영역 견본: 산출이 쓰거나 비운다
    ws["A7"], ws["B7"] = "8", "또 다른"
    for ref, v in (extra or {}).items():
        ws[ref] = v
    if comment_at:
        ws[comment_at].comment = Comment("남의 메모", "누군가")
    o = wb.create_sheet("Other")
    if other:
        o["A1"] = other
    buf = io.BytesIO()
    wb.save(buf)
    patched = _patch(buf.getvalue(), drawing_at, picture, dv, hf, empty_inline)
    path.write_bytes(patched)


def _patch(data: bytes, drawing_at: str | None, picture: bool, dv: str | None, hf: str | None, empty_inline: bool) -> bytes:
    if not (drawing_at or picture or dv or hf or empty_inline):
        return data
    with zipfile.ZipFile(io.BytesIO(data)) as zi:
        items = [(i, zi.read(i.filename)) for i in zi.infolist()]
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        new_rels = ""
        if drawing_at:
            new_rels += f'<Relationship Id="rId8" Type="{REL}drawing" Target="../drawings/drawing1.xml"/>'
        if picture:
            new_rels += f'<Relationship Id="rId7" Type="{REL}image" Target="../media/bg.png"/>'
        for info, body in items:
            if info.filename == "xl/worksheets/_rels/sheet1.xml.rels":  # openpyxl 이 메모 때문에 이미 만든 관계 파트에 덧붙인다
                body = body.replace(b"</Relationships>", new_rels.encode() + b"</Relationships>")
                new_rels = ""
            if info.filename == "xl/worksheets/sheet1.xml":
                x = body.decode("utf-8")
                x = x.replace("<worksheet ", f'<worksheet xmlns:r="{REL.rstrip("/")}" ', 1)
                if empty_inline:
                    x = x.replace("</sheetData>", '<row r="30"><c r="Z30" t="inlineStr"><is><t/></is></c></row></sheetData>', 1)
                tail = ""
                if dv:
                    tail += f'<dataValidations count="1"><dataValidation type="list" sqref="{dv}"><formula1>"Y,N"</formula1></dataValidation></dataValidations>'
                if hf:
                    tail += f"<headerFooter><oddFooter>{hf}</oddFooter></headerFooter>"
                if tail:
                    x = x.replace("<pageMargins", tail + "<pageMargins", 1) if "<pageMargins" in x else x.replace("</worksheet>", tail + "</worksheet>")
                if picture:
                    x = x.replace("</worksheet>", '<picture r:id="rId7"/></worksheet>')
                if drawing_at:
                    x = x.replace("</worksheet>", '<drawing r:id="rId8"/></worksheet>')
                body = x.encode("utf-8")
            zo.writestr(info.filename, body)
        if drawing_at:
            col = ord(drawing_at[0]) - 65
            row = int(drawing_at[1:]) - 1
            zo.writestr("xl/drawings/drawing1.xml", (
                '<?xml version="1.0"?><xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
                'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><xdr:twoCellAnchor>'
                f"<xdr:from><xdr:col>{col}</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{row}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>"
                f"<xdr:to><xdr:col>{col + 2}</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{row + 2}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>"
                "<xdr:sp><xdr:txBody><a:p><a:r><a:t>남의 도형</a:t></a:r></a:p></xdr:txBody></xdr:sp></xdr:twoCellAnchor></xdr:wsDr>"
            ).encode())
        if picture:
            zo.writestr("xl/media/bg.png", b"\x89PNG\r\n\x1a\nFAKE")
        if new_rels:
            zo.writestr("xl/worksheets/_rels/sheet1.xml.rels", f'<?xml version="1.0"?><Relationships xmlns="{RELS_NS}">{new_rels}</Relationships>')
    # [Content_Types] 는 시험에 필요 없다(엔진은 읽기만 하고 openpyxl 은 이 파트를 열지 않는다)
    return out.getvalue()


def _setup(tmp_path: Path, *, top: str = "견본_검사: true\n", main: str = "", other: str = "", **tpl) -> Path:
    (tmp_path / "templates" / "forms").mkdir(parents=True)
    _template(tmp_path / "templates" / "시험.xlsx", **tpl)
    _write_map(tmp_path, top=top, main=main, other=other)
    return tmp_path


def _write_map(root: Path, *, top: str = "견본_검사: true\n", main: str = "", other: str = "") -> None:
    other_body = other or "    설명: 고정 글자뿐인 시트\n"
    (root / "templates" / "forms" / "시험양식.yaml").write_text(MAP % {"top": top, "main": main, "other": other_body}, encoding="utf-8")


def _run(root: Path, capsys, *extra: str, spec: str = SPEC):
    sp = root / "spec.md"
    sp.write_text(spec, encoding="utf-8")
    code = export.main(["--spec", str(sp), "--root", str(root), *extra])
    out = capsys.readouterr().out
    assert out.count("\n") == 1
    return code, json.loads(out)


def _leftovers(report: dict) -> list[dict]:
    return [p for p in report["problems"] if p.get("type") == "leftover-sample"]


KEEP_MAIN = "    유지: [A1]\n"
KEEP_OTHER = "    유지: [A1]\n"


def _ok_setup(tmp_path: Path, **tpl) -> Path:
    return _setup(tmp_path, main=KEEP_MAIN, other=KEEP_OTHER, **tpl)


# ---------------------------------------------------------------- 거부와 허용

def test_지도_밖_견본이_남으면_거부하고_좌표와_값과_다음_안내를_준다(tmp_path, capsys):
    root = _setup(tmp_path, other=None)
    code, rep = _run(root, capsys)
    assert code == 1 and rep["status"] == "rejected" and rep["files"] == []
    cells = {c["cell"]: c["value"] for p in _leftovers(rep) if p["sheet"] == "Main" for c in p["cells"]}
    assert cells == {"A1": "라벨"}  # 칸·표 데이터·머리글 행·자동 칸은 견본이 아니다
    assert "A1" in _leftovers(rep)[0]["reason"] and "'라벨'" in _leftovers(rep)[0]["reason"]
    assert "blank_template" in rep["next"] and "`유지:`" in rep["next"]
    assert not (root / "docs").exists()


def test_유지_범위는_고정_문구로_허용한다(tmp_path, capsys):
    root = _ok_setup(tmp_path)
    code, rep = _run(root, capsys)
    assert code == 0, rep
    assert rep["status"] == "ok" and rep["problems"] == []
    assert (root / rep["files"][0]).exists()


def test_산출이_쓴_셀과_자동_칸과_머리글_행은_견본이_아니다(tmp_path, capsys):
    root = _ok_setup(tmp_path)
    code, rep = _run(root, capsys)
    assert code == 0, rep
    ws = openpyxl.load_workbook(root / rep["files"][0])["Main"]
    assert ws["B2"].value == "X1" and ws["B6"].value == "가" and ws["B7"].value is None
    assert ws["B3"].value == "=TODAY()" and ws["B5"].value == "이름"  # 자동 칸과 머리글은 템플릿 그대로


def test_템플릿_값과_같은_값을_쓰면_산출이_쓴_셀로_본다(tmp_path, capsys):
    root = _ok_setup(tmp_path, extra={"B2": "X1"})
    code, rep = _run(root, capsys)
    assert code == 0, rep


def test_머리글_행_밖의_표_라벨은_견본으로_잡힌다(tmp_path, capsys):
    root = _ok_setup(tmp_path, extra={"C5": "머리글 행의 다른 칸"})
    code, rep = _run(root, capsys)
    assert code == 0, rep  # 머리글 행은 행 전체가 고정 줄이다
    root2 = _ok_setup(tmp_path / "b", extra={"C4": "머리글 위 줄"})
    code, rep = _run(root2, capsys)
    assert code == 1 and {c["cell"] for p in _leftovers(rep) for c in p.get("cells", [])} == {"C4"}


def test_해당_없음_시트도_템플릿_글자가_남으면_거부한다(tmp_path, capsys):
    root = _setup(tmp_path, main=KEEP_MAIN)  # Other 는 해당 없음이고 A1 '고정 글자' 가 있다
    code, rep = _run(root, capsys)
    assert code == 1
    assert [p["sheet"] for p in _leftovers(rep)] == ["Other"]
    assert not any("템플릿 샘플" in w["reason"] for w in rep["warnings"])  # 예전 해당 없음 한정 경고는 대체됐다


def test_해당_없음_시트의_유지_범위는_통과한다(tmp_path, capsys):
    root = _ok_setup(tmp_path)
    code, rep = _run(root, capsys)
    assert code == 0, rep


def test_지도에_없는_시트의_값도_본다(tmp_path, capsys):
    root = _ok_setup(tmp_path)
    wb = openpyxl.load_workbook(root / "templates" / "시험.xlsx")
    wb.create_sheet("Lists")["A1"] = "목록값"
    wb.save(root / "templates" / "시험.xlsx")
    code, rep = _run(root, capsys)
    assert code == 1 and [p["sheet"] for p in _leftovers(rep)] == ["Lists"]
    _write_map(root, main=KEEP_MAIN, other=KEEP_OTHER)
    (root / "templates" / "forms" / "시험양식.yaml").write_text(
        (root / "templates" / "forms" / "시험양식.yaml").read_text(encoding="utf-8") + "  - 이름: Lists\n    유지: [A1:A9]\n", encoding="utf-8")
    code, rep = _run(root, capsys)
    assert code == 0, rep


def test_빈_문자열_셀은_값이_아니다(tmp_path, capsys):
    root = _ok_setup(tmp_path, empty_inline=True)
    code, rep = _run(root, capsys)
    assert code == 0, rep


# ---------------------------------------------------------------- 그림·메모·개체

def test_그림이_남으면_거부하고_앵커가_유지_범위_안이면_허용한다(tmp_path, capsys):
    root = _ok_setup(tmp_path, drawing_at="D5")
    code, rep = _run(root, capsys)
    assert code == 1
    p = [x for x in _leftovers(rep) if x["kind"] == "drawings"][0]
    assert p["sheet"] == "Main" and "D5" in p["reason"] and "남의 도형" in p["reason"] and p["count"] == 1
    _write_map(root, main="    유지: [A1, D5:E6]\n", other=KEEP_OTHER)
    code, rep = _run(root, capsys)
    assert code == 0, rep


def test_앵커_없는_그림은_위치_불명으로_막는다(tmp_path, capsys):
    root = _ok_setup(tmp_path, drawing_at="D5")
    p = root / "templates" / "시험.xlsx"
    with zipfile.ZipFile(p) as zi:
        items = [(i.filename, zi.read(i.filename)) for i in zi.infolist()]
    with zipfile.ZipFile(p, "w") as zo:
        for n, d in items:
            if n == "xl/drawings/drawing1.xml":
                d = b'<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"><xdr:sp/></xdr:wsDr>'
            zo.writestr(n, d)
    _write_map(root, main="    유지: [A1:Z99]\n", other=KEEP_OTHER)
    code, rep = _run(root, capsys)
    assert code == 1 and "위치 불명" in _leftovers(rep)[0]["reason"]


def test_메모가_남으면_거부하고_셀이_유지_범위_안이면_허용한다(tmp_path, capsys):
    root = _ok_setup(tmp_path, comment_at="D3")
    code, rep = _run(root, capsys)
    assert code == 1
    p = [x for x in _leftovers(rep) if x["kind"] == "comments"][0]
    assert p["sheet"] == "Main" and "D3" in p["reason"] and "남의 메모" in p["reason"]
    _write_map(root, main="    유지: [A1, D3]\n", other=KEEP_OTHER)
    code, rep = _run(root, capsys)
    assert code == 0, rep


def test_위치를_판정할_수_없는_개체는_막는다(tmp_path, capsys):
    root = _ok_setup(tmp_path, picture=True)
    code, rep = _run(root, capsys)
    assert code == 1
    assert any(p["kind"] == "objects" and "시트 배경 그림" in p["reason"] for p in _leftovers(rep))


# ---------------------------------------------------------------- 알림(거부 아님)

def test_직접_적은_드롭다운_목록과_머리글_바닥글은_경고로_알린다(tmp_path, capsys):
    root = _ok_setup(tmp_path, dv="D2:D4", hf="&amp;CAcme 고객사 &amp;P")
    code, rep = _run(root, capsys)
    assert code == 3 and rep["status"] == "warning" and rep["problems"] == []
    kinds = {w["item"] for w in rep["warnings"] if w["type"] == "leftover-review"}
    assert kinds == {"드롭다운 목록", "머리글·바닥글"}
    assert "leftover-review" in rep["next"]
    assert (root / rep["files"][0]).exists()  # 파일은 만들었다(경고)


def test_유지_범위_안의_드롭다운_목록은_알리지_않고_쪽번호_코드뿐인_바닥글도_알리지_않는다(tmp_path, capsys):
    root = _setup(tmp_path, main="    유지: [A1, D2:D4]\n", other=KEEP_OTHER, dv="D2:D4", hf="&amp;C&amp;P / &amp;N")
    code, rep = _run(root, capsys)
    assert code == 0 and not [w for w in rep["warnings"] if w.get("type") == "leftover-review"], rep


# ---------------------------------------------------------------- 끈 양식과 --check

@pytest.mark.parametrize("top", ["", "견본_검사: false\n"])
def test_견본_검사를_끄면_예전_동작_그대로다(tmp_path, capsys, top):
    root = _setup(tmp_path, top=top, drawing_at="D5", comment_at="D3")
    code, rep = _run(root, capsys)
    assert code == 0 and rep["problems"] == []
    # 예전의 해당 없음 한정 경고는 그대로 난다
    assert any("템플릿 샘플" in w["reason"] and w["sheet"] == "Other" for w in rep["warnings"])


def test_check_도_같은_검사로_거부하고_파일을_쓰지_않는다(tmp_path, capsys):
    root = _setup(tmp_path, main=KEEP_MAIN, drawing_at="D5")
    code, rep = _run(root, capsys, "--check")
    assert code == 1 and rep["status"] == "rejected"
    assert {p["sheet"] for p in _leftovers(rep)} == {"Main", "Other"}
    assert not (root / "docs").exists()
    _write_map(root, main="    유지: [A1, D5]\n", other=KEEP_OTHER)
    code, rep = _run(root, capsys, "--check")
    assert code == 0 and rep["dry_run"] is True and not (root / "docs").exists()


def test_산출물_바이트는_견본_검사_여부와_상관없다(tmp_path, capsys):
    a = _ok_setup(tmp_path / "a")
    _, ra = _run(a, capsys)
    b = _setup(tmp_path / "b", top="", main=KEEP_MAIN, other=KEEP_OTHER)
    # openpyxl 은 docProps/core.xml 에 저장 시각(초 단위)을 쓴다 — 템플릿을 따로 만들면 초가 바뀔 때 바이트가 달라져 간헐 실패한다. 같은 템플릿을 쓴다.
    (b / "templates" / "시험.xlsx").write_bytes((a / "templates" / "시험.xlsx").read_bytes())
    _, rb = _run(b, capsys)
    assert (a / ra["files"][0]).read_bytes() == (b / rb["files"][0]).read_bytes()


# ---------------------------------------------------------------- 지도 형식

@pytest.mark.parametrize("top", ["견본_검사: 예\n", "견본_검사: [true]\n", "견본_검사: 1\n"])
def test_견본_검사_값이_불리언이_아니면_지도_오류다(tmp_path, capsys, top):
    root = _ok_setup(tmp_path, top=top)
    code, rep = _run(root, capsys)
    assert code == 2 and "견본_검사" in rep["problems"][0]["reason"]


def test_유지_범위_모양_오류는_지도_오류다(tmp_path, capsys):
    root = _setup(tmp_path, main="    유지: [A1:]\n", other=KEEP_OTHER)
    code, rep = _run(root, capsys)
    assert code == 2 and "유지" in rep["problems"][0]["reason"]


def test_map_problems_가_견본_검사와_유지를_함께_본다():
    probs = blank_template.map_problems(yaml.safe_load("견본_검사: x\n시트:\n  - 이름: S\n    유지: B\n"))
    texts = " ".join(m for _, m in probs)
    assert "견본_검사" in texts and "유지" in texts
    assert blank_template.map_problems(yaml.safe_load("견본_검사: true\n시트:\n  - 이름: S\n    유지: [A1:C3]\n")) == []


def test_form_columns_도_견본_검사_값_오류를_검사_불능으로_본다():
    from checker.model import ConfigError
    from checker.rules import form_columns

    ok = {"견본_검사": True, "시트": [{"이름": "S", "유지": ["A1:B2"]}]}
    assert form_columns._valid_map(ok, Path("m.yaml")) is ok
    with pytest.raises(ConfigError, match="견본_검사"):
        form_columns._valid_map({"견본_검사": "예", "시트": [{"이름": "S"}]}, Path("m.yaml"))


# ---------------------------------------------------------------- 읽지 못한 것은 통과가 아니다

def _mutate(path: Path, edits: dict[str, object], add: dict[str, bytes] | None = None) -> None:
    """zip 안의 파트를 이름(접미사)으로 골라 고친다. edits 값은 bytes -> bytes 함수이거나 None(삭제)."""
    with zipfile.ZipFile(path) as zi:
        items = [(i.filename, zi.read(i.filename)) for i in zi.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zo:
        for name, data in items:
            hit = next((k for k in edits if name.endswith(k)), None)
            if hit is not None:
                if edits[hit] is None:
                    continue
                data = edits[hit](data)
            zo.writestr(name, data)
        for name, data in (add or {}).items():
            zo.writestr(name, data)


def _tpl(root: Path) -> Path:
    return root / "templates" / "시험.xlsx"


def test_공유_문자열_번호가_범위_밖이면_검사_불능으로_거부한다(tmp_path, capsys):
    root = _ok_setup(tmp_path, extra={"D9": "x"})
    _mutate(_tpl(root), {"worksheets/sheet1.xml": lambda d: d.replace(b'<c r="D9" t="inlineStr"><is><t>x</t></is></c>', b'<c r="D9" t="s"><v>999</v></c>')},
            add={"xl/sharedStrings.xml": b'<?xml version="1.0"?><sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="1" uniqueCount="1"><si><t>a</t></si></sst>'})
    code, rep = _run(root, capsys)
    assert code == 1
    bad = [p for p in rep["problems"] if p["type"] == "leftover-unchecked"]
    assert bad and bad[0]["sheet"] == "Main" and "D9" in bad[0]["reason"] and "999" in bad[0]["reason"]
    assert not _leftovers(rep)  # 위반으로 뭉개지도, 통과로 넘기지도 않는다


def test_sharedStrings_가_없는데_공유_문자열_셀이_있으면_검사_불능이다(tmp_path, capsys):
    root = _ok_setup(tmp_path, extra={"D9": "x"})
    _mutate(_tpl(root), {"worksheets/sheet1.xml": lambda d: d.replace(b'<c r="D9" t="inlineStr"><is><t>x</t></is></c>', b'<c r="D9" t="s"><v>0</v></c>')})
    code, rep = _run(root, capsys)
    assert code == 1 and any(p["type"] == "leftover-unchecked" and "sharedStrings" in p["reason"] for p in rep["problems"])


def test_v_요소에_속성이_붙은_값도_읽는다(tmp_path, capsys):
    root = _ok_setup(tmp_path, extra={"D9": "x"})
    _mutate(_tpl(root), {"worksheets/sheet1.xml": lambda d: d.replace(b'<c r="D9" t="inlineStr"><is><t>x</t></is></c>', b'<c r="D9"><v xml:space="preserve">77</v></c>')})
    code, rep = _run(root, capsys)
    assert code == 1
    assert {c["cell"]: c["value"] for p in _leftovers(rep) for c in p.get("cells", [])} == {"D9": "77"}


def test_메모_시트의_VML_에_메모가_아닌_개체가_있으면_견본이다(tmp_path, capsys):
    root = _ok_setup(tmp_path, comment_at="D3")
    _mutate(_tpl(root), {".vml": lambda d: d.replace(b"</xml>", b'<v:shape><x:ClientData ObjectType="Button"/></v:shape></xml>')})
    _write_map(root, main="    유지: [A1, D3]\n", other=KEEP_OTHER)  # 메모는 허용해도 버튼은 못 지나간다
    code, rep = _run(root, capsys)
    assert code == 1 and any(p["kind"] == "objects" and "Button" in p["reason"] for p in _leftovers(rep)), rep["problems"]


def test_메모만_있는_VML_은_개체로_보지_않는다(tmp_path, capsys):
    root = _ok_setup(tmp_path, comment_at="D3")
    _write_map(root, main="    유지: [A1, D3]\n", other=KEEP_OTHER)
    code, rep = _run(root, capsys)
    assert code == 0, rep


def test_차트_시트는_지원하지_않는다는_사유로_거부한다(tmp_path, capsys):
    root = _ok_setup(tmp_path)
    _mutate(
        _tpl(root),
        {
            "xl/workbook.xml": lambda d: d.replace(b"</sheets>", b'<sheet name="Chart1" sheetId="9" r:id="rIdChart"/></sheets>'),
            "xl/_rels/workbook.xml.rels": lambda d: d.replace(b"</Relationships>", f'<Relationship Id="rIdChart" Type="{REL}chartsheet" Target="chartsheets/sheet1.xml"/></Relationships>'.encode()),
        },
        add={"xl/chartsheets/sheet1.xml": b'<?xml version="1.0"?><chartsheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"/>'},
    )
    code, rep = _run(root, capsys)
    assert code == 1
    bad = [p for p in rep["problems"] if p["sheet"] == "Chart1"]
    assert bad and bad[0]["type"] == "leftover-unchecked" and "차트 시트" in bad[0]["reason"] and "sheetData" not in bad[0]["reason"]


def test_해석되지_않은_목록_유효성은_경고하고_읽다_예외가_나면_거부한다(tmp_path, capsys, monkeypatch):
    root = _ok_setup(tmp_path)
    _mutate(_tpl(root), {"worksheets/sheet1.xml": lambda d: d.replace(b"</worksheet>", b'<dataValidations count="1"><dataValidation type="list"><formula1>"A,B"</formula1></dataValidation></dataValidations></worksheet>')})
    code, rep = _run(root, capsys)
    assert code == 3 and any(w["type"] == "leftover-review" and "해석하지 못했다" in w["reason"] for w in rep["warnings"]), rep

    from checker import leftover

    def boom(self, sheet):
        raise ValueError("깨진 유효성")

    monkeypatch.setattr(leftover.ListReader, "validations", boom)
    code, rep = _run(root, capsys)
    assert code == 1 and any(p["type"] == "leftover-unchecked" and "깨진 유효성" in p["reason"] for p in rep["problems"])
    assert not [w for w in rep["warnings"] if w.get("type") == "leftover-review"]
