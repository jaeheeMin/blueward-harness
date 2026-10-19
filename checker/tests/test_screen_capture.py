"""Spec 의 화면 캡처 그림 넣기(`checker.screen_capture`)를 산출(`checker.export`) 끝에서 끝까지 검증한다(이슈 #348).

합성 템플릿(openpyxl)에 지도 `화면:` 의 `캡처:` 를 두고, 작업본 `#### 화면 1 캡처` 를 산출해 zip 안 파트·관계·시트 XML 과
openpyxl 로 다시 연 결과를 확인한다. PNG 는 표준 라이브러리로 만든다(Pillow 를 쓰지 않는다).
"""
from __future__ import annotations

import io
import json
import re
import struct
import zipfile
import zlib
from pathlib import Path

import openpyxl
import pytest

from checker import export, screen_capture, screen_draw

MAP = """\
양식: 시험양식
설명: 시험용
템플릿: ../시험.xlsx
%(top)s산출:
  폴더: docs/out
  파일명: "T_{ID}_v0.2.xlsx"
  자리:
    ID: {칸: 코드}
시트:
  - 이름: Screen
    칸:
      - 항목: 코드
        위치: B2
        필수: true
%(keep)s    비우기: [C7:S14, U7:V20]
    화면:
      - 이름: 화면 1
        영역: C7:S14
%(cap)s  - 이름: Other
%(other)s"""

CAP = "        캡처: {설명: U7:V20}\n"

SPEC_HEAD = """\
# DEV-001 시험

## 양식 항목

### 시험양식 · Screen

| 항목 | 값 |
|---|---|
| 코드 | X1 |

"""

IMG = "docs/spec/img/DEV-001_화면1_v0.2.png"


def png(w: int, h: int) -> bytes:
    """w x h 흰색 PNG."""
    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))

    raw = b"".join(b"\x00" + b"\xff\xff\xff" * w for _ in range(h))
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


def cap_section(rows: int = 3, path: str = IMG, name: str = "화면 1") -> str:
    table = "".join(f"| {k} | 설명 {k} |\n" for k in range(1, rows + 1))
    return f"#### {name} 캡처\n\n![화면 캡처]({path})\n\n| 번호 | 설명 |\n|---|---|\n{table}"


def _template(path: Path, sample: dict | None = None, drawing: bool = False) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Screen"
    ws["A2"], ws["B2"] = "코드", "SAMPLE"
    ws["C8"] = "견본 글자"
    ws["C30"] = "영역 밖"
    for ref, v in (sample or {}).items():
        ws[ref] = v
    wb.create_sheet("Other")["A1"] = "그대로"
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    if drawing:  # 시트에 이미 drawing 이 걸린 템플릿(zip 수준으로 흉내)
        out = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(data)) as zi, zipfile.ZipFile(out, "w") as zo:
            for info in zi.infolist():
                body = zi.read(info.filename)
                if info.filename == "xl/worksheets/sheet1.xml":
                    body = body.replace(b"</worksheet>", b'<drawing r:id="rId9"/></worksheet>')
                zo.writestr(info, body)
        data = out.getvalue()
    path.write_bytes(data)


def _write_map(proj: Path, cap: str = CAP, top: str = "", keep: str = "", other: str = "") -> None:
    body = MAP % {"top": top, "keep": keep, "cap": cap, "other": other}
    (proj / "templates" / "forms" / "시험양식.yaml").write_text(body, encoding="utf-8")


@pytest.fixture
def proj(tmp_path: Path) -> Path:
    (tmp_path / "templates" / "forms").mkdir(parents=True)
    _template(tmp_path / "templates" / "시험.xlsx")
    _write_map(tmp_path)
    (tmp_path / "docs" / "spec" / "img").mkdir(parents=True)
    (tmp_path / IMG).write_bytes(png(400, 100))
    return tmp_path


def _run(proj: Path, capsys, body: str, *extra: str):
    spec = proj / "spec.md"
    spec.write_text(SPEC_HEAD + body, encoding="utf-8")
    code = export.main(["--spec", str(spec), "--root", str(proj), *extra])
    out = capsys.readouterr().out
    assert out.count("\n") == 1
    return code, json.loads(out)


def _zip(proj: Path, report: dict) -> zipfile.ZipFile:
    assert len(report["files"]) == 1, report
    return zipfile.ZipFile(proj / report["files"][0])


def _reasons(report: dict) -> str:
    return " / ".join(p["reason"] for p in report.get("problems", [])) + " " + str(report.get("reason", ""))


# ---------------------------------------------------------------- 정상 삽입

def test_그림_파트_관계_콘텐츠_타입_시트_위치(proj, capsys):
    code, report = _run(proj, capsys, cap_section())
    assert code == 0, report
    z = _zip(proj, report)
    names = z.namelist()
    assert "xl/drawings/drawing1.xml" in names and "xl/media/image1.png" in names
    assert "xl/drawings/_rels/drawing1.xml.rels" in names
    assert z.read("xl/media/image1.png") == (proj / IMG).read_bytes()
    assert 'Target="../media/image1.png"' in z.read("xl/drawings/_rels/drawing1.xml.rels").decode()
    ct = z.read("[Content_Types].xml").decode()
    assert '/xl/drawings/drawing1.xml' in ct and 'drawing+xml' in ct
    assert len(re.findall(r'Extension="png"', ct, re.I)) == 1
    rels = z.read("xl/worksheets/_rels/sheet1.xml.rels").decode()
    assert 'Target="../drawings/drawing1.xml"' in rels
    rid = re.search(r'Id="([^"]+)"[^>]*/drawing"|Type="[^"]*/drawing"[^>]*Id="([^"]+)"', rels)
    rid = next(g for g in rid.groups() if g)
    xml = z.read("xl/worksheets/sheet1.xml").decode()
    assert f'<drawing r:id="{rid}"/>' in xml and 'xmlns:r=' in xml
    assert xml.index("<drawing ") > xml.index("</sheetData>")
    assert xml.index("<drawing ") < xml.index("</worksheet>")
    for later in ("<legacyDrawing", "<tableParts", "<extLst"):
        if later in xml:
            assert xml.index("<drawing ") < xml.index(later)
    for earlier in ("<pageMargins", "<pageSetup"):
        if earlier in xml:
            assert xml.index(earlier) < xml.index("<drawing ")


def test_그림은_영역_왼쪽_위에_원본_크기로_들어간다(proj, capsys):
    code, report = _run(proj, capsys, cap_section())
    assert code == 0, report
    d = _zip(proj, report).read("xl/drawings/drawing1.xml").decode()
    assert "<xdr:col>2</xdr:col>" in d and "<xdr:row>6</xdr:row>" in d  # C7 → 0부터 센 열 2, 행 6
    # 영역 1088x160px 안에 400x100px: 늘리지 않는다
    assert f'<xdr:ext cx="{400 * 9525}" cy="{100 * 9525}"/>' in d
    assert 'name="화면 1 캡처"' in d and 'descr="화면 캡처"' in d


def test_영역보다_크면_비율을_지켜_줄인다(proj, capsys):
    (proj / IMG).write_bytes(png(1000, 200))  # 영역 1088x160px → 높이에 맞춰 0.8배
    code, report = _run(proj, capsys, cap_section())
    assert code == 0, report
    d = _zip(proj, report).read("xl/drawings/drawing1.xml").decode()
    assert f'<xdr:ext cx="{800 * 9525}" cy="{160 * 9525}"/>' in d


def test_열_너비와_행_높이를_반영한다(proj, capsys):
    wb = openpyxl.load_workbook(proj / "templates" / "시험.xlsx")
    ws = wb["Screen"]
    for c in "CDEFGHIJKLMNOPQRS":
        ws.column_dimensions[c].width = 10  # 10*7+0.5 → 70px
    for r in range(7, 15):
        ws.row_dimensions[r].height = 30  # 30pt = 40px
    wb.save(proj / "templates" / "시험.xlsx")
    (proj / IMG).write_bytes(png(2000, 100))
    code, report = _run(proj, capsys, cap_section())
    assert code == 0, report
    d = _zip(proj, report).read("xl/drawings/drawing1.xml").decode()
    box_w = 17 * 70 * 9525  # 1190px
    assert f'<xdr:ext cx="{box_w}" cy="{int(100 * 9525 * box_w / (2000 * 9525))}"' in d


def test_설명은_설명_범위_첫_열에_원문자_번호로_쓴다(proj, capsys):
    code, report = _run(proj, capsys, cap_section(rows=3))
    assert code == 0, report
    ws = openpyxl.load_workbook(proj / report["files"][0])["Screen"]
    assert [ws[f"U{r}"].value for r in (7, 8, 9, 10)] == ["① 설명 1", "② 설명 2", "③ 설명 3", None]
    assert ws["C8"].value is None and ws["C30"].value == "영역 밖"


def test_21번부터는_괄호_숫자다():
    assert [screen_capture.circled(n) for n in (1, 10, 20, 21, 30)] == ["①", "⑩", "⑳", "(21)", "(30)"]


def test_openpyxl이_그림_하나를_본다(proj, capsys):
    pytest.importorskip("PIL")  # openpyxl 이 이미지를 읽을 때만 필요하다(엔진은 쓰지 않는다)
    code, report = _run(proj, capsys, cap_section())
    assert code == 0, report
    ws = openpyxl.load_workbook(proj / report["files"][0])["Screen"]
    assert len(ws._images) == 1
    assert ws._images[0].anchor._from.col == 2 and ws._images[0].anchor._from.row == 6


def test_기존_파트_이름과_겹치지_않는_번호를_쓴다(proj, capsys):
    tpl = proj / "templates" / "시험.xlsx"
    out = io.BytesIO()
    with zipfile.ZipFile(tpl) as zi, zipfile.ZipFile(out, "w") as zo:
        for info in zi.infolist():
            zo.writestr(info, zi.read(info.filename))
        zo.writestr("xl/drawings/drawing1.xml", b"<x/>")  # 다른 시트가 쓰는 drawing 이 이미 있다고 치고 번호만 점유
        zo.writestr("xl/media/image1.png", b"x")
        zo.writestr("xl/media/image3.jpeg", b"x")
    tpl.write_bytes(out.getvalue())
    code, report = _run(proj, capsys, cap_section())
    assert code == 0, report
    z = _zip(proj, report)
    assert "xl/drawings/drawing2.xml" in z.namelist() and "xl/media/image4.png" in z.namelist()
    assert z.read("xl/drawings/drawing1.xml") == b"<x/>"


def test_캡처가_없는_지도_화면은_그림_절이_없으면_경고만_한다(proj, capsys):
    code, report = _run(proj, capsys, "")
    assert code == 0, report
    assert any("캡처" in w["reason"] and w["item"] == "화면 1 캡처" for w in report["warnings"])
    assert "xl/drawings/drawing1.xml" not in _zip(proj, report).namelist()


def test_남은_견본_검사와_함께_써도_새_그림을_견본으로_보지_않는다(proj, capsys):
    _write_map(proj, top="견본_검사: true\n", keep="    유지: [A2, C30]\n", other="    유지: [A1]\n")
    code, report = _run(proj, capsys, cap_section())
    assert code == 0, report
    assert "xl/drawings/drawing1.xml" in _zip(proj, report).namelist()


def test_check는_파일을_쓰지_않고_같은_검사를_한다(proj, capsys):
    code, report = _run(proj, capsys, cap_section(), "--check")
    assert code == 0 and report["dry_run"], report
    assert not (proj / "docs" / "out").exists()


# ---------------------------------------------------------------- 멈추는 경우

def test_그림_파일이_없으면_멈춘다(proj, capsys):
    (proj / IMG).unlink()
    code, report = _run(proj, capsys, cap_section())
    assert code == 1 and "없다" in _reasons(report), report
    assert not (proj / "docs" / "out").exists()


def test_PNG가_아니면_멈춘다(proj, capsys):
    (proj / IMG).write_bytes(b"GIF89a....")
    code, report = _run(proj, capsys, cap_section())
    assert code == 1 and "PNG" in _reasons(report), report


def test_캡처가_옛_판이면_멈춘다(proj, capsys):
    old = "docs/spec/img/DEV-001_화면1_v0.1.png"
    (proj / old).write_bytes(png(400, 100))
    code, report = _run(proj, capsys, cap_section(path=old))
    assert code == 1 and "옛 판" in _reasons(report) and "v0.1" in _reasons(report) and "v0.2" in _reasons(report), report
    assert not (proj / "docs" / "out").exists()


def test_파일_이름에_판이_없으면_멈춘다(proj, capsys):
    bare = "docs/spec/img/화면1.png"
    (proj / bare).write_bytes(png(400, 100))
    code, report = _run(proj, capsys, cap_section(path=bare))
    assert code == 1 and "판 번호" in _reasons(report), report


def test_대문자_V도_같은_판으로_본다(proj, capsys):
    cap = "docs/spec/img/화면1_V0.2.png"
    (proj / cap).write_bytes(png(400, 100))
    code, report = _run(proj, capsys, cap_section(path=cap))
    assert code == 0, report


def test_그림과_캡처_영역이_겹치면_멈춘다(proj, capsys):
    shape = "#### 화면 1 그림\n\n##### 결과 테이블\n\n| 열 |\n|---|\n| 판매문서 |\n\n"
    code, report = _run(proj, capsys, shape + cap_section())
    assert code == 1 and "둘 다" in _reasons(report) and "겹친다" in _reasons(report), report


def test_그림과_캡처_영역이_따로_잡히면_둘_다_넣는다(proj, capsys):
    _write_map(proj, cap="        캡처: {영역: C16:S24, 설명: U7:V20}\n")
    (proj / "templates" / "forms" / "시험양식.yaml").write_text(
        (proj / "templates" / "forms" / "시험양식.yaml").read_text(encoding="utf-8").replace("U7:V20]", "U7:V20, C16:S24]"), encoding="utf-8")
    shape = "#### 화면 1 그림\n\n##### 결과 테이블\n\n| 열 |\n|---|\n| 판매문서 |\n\n"
    code, report = _run(proj, capsys, shape + cap_section())
    assert code == 0, report
    ws = openpyxl.load_workbook(proj / report["files"][0])["Screen"]
    assert ws["C7"].value == "판매문서" and ws["U7"].value == "① 설명 1"
    d = _zip(proj, report).read("xl/drawings/drawing1.xml").decode()
    assert "<xdr:row>15</xdr:row>" in d  # C16


def test_너무_많이_줄여야_하면_멈춘다(proj, capsys):
    (proj / IMG).write_bytes(png(2000, 400))  # 영역 1088x160 → 0.4배
    code, report = _run(proj, capsys, cap_section())
    assert code == 1, report
    r = _reasons(report)
    assert "40%" in r and "50%" in r and "C7:S14" in r and "영역을 넓히거나" in r, r


def test_정확히_50퍼센트면_들어간다(proj, capsys):
    (proj / IMG).write_bytes(png(2000, 320))  # 높이 160 → 0.5배
    code, report = _run(proj, capsys, cap_section())
    assert code == 0, report


def test_설명_행이_모자라면_멈춘다(proj, capsys):
    code, report = _run(proj, capsys, cap_section(rows=15))  # U7:V20 은 14행
    assert code == 1 and "15행" in _reasons(report) and "14행" in _reasons(report), report


def test_설명_범위에_견본_값이_있고_비우기_대상이_아니면_멈춘다(proj, capsys):
    _template(proj / "templates" / "시험.xlsx", sample={"U9": "견본 설명"})
    _write_map(proj)
    (proj / "templates" / "forms" / "시험양식.yaml").write_text(
        (proj / "templates" / "forms" / "시험양식.yaml").read_text(encoding="utf-8").replace("    비우기: [C7:S14, U7:V20]\n", "    비우기: [C7:S14]\n"), encoding="utf-8")
    code, report = _run(proj, capsys, cap_section())
    assert code == 1 and "U9" in _reasons(report), report


def test_설명_범위의_견본은_비우기가_지운다(proj, capsys):
    _template(proj / "templates" / "시험.xlsx", sample={"U9": "견본 설명", "V12": "견본"})
    code, report = _run(proj, capsys, cap_section(rows=1))
    assert code == 0, report
    ws = openpyxl.load_workbook(proj / report["files"][0])["Screen"]
    assert ws["U9"].value is None and ws["V12"].value is None and ws["U7"].value == "① 설명 1"


def test_시트에_이미_그림이_있으면_멈춘다(proj, capsys):
    _template(proj / "templates" / "시험.xlsx", drawing=True)
    code, report = _run(proj, capsys, cap_section())
    assert code == 1 and "이미 그림" in _reasons(report), report


def test_설명이_비면_멈춘다(proj, capsys):
    body = cap_section(rows=2).replace("| 2 | 설명 2 |", "| 2 | 미정 |")
    code, report = _run(proj, capsys, body)
    assert code == 1 and "2번 설명" in _reasons(report), report


# ---------------------------------------------------------------- Spec 절 형식(저장·PR 에서도 같은 검사)

def test_지도에_캡처가_없는_화면의_캡처_절은_양식_밖이다(proj, capsys):
    _write_map(proj, cap="")
    code, report = _run(proj, capsys, cap_section())
    assert code == 2 and "양식 밖" in _reasons(report) and "화면 1 캡처" in _reasons(report), report


def _item():
    return {"이름": "화면 1"}


@pytest.mark.parametrize(
    "body, kind",
    [
        ("![a](docs/x_v0.2.png)\n\n| 번호 | 설명 |\n|---|---|\n| 1 | a |\n\n설명 글\n", "그림·표 밖 글"),
        ("| 번호 | 설명 |\n|---|---|\n| 1 | a |\n", "캡처 그림"),
        ("![a](docs/x_v0.2.png)\n![b](docs/y_v0.2.png)\n\n| 번호 | 설명 |\n|---|---|\n| 1 | a |\n", "캡처 그림"),
        ("![a](docs/x_v0.2.jpg)\n\n| 번호 | 설명 |\n|---|---|\n| 1 | a |\n", "캡처 그림 경로"),
        ("![a](https://x.test/x_v0.2.png)\n\n| 번호 | 설명 |\n|---|---|\n| 1 | a |\n", "캡처 그림 경로"),
        ("![a](/abs/x_v0.2.png)\n\n| 번호 | 설명 |\n|---|---|\n| 1 | a |\n", "캡처 그림 경로"),
        ("![a](../x_v0.2.png)\n\n| 번호 | 설명 |\n|---|---|\n| 1 | a |\n", "캡처 그림 경로"),
        ("![a](docs/x_v0.2.png)\n", "캡처 표 머리"),
        ("![a](docs/x_v0.2.png)\n\n| 번호 | 내용 |\n|---|---|\n| 1 | a |\n", "캡처 표 머리"),
        ("![a](docs/x_v0.2.png)\n\n| 번호 | 설명 |\n|---|---|\n", "캡처 표 행"),
        ("![a](docs/x_v0.2.png)\n\n| 번호 | 설명 |\n|---|---|\n| 1 | a | b |\n", "캡처 표 행"),
        ("![a](docs/x_v0.2.png)\n\n| 번호 | 설명 |\n|---|---|\n| 2 | a |\n", "캡처 번호"),
        ("![a](docs/x_v0.2.png)\n\n| 번호 | 설명 |\n|---|---|\n| 1 | a |\n| 3 | b |\n", "캡처 번호"),
        ("| 번호 | 설명 |\n|---|---|\n| 1 | a |\n\n![a](docs/x_v0.2.png)\n", "캡처 순서"),
    ],
)
def test_캡처_절_형식이_틀리면_양식_밖이다(body, kind):
    off = screen_draw.capture_offenders(_item(), body.splitlines())
    assert any(kind in k for k, _, _ in off), off


def test_올바른_캡처_절은_위반이_없다():
    body = "![화면](docs/spec/img/a_v0.2.png)\n\n| 번호 | 설명 |\n|---|---|\n| 1 | a |\n| 2 | b |\n".splitlines()
    assert screen_draw.capture_offenders(_item(), body) == []


def test_형식이_틀린_캡처_절은_산출도_양식_밖으로_거절한다(proj, capsys):
    code, report = _run(proj, capsys, "#### 화면 1 캡처\n\n설명만 있다\n")
    assert code == 2 and "양식 밖" in _reasons(report), report


def test_spec_offenders가_캡처_절을_검사한다(proj):
    smap = {
        "시트": [{"이름": "Screen", "화면": [{"이름": "화면 1", "영역": "C7:S14", "캡처": {"설명": "U7:V20"}}]}]
    }["시트"][0]
    sec, subs = export.parse_spec(SPEC_HEAD + "#### 화면 1 캡처\n\n설명만 있다\n")["시험양식"]["Screen"]
    assert export.spec_offenders(sec, subs, smap)


# ---------------------------------------------------------------- 지도 검증

def _fmap(*screens: dict, **sheet) -> dict:
    return {"양식": "x", "템플릿": "t.xlsx", "시트": [{"이름": "S", "화면": list(screens), **sheet}]}


def _p(*screens: dict, **sheet) -> list[str]:
    return [why for _, why in screen_draw.map_problems(_fmap(*screens, **sheet))]


S1 = {"이름": "화면 1", "영역": "C7:S14"}


def test_캡처_지도가_옳으면_문제가_없다():
    assert _p({**S1, "캡처": {"설명": "U7:V20"}}) == []
    assert _p({**S1, "캡처": {"영역": "C20:S30", "설명": "U7:V20"}}) == []


def test_같은_화면의_그림_영역과_캡처_영역이_같거나_겹쳐도_지도_오류는_아니다():
    assert _p({**S1, "캡처": {"영역": "C7:S14", "설명": "U7:V20"}}) == []
    assert _p({**S1, "캡처": {"영역": "D8:E9", "설명": "U7:V20"}}) == []


@pytest.mark.parametrize(
    "cap, why",
    [
        ({"영역": "C20:S30"}, "설명"),
        ({"설명": "U7:V20", "모양": 1}, "모르는 키"),
        ({"설명": "U7"}, None),
        ({"설명": "7U"}, "모양이 틀렸다"),
        ({"영역": "S30:C20", "설명": "U7:V20"}, "시작이 끝보다 뒤"),
        ("U7:V20", "목록"),
        ({"설명": "D8:E9"}, "겹친다"),  # 같은 화면의 그림 영역
        ({"영역": "C7:S14", "설명": "S14:T15"}, "겹친다"),
    ],
)
def test_캡처_지도가_틀리면_알린다(cap, why):
    got = _p({**S1, "캡처": cap})
    if why is None:
        assert got == []
    else:
        assert any(why in g for g in got), got


def test_캡처_범위가_맡은_자리와_겹치면_알린다():
    assert any("겹친다" in g for g in _p({**S1, "캡처": {"설명": "U7:V20"}}, 칸=[{"항목": "코드", "위치": "V10"}]))
    assert any("겹친다" in g for g in _p({**S1, "캡처": {"영역": "C20:S30", "설명": "U7:V20"}}, 유지=["D25"]))


def test_다른_화면과_겹치면_알린다():
    s2 = {"이름": "화면 2", "영역": "C43:Y51", "캡처": {"설명": "AA43:AB60"}}
    assert _p({**S1, "캡처": {"설명": "U7:V20"}}, s2) == []
    assert any("겹친다" in g for g in _p({**S1, "캡처": {"설명": "U7:V20"}}, {**s2, "캡처": {"설명": "V10:W12"}}))
    assert any("겹친다" in g for g in _p({**S1, "캡처": {"영역": "C43:D44", "설명": "U7:V20"}}, s2))
    assert any("겹친다" in g for g in _p({**S1, "캡처": {"설명": "U7:V20"}}, {**s2, "영역": "V8:W9"}))


def test_캡처_절_이름이_표나_서술과_같으면_알린다():
    got = _p({**S1, "캡처": {"설명": "U7:V20"}}, 표=[{"이름": "화면 1 캡처", "열": {"a": "AA"}, "시작_행": 3}])
    assert any("구분할 수 없다" in g for g in got), got


def test_items가_캡처_영역_기본값을_채운다():
    items = screen_draw.items(_fmap({**S1, "캡처": {"설명": " U7:V20 "}})["시트"][0])
    assert items[0]["캡처"] == {"영역": "C7:S14", "설명": "U7:V20"}
    assert screen_draw.items(_fmap(S1)["시트"][0])[0]["캡처"] is None


def test_png_크기는_헤더에서_읽는다():
    assert screen_capture.png_size(png(37, 5)) == (37, 5)
    assert screen_capture.png_size(b"\x89PNG\r\n\x1a\n" + b"x" * 40) is None
    assert screen_capture.png_size(b"not a png at all, long enough to pass length" * 2) is None


def test_파일_이름의_판():
    f = screen_capture.file_version
    assert f("a_v0.1.png") == "0.1" and f("a_V0.12.png") == "0.12" and f("a_v1.png") == "1"
    assert f("a_v0.1_v0.2.png") == "0.2" and f("a.png") is None and f("avatar.png") is None
    assert f("a_v0.1_최종.png") is None and f("a_v0.1x.png") is None and f("a_v0.png") == "0"


def test_판이_파일_이름_끝이_아니면_멈춘다(proj, capsys):
    mid = "docs/spec/img/화면1_v0.2_최종.png"
    (proj / mid).write_bytes(png(400, 100))
    code, report = _run(proj, capsys, cap_section(path=mid))
    assert code == 1 and "판 번호는 파일 이름 끝" in _reasons(report), report
    assert not (proj / "docs" / "out").exists()


_HEAD = '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"/></sheetData>'


def test_drawing은_cfRule_안_extLst_앞이_아니라_최상위_자리에_들어간다():
    xml = (
        _HEAD + '<conditionalFormatting sqref="A1"><cfRule type="dataBar" priority="1"><dataBar/>'
        '<extLst><ext uri="{x}"><x14:id>1</x14:id></ext></extLst></cfRule></conditionalFormatting>'
        '<pageMargins left="1"/></worksheet>'
    )
    got = screen_capture._with_drawing(xml, "rId9")
    assert got.index("<drawing ") > got.index("</conditionalFormatting>")
    assert got.index("<drawing ") > got.index("<pageMargins") and got.endswith('<drawing r:id="rId9"/></worksheet>')


def test_drawing은_최상위_tableParts_extLst_앞에_들어간다():
    tail = '<tableParts count="1"><tablePart r:id="rId1"/></tableParts><extLst><ext uri="{y}"/></extLst></worksheet>'
    xml = _HEAD + '<conditionalFormatting sqref="A1"><cfRule type="x"><extLst><ext uri="{z}"/></extLst></cfRule></conditionalFormatting>' + tail
    got = screen_capture._with_drawing(xml, "rId9")
    assert got.endswith('<drawing r:id="rId9"/>' + tail)
    only_ext = _HEAD + '<extLst><ext uri="{y}"/></extLst></worksheet>'
    got = screen_capture._with_drawing(only_ext, "rId9")
    assert got.index("<drawing ") < got.index("<extLst")


def test_작은따옴표_속성의_png_Default와_rId도_알아본다():
    raw = {
        "[Content_Types].xml": b"<Types xmlns='x'><Default Extension='png' ContentType='image/png'/></Types>",
        "xl/worksheets/_rels/sheet1.xml.rels": b"<Relationships xmlns='x'><Relationship Id='rId1' Type='t' Target='a'/></Relationships>",
    }
    pl = screen_capture.Placement("화면 1 캡처", b"p", "a", 1, 1, 10, 10, "0.1", "x")
    changed, _ = screen_capture.assemble(raw, {"S": "xl/worksheets/sheet1.xml"}, {"xl/worksheets/sheet1.xml": _HEAD + "</worksheet>"}, {"S": [pl]})
    assert changed["[Content_Types].xml"].decode().count("Extension=") == 1
    assert "Id=\"rId2\"" in changed["xl/worksheets/_rels/sheet1.xml.rels"].decode()
