"""Spec 의 화면 구성 그리기(`checker.screen_draw`)를 산출(`checker.export`) 끝에서 끝까지 검증한다(이슈 #335, PR C2).

합성 템플릿(openpyxl)에 지도 `화면:` 영역을 두고, 작업본 `#### 화면 1 그림` 을 산출해 openpyxl 로 다시 열어 확인한다.
기대 좌표는 사용자가 엑셀로 보고 승인한 시안(영역 C7:S14, 필터 바 3·액션 2·결과 테이블 4열)에서 가져왔다.
"""
from __future__ import annotations

import io
import json
import re
import zipfile
from pathlib import Path

import openpyxl
import pytest

from checker import export, screen_draw

MAP = """\
양식: 시험양식
설명: 시험용
템플릿: ../시험.xlsx
산출:
  폴더: docs/out
  파일명: "T_{ID}_v1.xlsx"
  자리:
    ID: {칸: 코드}
시트:
  - 이름: Screen
    칸:
      - 항목: 코드
        위치: B2
        필수: true
    비우기: [C7:S14]
    화면:
      - 이름: 화면 1
        영역: C7:S14
%(extra)s  - 이름: Other
    설명: 건드리지 않는다
"""

SPEC_HEAD = """\
# DEV-001 시험

## 양식 항목

### 시험양식 · Screen

| 항목 | 값 |
|---|---|
| 코드 | X1 |

"""

COND = "| 이름 | 필수 |\n|---|---|\n| 판매 조직 | Y |\n| 기간 | Y |\n| 판매처 | |\n"
BTN = "| 이름 |\n|---|\n| 취소 |\n| 일괄 취소 |\n"
RES = "| 열 |\n|---|\n| 판매문서 |\n| 품목 |\n| 수량 |\n| 상태 |\n"


def screen(cond: str | None = COND, btn: str | None = BTN, res: str | None = RES, name: str = "화면 1", settings: str | None = None) -> str:
    out = f"#### {name} 그림\n\n" + (f"{settings}\n" if settings else "")
    for title, body in (("필터 바", cond), ("액션", btn), ("결과 테이블", res)):
        if body is not None:
            out += f"##### {title}\n\n{body}\n"
    return out


def _template(path: Path, merge: str | None = None, sample: dict | None = None) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Screen"
    ws["A2"], ws["B2"] = "코드", "SAMPLE"
    ws["C8"] = "견본 글자"  # 지도 `비우기` 가 먼저 지운 뒤 그림을 그려야 한다
    ws["C30"] = "영역 밖"
    for ref, v in (sample or {}).items():
        ws[ref] = v
    if merge:
        ws.merge_cells(merge)
    other = wb.create_sheet("Other")
    other["A1"] = "그대로"
    buf = io.BytesIO()
    wb.save(buf)
    path.write_bytes(buf.getvalue())


@pytest.fixture
def proj(tmp_path: Path) -> Path:
    (tmp_path / "templates" / "forms").mkdir(parents=True)
    _template(tmp_path / "templates" / "시험.xlsx")
    _map(tmp_path)
    return tmp_path


def _map(proj: Path, extra: str = "", text: str | None = None) -> None:
    body = (text if text is not None else MAP) % {"extra": extra}
    (proj / "templates" / "forms" / "시험양식.yaml").write_text(body, encoding="utf-8")


def _run(proj: Path, capsys, body: str, *extra: str):
    spec = proj / "spec.md"
    spec.write_text(SPEC_HEAD + body, encoding="utf-8")
    code = export.main(["--spec", str(spec), "--root", str(proj), *extra])
    out = capsys.readouterr().out
    assert out.count("\n") == 1
    return code, json.loads(out)


def _sheet(proj: Path, report: dict):
    assert len(report["files"]) == 1, report
    return openpyxl.load_workbook(proj / report["files"][0])["Screen"]


def _nothing_written(proj: Path) -> bool:
    return not (proj / "docs").exists()


def _merges(ws) -> set[str]:
    return {str(r) for r in ws.merged_cells.ranges}


# ---------------------------------------------------------------- 시안과 같은 그림

def test_시안과_같은_좌표_병합_글자(proj, capsys):
    code, report = _run(proj, capsys, screen())
    assert code == 0, report
    ws = _sheet(proj, report)
    assert _merges(ws) == {
        # 필터 바 3개: 라벨 2칸 + 입력 3칸, 묶음 사이 1칸
        "C7:D7", "E7:G7", "I7:J7", "K7:M7", "O7:P7", "Q7:S7",
        # 액션 2개(2칸, 사이 1칸): 필터 바 뒤 빈 줄 한 줄(8행) 다음 9행, 영역 오른쪽 끝(S열)에 붙여 오른쪽 정렬
        "O9:P9", "R9:S9",
        # 결과 테이블: 머리글(바로 아래 10행) + 빈 행 2줄, 열마다 2칸
        "C10:D10", "E10:F10", "G10:H10", "I10:J10",
        "C11:D11", "E11:F11", "G11:H11", "I11:J11",
        "C12:D12", "E12:F12", "G12:H12", "I12:J12",
    }
    got = {c: ws[c].value for c in ("C7", "I7", "O7", "O9", "R9", "C10", "E10", "G10", "I10")}
    assert got == {
        "C7": "판매 조직 *", "I7": "기간 *", "O7": "판매처", "O9": "취소", "R9": "일괄 취소",
        "C10": "판매문서", "E10": "품목", "G10": "수량", "I10": "상태",
    }
    assert ws["C8"].value is None  # 견본 글자는 비우기로 지워졌고 그림이 그 뒤에 그려졌다
    assert ws["C30"].value == "영역 밖"
    assert ws["B2"].value == "X1"


def test_서식은_시안과_같다(proj, capsys):
    _, report = _run(proj, capsys, screen())
    ws = _sheet(proj, report)
    label, field, button, head, row = ws["C7"], ws["E7"], ws["O9"], ws["C10"], ws["C11"]
    assert label.fill.fgColor.rgb == "FFD9E1F2" and label.font.b and label.border.left.style == "thin"
    assert field.fill.fgColor.rgb == "FFFFFFFF" and field.border.top.style == "thin"
    assert button.fill.fgColor.rgb == "FFD9D9D9" and button.alignment.horizontal == "center"
    assert head.fill.fgColor.rgb == "FF1F4E78" and head.font.b and head.font.color.rgb == "FFFFFFFF"
    assert head.alignment.horizontal == "center"
    assert row.border.bottom.style == "thin" and row.fill.fill_type is None
    # 병합 칸은 범위 전체에 같은 서식이 있다(오른쪽 칸 테두리도 보인다). openpyxl 은 병합된 칸의 서식을 안 읽어 XML 로 본다.
    with zipfile.ZipFile(proj / report["files"][0]) as z:
        xml = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
    s = {ref: re.search(rf'<c r="{ref}"[^>]*?\bs="(\d+)"', xml).group(1) for ref in ("C7", "D7", "E7", "G7", "O9", "P9")}
    assert s["C7"] == s["D7"] != s["E7"] == s["G7"] != s["O9"] == s["P9"]


def test_조건이_한_줄을_넘으면_다음_줄로_줄바꿈한다(proj, capsys):
    cond = "| 이름 | 필수 |\n|---|---|\n" + "".join(f"| 조건{i} | |\n" for i in range(1, 5))
    code, report = _run(proj, capsys, screen(cond=cond))
    assert code == 0, report
    ws = _sheet(proj, report)
    assert ws["C7"].value == "조건1" and ws["I7"].value == "조건2" and ws["O7"].value == "조건3"
    assert ws["C8"].value == "조건4"  # 4번째는 다음 줄 첫 자리
    assert "C8:D8" in _merges(ws) and "E8:G8" in _merges(ws)
    # 필터 바가 두 줄이면 빈 줄 하나(9행) 뒤에 액션(10행, 오른쪽 정렬), 바로 아래 결과 테이블(11행)
    assert ws["O10"].value == "취소" and ws["R10"].value == "일괄 취소" and ws["C11"].value == "판매문서"


def test_결과_테이블만_있으면_영역_맨_위부터_그린다(proj, capsys):
    code, report = _run(proj, capsys, screen(cond=None, btn=None))
    assert code == 0, report
    ws = _sheet(proj, report)
    assert ws["C7"].value == "판매문서"
    assert "C7:D7" in _merges(ws)


def test_필터_바만_있으면_액션과_테이블을_건너뛴다(proj, capsys):
    code, report = _run(proj, capsys, screen(btn=None, res=None))
    assert code == 0, report
    ws = _sheet(proj, report)
    assert _merges(ws) == {"C7:D7", "E7:G7", "I7:J7", "K7:M7", "O7:P7", "Q7:S7"}


def test_결과_테이블이_없으면_액션은_필터_바_뒤_빈_줄_다음_줄에_오른쪽_정렬(proj, capsys):
    code, report = _run(proj, capsys, screen(res=None))
    assert code == 0, report
    ws = _sheet(proj, report)
    assert ws["O9"].value == "취소" and ws["R9"].value == "일괄 취소" and ws["C9"].value is None


def test_필터_바도_결과_테이블도_없으면_액션은_영역_맨_위_오른쪽에(proj, capsys):
    code, report = _run(proj, capsys, screen(cond=None, res=None))
    assert code == 0, report
    ws = _sheet(proj, report)
    assert _merges(ws) == {"O7:P7", "R7:S7"} and ws["R7"].value == "일괄 취소"


def test_액션_한_개는_마지막_열에서_끝난다(proj, capsys):
    code, report = _run(proj, capsys, screen(cond=None, res=None, btn="| 이름 |\n|---|\n| 열기 |\n"))
    assert code == 0, report
    ws = _sheet(proj, report)
    assert _merges(ws) == {"R7:S7"} and ws["R7"].value == "열기"


def test_액션이_영역_너비보다_많으면_거부한다(proj, capsys):
    btn = "| 이름 |\n|---|\n" + "".join(f"| 액션{i} |\n" for i in range(1, 8))  # 7개 x 3칸 - 1 = 20칸 > 17칸
    code, report = _run(proj, capsys, screen(cond=None, res=None, btn=btn))
    assert code == 1 and "액션 '액션1'(열)" in report["problems"][0]["reason"]


# ---------------------------------------------------------------- 조회 방식(자동 / Go)

SET_GO = "| 항목 | 값 |\n|---|---|\n| 조회 | Go |\n"


def test_Go_조회는_필터_바_마지막_줄_오른쪽_끝에_Go_액션을_그린다(proj, capsys):
    cond = "| 이름 | 필수 |\n|---|---|\n| 판매 조직 | Y |\n| 기간 | Y |\n"  # 7행: C~G, I~M 까지 쓰고 Q~S 가 남는다
    code, report = _run(proj, capsys, screen(cond=cond, settings=SET_GO))
    assert code == 0, report
    ws = _sheet(proj, report)
    assert ws["R7"].value == "Go" and "R7:S7" in _merges(ws)
    assert ws["R7"].fill.fgColor.rgb == "FFD9D9D9"  # 액션 스타일
    assert ws["O9"].value == "취소" and ws["C10"].value == "판매문서"  # 필터 바 뒤 빈 줄(8행) 그대로


def test_Go_자리가_없으면_다음_줄_오른쪽_끝(proj, capsys):
    cond = "| 이름 | 필수 |\n|---|---|\n| 가 | |\n| 나 | |\n| 다 | |\n"  # 7행이 Q~S 까지 꽉 참
    code, report = _run(proj, capsys, screen(cond=cond, settings=SET_GO))
    assert code == 0, report
    ws = _sheet(proj, report)
    assert ws["R7"].value is None and ws["R8"].value == "Go" and "R8:S8" in _merges(ws)
    assert ws["O10"].value == "취소"  # 빈 줄(9행) 다음 줄


def test_Go만_있고_필터_바가_없으면_영역_맨_위_오른쪽(proj, capsys):
    code, report = _run(proj, capsys, screen(cond=None, btn=None, res=None, settings=SET_GO))
    assert code == 0, report
    assert _sheet(proj, report)["R7"].value == "Go"


def test_Go가_영역을_넘으면_거부한다(proj, capsys):
    _map(proj, text=MAP.replace("영역: C7:S14", "영역: C7:S8"))
    cond = "| 이름 |\n|---|\n| 가 |\n| 나 |\n| 다 |\n| 라 |\n"  # 7행 3개, 8행 1개 → Go 는 8행에 자리가 있다
    code, report = _run(proj, capsys, screen(cond=cond, btn=None, res=None, settings=SET_GO), "--check")
    assert code == 0, report
    cond = "| 이름 |\n|---|\n| 가 |\n| 나 |\n| 다 |\n| 라 |\n| 마 |\n| 바 |\n"  # 8행도 꽉 참 → Go 가 9행
    code, report = _run(proj, capsys, screen(cond=cond, btn=None, res=None, settings=SET_GO), "--check")
    assert code == 1 and "Go 액션(행)" in report["problems"][0]["reason"]


def test_조회_자동은_기본이고_표시해도_Go를_그리지_않는다(proj, capsys):
    code, report = _run(proj, capsys, screen(settings="| 항목 | 값 |\n|---|---|\n| 조회 | 자동 |\n"))
    assert code == 0, report
    ws = _sheet(proj, report)
    assert all(ws.cell(7, c).value != "Go" for c in range(3, 20))


@pytest.mark.parametrize("settings,why", [
    ("| 항목 | 값 |\n|---|---|\n| 모드 | Go |\n", "화면 1 그림: 모드"),          # 모르는 키
    ("| 항목 | 값 |\n|---|---|\n| 조회 | 수동 |\n", "화면 항목 값"),              # 모르는 값
    ("| 항목 | 값 |\n|---|---|\n| 조회 | go |\n", "화면 항목 값"),                # 대소문자도 정확히
    ("| 항목 | 값 |\n|---|---|\n| 조회 | Go |\n| 조회 | 자동 |\n", "두 번"),
    ("| 이름 | 값 |\n|---|---|\n| 조회 | Go |\n", "화면 항목 표 머리"),
    ("| 항목 | 값 |\n|---|---|\n| 조회 | Go | 추가 |\n", "화면 표 행"),
])
def test_조회_표의_모르는_키_값은_거부한다(proj, capsys, settings, why):
    for extra in ((), ("--check",)):
        code, report = _run(proj, capsys, screen(settings=settings), *extra)
        assert code == 2, report
        assert why in report["problems"][0]["reason"]
    assert _nothing_written(proj)


def test_필수_열을_생략해도_된다(proj, capsys):
    cond = "| 이름 |\n|---|\n| 판매 조직 |\n"
    code, report = _run(proj, capsys, screen(cond=cond, btn=None, res=None))
    assert code == 0, report
    assert _sheet(proj, report)["C7"].value == "판매 조직"


def test_그림_절이_없으면_영역을_그리지_않고_비워_둔다(proj, capsys):
    code, report = _run(proj, capsys, "")
    assert code == 0, report
    ws = _sheet(proj, report)
    assert _merges(ws) == set()
    assert all(ws.cell(r, c).value is None for r in range(7, 15) for c in range(3, 20))  # 비우기가 견본을 지웠다
    assert ws["C7"].fill.fill_type is None


def test_산출물은_openpyxl로_열리고_다른_시트와_파트는_그대로다(proj, capsys):
    with zipfile.ZipFile(proj / "templates" / "시험.xlsx") as z:
        before = {n: z.read(n) for n in z.namelist()}
    code, report = _run(proj, capsys, screen())
    assert code == 0, report
    with zipfile.ZipFile(proj / report["files"][0]) as z:
        after = {n: z.read(n) for n in z.namelist()}
    assert after["xl/styles.xml"] != before["xl/styles.xml"]  # 새 서식이 styles.xml 에 덧붙었다
    assert after["xl/styles.xml"].count(b'rgb="FF1F4E78"') == 1
    assert after["xl/worksheets/sheet2.xml"] == before["xl/worksheets/sheet2.xml"]  # Other 시트
    assert openpyxl.load_workbook(proj / report["files"][0])["Other"]["A1"].value == "그대로"


# ---------------------------------------------------------------- 지도 값으로 덮어쓰기

def test_지도의_스타일과_칸_수가_기본값을_덮어쓴다(proj, capsys):
    _map(proj, extra=(
        "        칸_수: {라벨: 1, 입력: 2, 간격: 0, 열: 3, 빈_행: 1}\n"
        "        스타일:\n"
        "          머리글: {채우기: FF0000, 글꼴_색: '000000'}\n"
    ))
    code, report = _run(proj, capsys, screen(btn=None))
    assert code == 0, report
    ws = _sheet(proj, report)
    head = ws["C9"]
    assert head.fill.fgColor.rgb == "FFFF0000" and head.font.color.rgb == "FF000000"
    assert not head.font.b  # 스타일은 이름 단위로 통째로 바뀐다(굵게도 기본값이 아니라 지도 값)
    # 라벨 1칸(병합 없음) + 입력 2칸, 간격 없음 → 묶음이 3칸씩. 열 3칸, 빈 행 1줄.
    assert _merges(ws) == {"D7:E7", "G7:H7", "J7:K7", "C9:E9", "F9:H9", "I9:K9", "L9:N9", "C10:E10", "F10:H10", "I10:K10", "L10:N10"}
    assert ws["C7"].value == "판매 조직 *" and ws["D7"].value is None and ws["F7"].value == "기간 *"


# ---------------------------------------------------------------- 영역을 넘으면 거절

def test_열이_영역을_넘으면_거절하고_넘친_항목을_알린다(proj, capsys):
    res = "| 열 |\n|---|\n" + "".join(f"| 열{i} |\n" for i in range(1, 10))  # 9열 x 2칸 = 18칸 > 17칸
    code, report = _run(proj, capsys, screen(res=res))
    assert code == 1
    assert report["status"] == "rejected" and report["files"] == []
    reason = report["problems"][0]["reason"]
    assert "화면 1 그림" in reason and "C7:S14" in reason and "열9" in reason and "열)" in reason
    assert "열1'" not in reason
    assert _nothing_written(proj)


def test_행이_영역을_넘으면_거절한다(proj, capsys):
    cond = "| 이름 |\n|---|\n" + "".join(f"| 조건{i} |\n" for i in range(1, 11))  # 3개씩 4줄(7~10행)
    code, report = _run(proj, capsys, screen(cond=cond))
    assert code == 1
    reason = report["problems"][0]["reason"]
    assert "(행)" in reason and "결과 테이블" in reason  # 표 머리글·빈 행이 14행을 넘는다
    assert _nothing_written(proj)


def test_조건_묶음이_영역_너비보다_넓으면_모두_넘친_것으로_본다(proj, capsys):
    _map(proj, text=MAP.replace("영역: C7:S14", "영역: C7:F14"))
    code, report = _run(proj, capsys, screen(btn=None, res=None))
    assert code == 1
    assert "판매 조직" in report["problems"][0]["reason"]


def test_check도_같은_이유로_거절한다(proj, capsys):
    res = "| 열 |\n|---|\n" + "".join(f"| 열{i} |\n" for i in range(1, 10))
    code, report = _run(proj, capsys, screen(res=res), "--check")
    assert code == 1 and report["status"] == "rejected"
    assert _nothing_written(proj)
    code, report = _run(proj, capsys, screen(), "--check")
    assert code == 0 and report["dry_run"] is True and report["files"]
    assert _nothing_written(proj)


# ---------------------------------------------------------------- 지도에 없는 것은 거절(읽지 못함)

@pytest.mark.parametrize("body,why", [
    (screen(name="화면 9"), "화면 9 그림"),                               # 지도에 없는 그림 이름
    (screen().replace("##### 액션", "##### 필터"), "필터"),               # 소절 이름이 셋 밖
    (screen(cond="| 이름 | 비고 |\n|---|---|\n| a | b |\n"), "비고"),    # 표 머리가 틀림
    (screen(btn="| 버튼 |\n|---|\n| a |\n"), "버튼"),                    # 액션 표 머리가 틀림
    (screen(res="| 이름 |\n|---|\n| a |\n"), "결과 테이블"),             # 결과 테이블 머리는 `열`
    (screen().replace("##### 액션", "##### 버튼"), "버튼"),              # 옛 소절 이름은 받지 않는다
    (screen().replace("##### 필터 바", "##### 조회 조건"), "조회 조건"),
    (screen().replace("##### 결과 테이블", "##### 결과 표"), "결과 표"),
    (screen() + "\n그림 설명 글\n", "표 밖 글"),                          # 표·소절 밖 글
    (screen(btn=""), "표 없음"),                                           # 소절에 표가 없음
])
def test_지도에_없는_이름_소절_표머리는_조용히_무시하지_않고_거절한다(proj, capsys, body, why):
    code, report = _run(proj, capsys, body)
    assert code == 2, report
    assert why in report["problems"][0]["reason"]
    assert _nothing_written(proj)
    code, report = _run(proj, capsys, body, "--check")
    assert code == 2 and _nothing_written(proj)


def test_지도에_화면이_없으면_그림_절은_양식_밖이다(proj, capsys):
    _map(proj, text=MAP.replace("    화면:\n      - 이름: 화면 1\n        영역: C7:S14\n", ""))
    code, report = _run(proj, capsys, screen())
    assert code == 2
    assert "화면 1 그림" in report["problems"][0]["reason"]


def test_필수_값이_Y_N_빈칸이_아니면_구조_검사에서_거절한다(proj, capsys):
    body = screen(cond="| 이름 | 필수 |\n|---|---|\n| 기간 | 예 |\n")
    for extra in ((), ("--check",)):
        code, report = _run(proj, capsys, body, *extra)
        assert code == 2 and "화면 필수 값" in report["problems"][0]["reason"]
    # 저장 훅·PR 과 같은 검사(offenders)가 먼저 본다
    item = screen_draw.items({"화면": [{"이름": "화면 1", "영역": "C7:S14"}]})[0]
    off = screen_draw.offenders(item, body.splitlines()[2:])
    assert [k for k, _, _ in off] == ["화면 필수 값"]


def test_셀_수가_머리글과_다른_행도_구조_검사에서_거절한다(proj, capsys):
    code, report = _run(proj, capsys, screen(cond="| 이름 | 필수 |\n|---|---|\n| 기간 | Y | 추가 |\n"))
    assert code == 2 and "화면 표 행" in report["problems"][0]["reason"]
    assert _nothing_written(proj)


def test_이름_칸이_미정이거나_비면_거부한다(proj, capsys):
    for bad, word in (("미정", "미정"), ("", "비어")):
        code, report = _run(proj, capsys, screen(btn=f"| 이름 |\n|---|\n| 취소 |\n| {bad} |\n"))
        assert code == 1, report
        assert word in report["problems"][0]["reason"] and "2번째 행" in report["problems"][0]["reason"]
        assert _nothing_written(proj)


# ---------------------------------------------------------------- 그리지 않은 화면은 경고

def _drawn_warnings(report: dict) -> list[dict]:
    return [w for w in report["warnings"] if "그리지 않았다" in w["reason"]]


def test_그림_절이_없으면_경고하되_거부하지_않는다(proj, capsys):
    code, report = _run(proj, capsys, "")
    assert code == 0
    ws = _drawn_warnings(report)
    assert len(ws) == 1 and ws[0]["item"] == "화면 1 그림" and ws[0]["sheet"] == "Screen" and "절이 없어" in ws[0]["reason"]


def test_소절이_하나도_없는_그림_절도_경고한다(proj, capsys):
    code, report = _run(proj, capsys, "#### 화면 1 그림\n\n")
    assert code == 0, report
    ws = _drawn_warnings(report)
    assert len(ws) == 1 and "비어 있어" in ws[0]["reason"]
    assert _merges(_sheet(proj, report)) == set()


def _check_spec_text(proj: Path, capsys, text: str) -> dict:
    spec = proj / "spec.md"
    spec.write_text(text, encoding="utf-8")
    export.main(["--spec", str(spec), "--root", str(proj), "--check"])
    return json.loads(capsys.readouterr().out)


def test_시트를_해당_없음으로_두면_경고하지_않는다(proj, capsys):
    # 의도된 생략이다(예: 객체 1건으로 바로 들어가는 앱이 아닐 때의 Initial Screen). 다른 이유의 거부(필수 칸)와 상관없이 경고는 없다.
    report = _check_spec_text(proj, capsys, "# D\n\n## 양식 항목\n\n### 시험양식 · Screen\n\n해당 없음\n")
    assert _drawn_warnings(report) == []


def test_작업본에_시트_절_자체가_없으면_경고한다(proj, capsys):
    report = _check_spec_text(proj, capsys, "# D\n\n## 양식 항목\n\n### 시험양식 · Other\n\n해당 없음\n")
    assert len(_drawn_warnings(report)) == 1, report


def test_해당_없음이_아닌데_그림_절만_없으면_경고한다(proj, capsys):
    report = _check_spec_text(proj, capsys, SPEC_HEAD)
    assert len(_drawn_warnings(report)) == 1, report


def test_그린_화면은_경고가_없다(proj, capsys):
    _, report = _run(proj, capsys, screen())
    assert _drawn_warnings(report) == []


# ---------------------------------------------------------------- 그림 글자는 늘 문자열

def test_숫자처럼_보이는_이름도_문자열_셀로_쓴다(proj, capsys):
    code, report = _run(proj, capsys, screen(cond="| 이름 |\n|---|\n| 2026 |\n", btn="| 이름 |\n|---|\n| 10 |\n", res="| 열 |\n|---|\n| 3 |\n"))
    assert code == 0, report
    ws = _sheet(proj, report)
    assert ws["C7"].value == "2026" and ws["R9"].value == "10" and ws["C10"].value == "3"
    assert all(isinstance(ws[c].value, str) for c in ("C7", "R9", "C10"))


# ---------------------------------------------------------------- 템플릿 상태

def test_영역_안에_견본_값이_남아_있으면_그리지_않는다(proj, capsys):
    _map(proj, text=MAP.replace("    비우기: [C7:S14]\n", ""))  # 견본 C8 이 그대로 남는다
    code, report = _run(proj, capsys, screen())
    assert code == 1
    assert "C8" in report["problems"][0]["reason"] and "비우기" in report["problems"][0]["reason"]
    assert _nothing_written(proj)


def test_영역_안에_기존_병합이_겹치면_멈춘다(proj, capsys):
    _template(proj / "templates" / "시험.xlsx", merge="D7:E8")
    code, report = _run(proj, capsys, screen())
    assert code == 2
    assert "겹친다" in report["problems"][0]["reason"]
    assert _nothing_written(proj)


# ---------------------------------------------------------------- 지도 검증

@pytest.mark.parametrize("extra,why", [
    ("        칸_수: {라벨: 0}\n", "칸_수 라벨"),
    ("        칸_수: {버튼수: 2}\n", "버튼수"),
    ("        칸_수: {버튼: 2}\n", "칸_수 키"),                                   # 옛 이름은 받지 않는다
    ("        스타일: {버튼: {채우기: FF0000}}\n", "스타일 이름"),
    ("        칸_수: {열: true}\n", "칸_수 열"),
    ("        스타일: {머리: {채우기: FF0000}}\n", "스타일 이름"),
    ("        스타일: {머리글: {색깔: 빨강}}\n", "색깔"),
    ("        스타일: {머리글: {채우기: 빨강}}\n", "채우기"),
    ("        스타일: {머리글: {}}\n", "비어 있다"),
    ("        색: 파랑\n", "모르는 키"),
])
def test_지도의_화면_값이_틀리면_산출이_멈춘다(proj, capsys, extra, why):
    _map(proj, extra=extra)
    code, report = _run(proj, capsys, screen())
    assert code == 2
    assert why in report["problems"][0]["reason"] and "화면 1" in report["problems"][0]["reason"]


@pytest.mark.parametrize("old,new,why", [
    ("영역: C7:S14", "영역: 가나", "영역"),
    ("        영역: C7:S14\n", "", "`영역` 이 없다"),
    ("영역: C7:S14", "영역: B2:C3", "겹친다"),  # 칸 코드(B2) 위치와 겹침
])
def test_영역이_틀리거나_지도가_맡은_자리와_겹치면_멈춘다(proj, capsys, old, new, why):
    _map(proj, text=MAP.replace(old, new))
    code, report = _run(proj, capsys, screen())
    assert code == 2 and why in report["problems"][0]["reason"]


def test_화면_영역끼리_겹치거나_이름이_같으면_멈춘다():
    fmap = {"시트": [{"이름": "S", "화면": [
        {"이름": "화면 1", "영역": "C7:S14"}, {"이름": "화면 2", "영역": "R14:T20"}, {"이름": "화면 1", "영역": "C30:D31"},
    ]}]}
    msgs = " | ".join(m for _, m in screen_draw.map_problems(fmap))
    assert "영역과 겹친다" in msgs and "같은 이름" in msgs


def test_화면_영역이_표_열과_겹치면_멈춘다():
    fmap = {"시트": [{"이름": "S", "표": [{"이름": "T", "시작_행": 9, "열": {"a": "D"}}],
                     "화면": [{"이름": "화면 1", "영역": "C7:S14"}]}]}
    assert any("표 'T'" in m for _, m in screen_draw.map_problems(fmap))


def test_SRS_원본_지도에_화면이_있으면_지도_오류다():
    fmap = {"시트": [{"이름": "S", "표": [{"이름": "T", "원본": "SRS", "시작_행": 4, "열": {"a": "A"}}],
                     "화면": [{"이름": "화면 1", "영역": "C7:S14"}]}]}
    assert any("원본이 SRS" in m for _, m in screen_draw.map_problems(fmap))


@pytest.mark.parametrize("key", ["칸", "서술", "표", "유지"])
def test_목록이_아닌_칸_표_서술_유지는_TypeError_없이_지도_오류다(key):
    fmap = {"시트": [{"이름": "S", key: 5, "화면": [{"이름": "화면 1", "영역": "C7:S14"}]}]}
    msgs = [m for _, m in screen_draw.map_problems(fmap)]
    assert any(f"`{key}`(목록이어야 한다)" in m for m in msgs)


def test_영역_문자열_오류는_range_종류다():
    kinds = {k for k, _ in screen_draw.map_problems({"시트": [{"이름": "S", "화면": [{"이름": "a", "영역": "zz"}]}]})}
    assert kinds == {"range"}


# ---------------------------------------------------------------- 배치 함수 단독

def test_layout_좌표():
    item = screen_draw.items({"화면": [{"이름": "화면 1", "영역": "C7:S14"}]})[0]
    boxes, over = screen_draw.layout(item, [("a", True), ("b", False)], ["취소"], ["x"])
    assert over == []
    assert [(b.col, b.row, b.width, b.kind, b.text) for b in boxes][:3] == [
        (3, 7, 2, "라벨", "a *"), (5, 7, 3, "입력", ""), (9, 7, 2, "라벨", "b"),
    ]
    assert (18, 9, 2, "액션", "취소") in boxes and (3, 10, 2, "머리글", "x") in boxes and (3, 12, 2, "행", "") in boxes
