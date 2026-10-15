"""견본 xlsx → 빈 양식(`checker.blank_template`)을 검증한다(이슈 #335, PR A).

고객 양식은 쓰지 않는다. 합성 xlsx 를 openpyxl 로 만든 뒤 zip 수준으로 drawing·그림·printerSettings·
x14 드롭다운·externalLinks 파트를 끼워 넣는다.
"""
from __future__ import annotations

import io
import json
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl
import pytest
import yaml

from checker import blank_template
from checker.model import ConfigError
from checker.rules import form_columns

REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
RELS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
DRAWING_CT = "application/vnd.openxmlformats-officedocument.drawing+xml"
PNG = b"\x89PNG\r\n\x1a\nFAKE"
PRINTER = b"\x00printer-settings\x01"


def _rels(*items: tuple[str, str, str]) -> bytes:
    body = "".join(f'<Relationship Id="{i}" Type="{t}" Target="{g}"/>' for i, t, g in items)
    return f'<?xml version="1.0"?><Relationships xmlns="{RELS_NS}">{body}</Relationships>'.encode()


def _drawing(text: str) -> bytes:
    return (
        '<?xml version="1.0"?><xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f"<xdr:sp><xdr:txBody><a:p><a:r><a:t>{text}</a:t></a:r></a:p></xdr:txBody></xdr:sp><xdr:pic/></xdr:wsDr>"
    ).encode()


def share_strings(path: Path) -> None:
    """openpyxl 이 인라인 문자열로 쓴 칸을 공유 문자열(sharedStrings)로 바꾼다(엑셀이 저장한 파일 모양)."""
    pool: list[str] = []
    refs = 0
    pat = re.compile(r'<c r="([A-Z]+\d+)"((?: s="\d+")?) t="inlineStr"><is><t(?: [^>]*)?>(.*?)</t></is></c>', re.S)

    def sub(m: re.Match) -> str:
        nonlocal refs
        text = m.group(3)
        if text not in pool:
            pool.append(text)
        refs += 1
        return f'<c r="{m.group(1)}"{m.group(2)} t="s"><v>{pool.index(text)}</v></c>'

    with zipfile.ZipFile(path) as zi:
        items = [(i, zi.read(i.filename)) for i in zi.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zo:
        for info, data in items:
            if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", info.filename):
                data = pat.sub(sub, data.decode("utf-8")).encode("utf-8")
            elif info.filename == "xl/_rels/workbook.xml.rels":
                data = data.replace(b"</Relationships>", f'<Relationship Id="rIdSst" Type="{REL}sharedStrings" Target="sharedStrings.xml"/></Relationships>'.encode())
            elif info.filename == "[Content_Types].xml":
                data = data.replace(b"</Types>", b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')
            zo.writestr(info.filename, data)
        body = "".join(f"<si><t>{t}</t></si>" for t in pool)
        zo.writestr(
            "xl/sharedStrings.xml",
            f'<?xml version="1.0"?><sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="{refs}" uniqueCount="{len(pool)}">{body}</sst>',
        )


def _make_src(
    path: Path, *, chart: bool = False, legacy: bool = False, ext: str | None = None, ext_name_use: str | None = None,
    hf_picture: bool = False, review: bool = False,
) -> None:
    """합성 견본. 시트: Info(sheet1) / Pic1(sheet2) / Pic2(sheet3) / Drop(sheet4, x14 드롭다운).

    ext: None 이면 externalLinks 없음. "dv"(표준 드롭다운이 외부 이름을 씀) / "x14" / "formula" / "free"(쓰는 곳 없음).
    """
    wb = openpyxl.Workbook()
    info = wb.active
    info.title = "Info"
    info["A1"], info["B2"], info["B3"], info["C2"], info["D2"] = "코드", "샘플A", "샘플B", "코드", 5
    info["C3"] = "=1+1"
    info["E5"] = "공통글자"
    wb.properties.title, wb.properties.creator, wb.properties.lastModifiedBy = "고객 견본", "홍길동", "김철수"
    wb.create_sheet("Pic1")["A1"] = "그림시트1"
    wb.create_sheet("Pic2")["A1"] = "그림시트2"
    wb.create_sheet("Drop")["A1"] = "목록"
    buf = io.BytesIO()
    wb.save(buf)

    def add_r(x: bytes) -> bytes:
        return x.replace(b"<worksheet ", b'<worksheet xmlns:r="' + REL.rstrip("/").encode() + b'" ', 1)

    mods = {}
    mods["xl/worksheets/sheet2.xml"] = lambda x: add_r(x).replace(
        b"</worksheet>",
        (b'<legacyDrawing r:id="rId9"/>' if legacy else b"") + b'<drawing r:id="rId1"/></worksheet>',
    )
    x14 = (
        '<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}" '
        'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main" '
        'xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main"><x14:dataValidations count="1">'
        '<x14:dataValidation type="list"><x14:formula1><xm:f>Drop!$A$1:$A$3</xm:f></x14:formula1>'
        "<xm:sqref>B1:B5</xm:sqref></x14:dataValidation></x14:dataValidations></ext></extLst>"
    ).encode()
    mods["xl/worksheets/sheet4.xml"] = lambda x: x.replace(b"</worksheet>", x14 + b"</worksheet>")
    if ext == "dv":
        mods["xl/worksheets/sheet1.xml"] = lambda x: x.replace(
            b"<pageMargins",
            b'<dataValidations count="1"><dataValidation type="list" sqref="B2:B3"><formula1>Pri</formula1>'
            b"</dataValidation></dataValidations><pageMargins", 1)
    if ext == "formula":
        mods["xl/worksheets/sheet1.xml"] = lambda x: x.replace(b"<f>1+1</f>", b"<f>[1]Defs!A1</f>")
    if ext == "cf":  # 접두사 없는 표준 조건부 서식
        mods["xl/worksheets/sheet1.xml"] = lambda x: x.replace(
            b"<pageMargins",
            b'<conditionalFormatting sqref="A1:A3"><cfRule type="expression" priority="1"><formula>[1]Defs!$A$1&gt;0</formula></cfRule>'
            b"</conditionalFormatting><pageMargins", 1)
    if review:
        dv = '<dataValidations count="1"><dataValidation type="list" sqref="F1:F3"><formula1>"가,나"</formula1></dataValidation></dataValidations>'
        hf = "<headerFooter><oddHeader>&amp;L고객사</oddHeader></headerFooter>"
        mods["xl/worksheets/sheet1.xml"] = lambda x: x.replace(
            b"<pageMargins", dv.encode(), 1).replace(b"</worksheet>", hf.encode() + b"</worksheet>")
    if hf_picture:
        prev = mods["xl/worksheets/sheet2.xml"]
        mods["xl/worksheets/sheet2.xml"] = lambda x: prev(x).replace(b"</worksheet>", b'<legacyDrawingHF r:id="rId8"/></worksheet>')
    if ext == "x14":
        mods["xl/worksheets/sheet4.xml"] = lambda x: x.replace(b"</worksheet>", x14.replace(b"Drop!$A$1:$A$3", b"Pri") + b"</worksheet>")

    if ext:
        def wbmod(x: bytes) -> bytes:
            if b"xmlns:r=" not in x:
                x = x.replace(b"<workbook ", b'<workbook xmlns:r="' + REL.rstrip("/").encode() + b'" ', 1)
            names = b'<definedNames><definedName name="Pri">[1]Defs!$A$1:$A$3</definedName>'
            names += b'<definedName name="_xlnm._FilterDatabase" localSheetId="3" hidden="1">Drop!#REF!</definedName>'
            names += b"</definedNames>"
            return x.replace(b"</sheets>", b'</sheets><externalReferences><externalReference r:id="rIdE1"/></externalReferences>' + names, 1)
        mods["xl/workbook.xml"] = wbmod
        mods["xl/_rels/workbook.xml.rels"] = lambda x: x.replace(
            b"</Relationships>",
            f'<Relationship Id="rIdE1" Type="{REL}externalLink" Target="externalLinks/externalLink1.xml"/></Relationships>'.encode())

    def ctmod(x: bytes) -> bytes:
        add = (
            b'<Default Extension="png" ContentType="image/png"/><Default Extension="bin" ContentType="application/octet-stream"/>'
            + b"".join(
                f'<Override PartName="/xl/drawings/drawing{i}.xml" ContentType="{DRAWING_CT}"/>'.encode() for i in (1, 2)
            )
        )
        if ext:
            add += b'<Override PartName="/xl/externalLinks/externalLink1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml"/>'
        return x.replace(b"</Types>", add + b"</Types>")
    mods["[Content_Types].xml"] = ctmod
    mods["docProps/app.xml"] = lambda x: x.replace(b"</Properties>", "<Company>고객사X</Company><Manager>관리자</Manager></Properties>".encode())

    adds = {
        "xl/worksheets/_rels/sheet2.xml.rels": _rels(
            ("rId1", REL + "drawing", "../drawings/drawing1.xml"),
            ("rId2", REL + "printerSettings", "../printerSettings/printerSettings1.bin"),
        ),
        "xl/worksheets/_rels/sheet3.xml.rels": _rels(("rId1", REL + "drawing", "../drawings/drawing2.xml")),  # 시트 XML 에는 <drawing> 이 없다
        "xl/printerSettings/printerSettings1.bin": PRINTER,
        "xl/drawings/drawing1.xml": _drawing("도형글자"),
        "xl/drawings/drawing2.xml": _drawing("공통글자"),
        "xl/drawings/_rels/drawing1.xml.rels": _rels(
            ("rId1", REL + "image", "../media/image1.png"),
            *([("rId2", REL + "chart", "../charts/chart1.xml")] if chart else []),
        ),
        "xl/drawings/_rels/drawing2.xml.rels": _rels(("rId1", REL + "image", "../media/image1.png")),  # 같은 이미지 공유
        "xl/media/image1.png": PNG,
    }
    if review:
        adds["xl/comments1.xml"] = b'<?xml version="1.0"?><comments/>'
    if chart:
        adds["xl/charts/chart1.xml"] = b'<?xml version="1.0"?><c:chartSpace xmlns:c="x"/>'
    if legacy:
        adds["xl/worksheets/_rels/sheet2.xml.rels"] = _rels(
            ("rId1", REL + "drawing", "../drawings/drawing1.xml"), ("rId9", REL + "vmlDrawing", "../drawings/vmlDrawing1.vml"))
        adds["xl/drawings/vmlDrawing1.vml"] = b"<xml/>"
    if ext:
        adds["xl/externalLinks/externalLink1.xml"] = (
            '<?xml version="1.0"?><externalLink xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><externalBook r:id="rId1">'
            '<sheetNames><sheetName val="Defs"/></sheetNames></externalBook></externalLink>'
        ).encode()
        adds["xl/externalLinks/_rels/externalLink1.xml.rels"] = (
            f'<?xml version="1.0"?><Relationships xmlns="{RELS_NS}"><Relationship Id="rId1" Type="{REL}externalLinkPath" '
            'Target="file:///\\\\server\\share\\other.xlsx" TargetMode="External"/></Relationships>'
        ).encode()
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as zi, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename in mods:
                data = mods[it.filename](data)
            zo.writestr(it.filename, data)
        for n, d in adds.items():
            zo.writestr(n, d)
    share_strings(path)


def _map(root: Path, *, clear=("B2:D3",), pics=("Pic1", "Pic2"), top=None, info_extra=None) -> None:
    info = {"이름": "Info", "칸": [{"항목": "코드", "위치": "A1"}]}
    if clear:
        info["비우기"] = list(clear)
    info.update(info_extra or {})
    sheets = [info] + [{"이름": p, "그림_지우기": True} for p in pics] + [{"이름": "Drop"}]
    data = {"양식": "빈양식", "템플릿": "../src.xlsx", "산출": {"폴더": "out", "파일명": "x.xlsx"}, **(top or {}), "시트": sheets}
    (root / "templates" / "forms").mkdir(parents=True, exist_ok=True)
    (root / "templates" / "forms" / "빈양식.yaml").write_text(yaml.safe_dump(data, allow_unicode=True), encoding="utf-8")


def _run(root: Path, capsys, *extra: str, out: str = "templates/빈.xlsx", src: str = "src.xlsx"):
    code = blank_template.main(["--root", str(root), "--form", "빈양식", "--src", src, "--out", out, *extra])
    text = capsys.readouterr().out
    assert text.count("\n") == 1, "stdout 은 JSON 한 줄이어야 한다"
    return code, json.loads(text)


def _parts(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as zf:
        return {n: zf.read(n) for n in zf.namelist()}


def _well_formed(parts: dict[str, bytes]) -> None:
    for n, d in parts.items():
        if n.endswith((".xml", ".rels")):
            ET.fromstring(d)


@pytest.fixture
def root(tmp_path: Path) -> Path:
    (tmp_path / "templates").mkdir()
    _make_src(tmp_path / "src.xlsx")
    _map(tmp_path)
    return tmp_path


def test_값과_그림을_지우고_openpyxl로_읽힌다(root, capsys):
    code, rep = _run(root, capsys)
    assert code == 0, rep
    assert rep["status"] == "ok" and rep["files"] == ["templates/빈.xlsx"]
    out = root / "templates" / "빈.xlsx"
    ws = openpyxl.load_workbook(out)["Info"]
    assert ws["A1"].value == "코드"
    assert ws["B2"].value is None and ws["B3"].value is None and ws["D2"].value is None
    assert ws["C3"].value == "=1+1"  # 수식 셀은 지우지 않는다
    assert any(n["type"] == "formula-kept" and n["cells"] == ["C3"] for n in rep["notices"])
    assert rep["cleared"] == {"Info": 4}  # B2 B3 C2 D2
    _well_formed(_parts(out))


def _sst_list(xml: str) -> list[str]:
    return [blank_template._si_text(m.group(0)) for m in blank_template._SI_RE.finditer(xml)]


def test_공유문자열은_더_이상_안_쓰는_것만_비운다(root, capsys):
    src_xml = _parts(root / "src.xlsx")["xl/sharedStrings.xml"].decode()
    _, rep = _run(root, capsys)
    sst_xml = _parts(root / "templates" / "빈.xlsx")["xl/sharedStrings.xml"].decode()
    before, after = _sst_list(src_xml), _sst_list(sst_xml)
    assert len(after) == len(before)  # 번호 유지
    assert [b for b, a in zip(before, after) if b != a] == ["샘플A", "샘플B"]
    assert "코드" in after  # A1 이 아직 가리킨다
    assert rep["strings_blanked"] == 2
    # 비운 칸 4개(B2 B3 C2 D2) 중 공유 문자열을 가리키던 3개만큼 count 가 줄었다
    count = lambda x: int(re.search(r' count="(\d+)"', x).group(1))
    assert count(sst_xml) == count(src_xml) - 3


def test_손대지_않은_파트는_바이트가_같다(root, capsys):
    _, _ = _run(root, capsys)
    a, b = _parts(root / "src.xlsx"), _parts(root / "templates" / "빈.xlsx")
    changed = {"xl/worksheets/sheet1.xml", "xl/sharedStrings.xml", "docProps/core.xml", "docProps/app.xml"}
    gone = {"xl/drawings/drawing1.xml", "xl/drawings/drawing2.xml", "xl/drawings/_rels/drawing1.xml.rels",
            "xl/drawings/_rels/drawing2.xml.rels", "xl/media/image1.png", "xl/worksheets/_rels/sheet3.xml.rels"}
    for n, d in a.items():
        if n in gone:
            assert n not in b
        elif n in changed or n in ("xl/worksheets/sheet2.xml", "xl/worksheets/_rels/sheet2.xml.rels", "[Content_Types].xml"):
            continue
        else:
            assert b[n] == d, n
    assert b["xl/worksheets/sheet4.xml"] == a["xl/worksheets/sheet4.xml"]  # x14 드롭다운 시트
    assert b"x14:dataValidation" in b["xl/worksheets/sheet4.xml"]


def test_그림_지우기는_drawing_관계_media_ContentTypes를_정리하고_printerSettings는_남긴다(root, capsys):
    _, rep = _run(root, capsys)
    b = _parts(root / "templates" / "빈.xlsx")
    assert b"<drawing" not in b["xl/worksheets/sheet2.xml"]
    rels = b["xl/worksheets/_rels/sheet2.xml.rels"].decode()
    assert "printerSettings1.bin" in rels and "drawing1.xml" not in rels
    assert b["xl/printerSettings/printerSettings1.bin"] == PRINTER
    assert "xl/worksheets/_rels/sheet3.xml.rels" not in b  # 관계만 있던 시트(<drawing> 없음)
    ct = b["[Content_Types].xml"].decode()
    assert "/xl/drawings/" not in ct
    assert 'Extension="png"' in ct  # Default 는 남긴다
    assert not any(n.startswith("xl/media/") or n.startswith("xl/drawings/") for n in b)
    assert "xl/media/image1.png" in rep["removed_parts"]


def test_다른_시트가_쓰는_drawing과_공유_이미지는_남긴다(tmp_path, capsys):
    (tmp_path / "templates").mkdir()
    _make_src(tmp_path / "src.xlsx")
    _map(tmp_path, pics=("Pic1",), clear=())
    code, rep = _run(tmp_path, capsys)
    b = _parts(tmp_path / "templates" / "빈.xlsx")
    assert code == 0, rep
    assert "xl/drawings/drawing1.xml" not in b and "xl/drawings/drawing2.xml" in b
    assert "xl/media/image1.png" in b  # drawing2 가 아직 쓴다
    assert "/xl/drawings/drawing2.xml" in b["[Content_Types].xml"].decode()
    assert "/xl/drawings/drawing1.xml" not in b["[Content_Types].xml"].decode()
    assert openpyxl.load_workbook(tmp_path / "templates" / "빈.xlsx")["Pic1"]["A1"].value == "그림시트1"


def test_지원하지_않는_drawing은_종료코드_2(tmp_path, capsys):
    (tmp_path / "templates").mkdir()
    _make_src(tmp_path / "src.xlsx", chart=True)
    _map(tmp_path, clear=())
    code, rep = _run(tmp_path, capsys)
    assert code == 2 and rep["status"] == "error"
    assert "chart" in rep["problems"][0]["reason"]
    assert not (tmp_path / "templates" / "빈.xlsx").exists()


def test_legacyDrawing이_있는_시트는_종료코드_2(tmp_path, capsys):
    (tmp_path / "templates").mkdir()
    _make_src(tmp_path / "src.xlsx", legacy=True)
    _map(tmp_path, clear=())
    code, rep = _run(tmp_path, capsys)
    assert code == 2 and "legacyDrawing" in rep["problems"][0]["reason"]


def test_out이_이미_있으면_거부_1(root, capsys):
    (root / "templates" / "빈.xlsx").write_bytes(b"keep")
    code, rep = _run(root, capsys)
    assert code == 1 and rep["status"] == "rejected"
    assert (root / "templates" / "빈.xlsx").read_bytes() == b"keep"


def test_범위_모양_오류는_거부_1(root, capsys):
    _map(root, clear=("ZZZ",))
    code, rep = _run(root, capsys)
    assert code == 1 and "ZZZ" in rep["problems"][0]["reason"]
    assert not (root / "templates" / "빈.xlsx").exists()


def test_읽지_못한_견본은_2(root, capsys):
    (root / "bad.xlsx").write_bytes(b"not a zip")
    code, rep = _run(root, capsys, src="bad.xlsx")
    assert code == 2 and rep["status"] == "error"


def test_비우기가_지도_칸과_겹치면_지도_오류_2(root, capsys):
    _map(root, clear=("A1:B2",))
    code, rep = _run(root, capsys)
    assert code == 2 and "겹친다" in rep["problems"][0]["reason"]


def test_비우기가_표와_서술과_유지와_겹쳐도_지도_오류():
    base = {"양식": "x", "템플릿": "t", "산출": {}}
    sheet = {"이름": "S", "표": [{"이름": "T", "머리글_행": 1, "시작_행": 2, "최대_행": 3, "열": {"a": "C"}}],
             "서술": [{"이름": "N", "위치": "E5:E9"}], "유지": ["G1"]}
    ok = {**base, "시트": [{**sheet, "비우기": ["C5:C9", "D1:D50", "F1:F9"]}]}  # C 열 5행 이후는 표(2~4행) 밖
    assert blank_template.map_problems(ok) == []
    for rng in ("C3", "E6:E7", "G1:H2"):
        bad = {**base, "시트": [{**sheet, "비우기": [rng]}]}
        assert any("겹친다" in m for _, m in blank_template.map_problems(bad)), rng


def test_form_columns는_새_키를_지도_오류로_보지_않고_겹침은_검사_불능(tmp_path):
    base = {"시트": [{"이름": "S", "칸": [{"항목": "a", "위치": "A1"}], "비우기": ["C1:C2"], "그림_지우기": True, "유지": ["E1"]}]}
    assert form_columns._valid_map(base, Path("m.yaml")) is base
    base["시트"][0]["비우기"] = ["A1"]
    with pytest.raises(ConfigError, match="겹친다"):
        form_columns._valid_map(base, Path("m.yaml"))


def test_check는_파일을_쓰지_않는다(root, capsys):
    code, rep = _run(root, capsys, "--check")
    assert code == 0 and rep["dry_run"] is True and rep["files"] == []
    assert not (root / "templates" / "빈.xlsx").exists()


def test_지운_견본_글자가_남은_drawing에_있으면_3(tmp_path, capsys):
    (tmp_path / "templates").mkdir()
    _make_src(tmp_path / "src.xlsx")
    _map(tmp_path, clear=("E5",), pics=("Pic1",))  # Pic2 의 drawing2 에 '공통글자' 가 남는다
    code, rep = _run(tmp_path, capsys)
    assert code == 3 and rep["status"] == "warning"
    assert {"part": "xl/drawings/drawing2.xml", "text": "공통글자"} in rep["leftovers"]
    assert (tmp_path / "templates" / "빈.xlsx").exists()  # 만들었으나 남은 견본이 있다


# ---------------------------------------------------------------- 외부 참조

def _ext_root(tmp_path: Path, ext: str, wipe: bool = True) -> Path:
    (tmp_path / "templates").mkdir()
    _make_src(tmp_path / "src.xlsx", ext=ext)
    _map(tmp_path, clear=(), pics=(), top={"외부_참조_지우기": True} if wipe else {})
    return tmp_path


def test_외부_참조를_쓰는_곳이_없으면_지운다(tmp_path, capsys):
    root = _ext_root(tmp_path, "free")
    code, rep = _run(root, capsys)
    assert code == 0, rep
    b = _parts(root / "templates" / "빈.xlsx")
    assert not any(n.startswith("xl/externalLinks/") for n in b)
    wb = b["xl/workbook.xml"].decode()
    assert "externalReferences" not in wb and 'name="Pri"' not in wb
    assert "_xlnm._FilterDatabase" in wb  # #REF! 이름은 외부 참조가 아니라 그대로
    assert "externalLink" not in b["xl/_rels/workbook.xml.rels"].decode()
    assert "externalLink" not in b["[Content_Types].xml"].decode()
    assert rep["external_links"] == {"removed": 1, "targets": 1, "names_removed": ["Pri"]}
    assert any(n["type"] == "broken-names" for n in rep["notices"])
    _well_formed(b)
    openpyxl.load_workbook(root / "templates" / "빈.xlsx")


@pytest.mark.parametrize("ext,what", [("dv", "B2:B3"), ("formula", "C3"), ("x14", "B1:B5"), ("cf", "A1:A3")])
def test_외부_이름을_쓰는_드롭다운이나_수식이_있으면_2(tmp_path, capsys, ext, what):
    root = _ext_root(tmp_path, ext)
    code, rep = _run(root, capsys)
    assert code == 2 and rep["status"] == "error"
    reason = rep["problems"][0]["reason"]
    assert what in reason
    assert not (root / "templates" / "빈.xlsx").exists()


def test_키가_없으면_외부_참조를_건드리지_않고_알린다(tmp_path, capsys):
    root = _ext_root(tmp_path, "free", wipe=False)
    code, rep = _run(root, capsys)
    assert code == 0
    b = _parts(root / "templates" / "빈.xlsx")
    assert "xl/externalLinks/externalLink1.xml" in b
    note = next(n for n in rep["notices"] if n["type"] == "external-links")
    assert note["links"] == 1 and note["targets"] == 1


def test_legacyDrawingHF가_있는_그림_지우기_시트는_2(tmp_path, capsys):
    (tmp_path / "templates").mkdir()
    _make_src(tmp_path / "src.xlsx", hf_picture=True)
    _map(tmp_path, clear=())
    code, rep = _run(tmp_path, capsys)
    assert code == 2 and "legacyDrawingHF" in rep["problems"][0]["reason"]


def test_문서_속성의_출처_정보를_비운다(root, capsys):
    code, rep = _run(root, capsys)
    assert code == 0
    b = _parts(root / "templates" / "빈.xlsx")
    core, app = b["docProps/core.xml"].decode(), b["docProps/app.xml"].decode()
    for gone in ("고객 견본", "홍길동", "김철수"):
        assert gone not in core
    assert "고객사X" not in app and "관리자" not in app
    assert sorted(rep["doc_props_cleared"]) == ["Company", "Manager", "cp:lastModifiedBy", "dc:creator", "dc:title"]
    _well_formed(b)


def test_견본_판정을_안_하는_곳은_확인_필요로_알린다(tmp_path, capsys):
    (tmp_path / "templates").mkdir()
    _make_src(tmp_path / "src.xlsx", review=True)
    _map(tmp_path)
    code, rep = _run(tmp_path, capsys)
    assert code == 0, rep
    kinds = {n["kind"] for n in rep["notices"] if n["type"] == "review-needed"}
    assert {"header-footer", "validation-literal", "comments", "sheet-names"} <= kinds
    assert all("확인 필요" in n["reason"] for n in rep["notices"] if n["type"] == "review-needed")


def test_family_같은_태그를_수식으로_오인하지_않는다():
    from checker import export
    assert not export.is_formula('<v>1</v><family val="2"/>')
    assert export.is_formula("<f>1+1</f><v></v>") and export.is_formula('<f t="shared" si="0"/>')
