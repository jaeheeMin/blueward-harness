"""작업본(md) → 고객 양식 pptx 산출(`checker.export_pptx`)을 검증한다(이슈 #561).

고객 양식 파일은 저장소에 넣지 않는다(다른 회사 화면이 들어 있다). 시험용 양식은 python-pptx 로 코드에서 만든다.
양식 글자는 일부러 뜻 없는 견본 문구(`견본_…`)로 쓴다 — 엔진에 특정 양식의 글자가 없다는 것도 함께 본다.
"""
from __future__ import annotations

import io
import json
import re
import zipfile
from pathlib import Path

import pytest
import yaml
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from checker import export, export_pptx

MAP = """\
양식: 시험매뉴얼
템플릿: ../시험.pptx
검사_규칙: ../../rules/test.yaml
산출:
  폴더: docs/out
  파일명: "M_{LoB}_{제목}_v{버전}.pptx"
  자리:
    버전: {메타: 버전}
쪽:
  - 이름: 표지
    견본: 1
    칸:
      - {항목: LoB, 찾기: "견본_표지LoB", 틀: "매뉴얼_{값}"}
      - {항목: 제목, 찾기: "견본_표지제목"}
    지우기: ["견본_표지안내"]
  - 이름: 목차
    견본: 2
    칸:
      - {항목: LoB, 찾기: "견본_목차LoB", 틀: "목차_{값}"}
    목록: {글: "견본_목차글", 번호: "견본_목차번호", 제목: 화면 제목}
  - 이름: 본문
    견본: 3
    반복: 화면
    칸:
      - {항목: 업무명, 찾기: "견본_업무명"}
      - {항목: 화면 제목, 찾기: "견본_화면제목"}
      - {항목: 기능 설명, 찾기: "견본_기능설명"}
    순번: {찾기: "견본_순번"}
    그림: {자리: 1}
    번호_설명: {번호: {글자: "견본_번호", 왼쪽_이상: 7000000}, 글: "견본_설명글", 쪽당_최대: 3}
    유지: [표]
  - 이름: 끝
    견본: 4
"""

RULES = """\
관할: "docs/out/**"
규칙:
  - 종류: filename
    패턴: '^M_.+_v\\d+\\.\\d+\\.pptx$'
"""

SPEC = """\
# DEV-001 시험

버전: 0.1

## 양식 항목

### 시험매뉴얼 · 표지

#### LoB

Production(생산)

#### 제목

생산계획 보드

### 시험매뉴얼 · 목차

해당 없음

### 시험매뉴얼 · 본문 · 첫째 화면

| 항목 | 값 |
|---|---|
| 업무명 | 계획 관리 |
| 화면 제목 | 첫째 제목 |
| 기능 설명 | 첫째 화면 설명입니다. |

#### 그림

docs/img/a.png

#### 번호 설명

1. 하나
2. 둘

### 시험매뉴얼 · 본문 · 둘째 화면

| 항목 | 값 |
|---|---|
| 업무명 | 계획 관리 |
| 화면 제목 | 둘째 제목 |
| 기능 설명 | 둘째 화면 설명입니다. |

#### 그림

docs/img/b.png

#### 번호 설명

1. 메뉴 경로를 따라 들어간다.
   다음 줄이다.
2. 둘
3. 셋
4. 넷
5. 다섯
6. 여섯
7. 일곱
"""

SAMPLES = [
    "견본_표지LoB", "견본_표지제목", "견본_표지안내", "견본_표지날짜", "견본_목차LoB", "견본_목차글", "견본_목차번호",
    "견본_업무명", "견본_화면제목", "견본_기능설명", "견본_순번", "견본_번호", "견본_설명글", "견본_엉뚱한도형",
]


def _png(color=(10, 120, 200), size=(160, 100)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, "PNG")
    return buf.getvalue()


def _tb(slide, left, top, w, h, text, pt=12):
    s = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(w), Emu(h))
    s.text_frame.text = text
    for p in s.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(pt)
    s.text_frame.word_wrap = True
    return s


def _make_template(path: Path, extra_slide: bool = False) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    blank = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank)
    _tb(s1, 100000, 100000, 5000000, 400000, "견본_표지LoB 한글")
    _tb(s1, 100000, 600000, 5000000, 400000, "견본_표지제목 xxx")
    _tb(s1, 100000, 1100000, 5000000, 400000, "견본_표지날짜 2025")
    _tb(s1, 100000, 1600000, 5000000, 400000, "견본_표지안내 지울 것")
    s2 = prs.slides.add_slide(blank)
    _tb(s2, 100000, 100000, 5000000, 400000, "견본_목차LoB")
    _tb(s2, 100000, 700000, 5000000, 2000000, "견본_목차글\n견본_목차글 둘\n견본_목차글 셋")
    _tb(s2, 100000, 2800000, 300000, 2000000, "견본_목차번호\n2\n3")
    s3 = prs.slides.add_slide(blank)
    _tb(s3, 100000, 100000, 2000000, 300000, "견본_업무명")
    _tb(s3, 100000, 500000, 4000000, 400000, "견본_화면제목")
    _tb(s3, 100000, 1000000, 800000, 400000, "견본_순번")
    s3.shapes.add_picture(io.BytesIO(_png((250, 0, 0))), Emu(500000), Emu(1800000), Emu(6000000), Emu(4000000))
    s3.shapes.add_table(2, 2, Emu(300000), Emu(1100000), Emu(11000000), Emu(300000))
    _tb(s3, 7500000, 1700000, 4000000, 800000, "견본_기능설명 한두 문장")
    _tb(s3, 7300000, 2900000, 280000, 260000, "견본_번호")
    _tb(s3, 100000, 4000000, 280000, 260000, "견본_번호")  # 왼쪽_이상 으로 가려져야 하는 다른 번호
    _tb(s3, 7600000, 2900000, 4200000, 340000, "견본_설명글 한 줄")
    _tb(s3, 3000000, 3000000, 500000, 300000, "견본_엉뚱한도형")  # 지도가 안 가리키므로 지워져야 한다
    page = _tb(s3, 11000000, 6500000, 500000, 300000, "7")
    run = page.text_frame.paragraphs[0].runs[0]._r
    fld = etree.SubElement(run.getparent(), qn("a:fld"))
    fld.set("id", "{00000000-1234-1234-1234-123412341234}")
    fld.set("type", "slidenum")
    t = etree.SubElement(fld, qn("a:t"))
    t.text = "7"
    run.getparent().remove(run)  # 쪽 번호 도형은 run 없이 a:fld 다
    s4 = prs.slides.add_slide(blank)
    _tb(s4, 100000, 100000, 5000000, 400000, "끝 인사")
    if extra_slide:
        s5 = prs.slides.add_slide(blank)
        _tb(s5, 100000, 100000, 5000000, 400000, "견본_안쓰는쪽")
    prs.save(str(path))


@pytest.fixture
def proj(tmp_path: Path) -> Path:
    (tmp_path / "templates" / "forms").mkdir(parents=True)
    (tmp_path / "rules").mkdir()
    (tmp_path / "docs" / "img").mkdir(parents=True)
    (tmp_path / "docs" / "spec").mkdir(parents=True)
    (tmp_path / "templates" / "forms" / "시험매뉴얼.yaml").write_text(MAP, encoding="utf-8")
    (tmp_path / "rules" / "test.yaml").write_text(RULES, encoding="utf-8")
    _make_template(tmp_path / "templates" / "시험.pptx")
    (tmp_path / "docs" / "img" / "a.png").write_bytes(_png((0, 150, 0), (320, 200)))
    (tmp_path / "docs" / "img" / "b.png").write_bytes(_png((150, 0, 150), (100, 200)))
    return tmp_path


def _write_map(proj: Path, text: str) -> None:
    (proj / "templates" / "forms" / "시험매뉴얼.yaml").write_text(text, encoding="utf-8")


def _run(proj: Path, capsys, spec: str = SPEC, *extra: str):
    p = proj / "docs" / "spec" / "DEV-001.md"
    p.write_text(spec, encoding="utf-8")
    code = export.main(["--spec", str(p), "--root", str(proj), *extra])
    out = capsys.readouterr().out
    assert out.count("\n") == 1, "stdout 은 JSON 한 줄이어야 한다"
    return code, json.loads(out)


def _out(proj: Path, report: dict) -> Path:
    assert len(report["files"]) == 1, report
    return proj / report["files"][0]


def _texts(slide) -> list[str]:
    return [export.norm(s.text_frame.text) for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]


def _produce(proj, capsys):
    code, rep = _run(proj, capsys)
    assert code == 3 and rep["status"] == "warning", rep
    return Presentation(str(_out(proj, rep))), rep


def test_산출하면_표지_목차_본문_끝_순서로_쪽이_놓인다(proj, capsys):
    prs, rep = _produce(proj, capsys)
    slides = list(prs.slides)
    assert [_texts(s)[0] for s in slides][:2] == ["매뉴얼_Production(생산)", "목차_Production(생산)"]
    assert _texts(slides[-1]) == ["끝 인사"]
    # 본문: 첫째 1쪽 + 둘째(번호 7개, 쪽당 최대 3) 3쪽
    titles = [t for s in slides[2:-1] for t in _texts(s) if t.endswith("제목")]
    assert titles == ["첫째 제목", "둘째 제목", "둘째 제목", "둘째 제목"]
    assert len(slides) == 2 + 4 + 1
    assert rep["files"][0].endswith("M_Production(생산)_생산계획 보드_v0.1.pptx")


def test_칸_치환_틀_지우기(proj, capsys):
    prs, _ = _produce(proj, capsys)
    cover = _texts(list(prs.slides)[0])
    assert "매뉴얼_Production(생산)" in cover and "생산계획 보드" in cover
    assert not any("견본_표지안내" in t for t in cover)  # 지우기
    assert any("견본_표지날짜" in t for t in cover), "지도가 안 가리킨 글은 그대로 남는다(지도의 책임)"


def test_글꼴_서식은_첫_문단_그대로(proj, capsys):
    prs, _ = _produce(proj, capsys)
    shape = [s for s in list(prs.slides)[0].shapes if s.has_text_frame and "생산계획 보드" in s.text_frame.text][0]
    assert shape.text_frame.paragraphs[0].runs[0].font.size == Pt(12)


def test_목차는_본문_쪽_제목으로_채운다(proj, capsys):
    prs, _ = _produce(proj, capsys)
    toc = list(prs.slides)[1]
    box = [s for s in toc.shapes if s.has_text_frame and "첫째 제목" in s.text_frame.text][0]
    nums = [s for s in toc.shapes if s.has_text_frame and s.text_frame.text.split("\n")[0] == "1"][0]
    assert [p.text for p in box.text_frame.paragraphs] == ["첫째 제목", "둘째 제목"]  # 이어진 쪽은 한 줄
    assert [p.text for p in nums.text_frame.paragraphs] == ["1", "2"]


def test_본문_쪽_순번은_두_자리이고_이어진_쪽은_같은_번호(proj, capsys):
    prs, _ = _produce(proj, capsys)
    body = list(prs.slides)[2:-1]
    seqs = [[t for t in _texts(s) if re.fullmatch(r"0\d", t)][0] for s in body]
    assert seqs == ["01", "02", "02", "02"]


def test_번호_설명은_쪽당_최대에서_넘기고_번호는_이어진다(proj, capsys):
    prs, rep = _produce(proj, capsys)
    body = list(prs.slides)[3:-1]
    nums = []
    for s in body:
        n = [export.norm(x.text_frame.text) for x in sorted(s.shapes, key=lambda x: x.top)
             if x.has_text_frame and x.left == 7300000]  # 번호 도형(쪽 번호 도형은 제외)
        nums.append(n)
    assert nums == [["1", "2", "3"], ["4", "5", "6"], ["7"]]
    first = [t for t in _texts(body[0]) if t.startswith("메뉴 경로")][0]
    assert "다음 줄이다." in first  # 이어지는 들여쓴 줄은 같은 항목
    assert any(w["type"] == "pages-split" for w in rep["warnings"])


def test_번호_글_상자는_위에서_아래로_겹치지_않게_쌓인다(proj, capsys):
    prs, _ = _produce(proj, capsys)
    body = list(prs.slides)[3]
    notes = sorted((s for s in body.shapes if s.has_text_frame and s.left == 7600000), key=lambda s: s.top)
    assert len(notes) == 3
    for a, b in zip(notes, notes[1:]):
        assert a.top + a.height <= b.top
    first_two_lines = notes[0].height > notes[1].height  # 두 줄짜리가 더 높다(추정)
    assert first_two_lines


def test_쪽_번호_도형은_유지되고_지도가_안_가리킨_도형은_지워진다(proj, capsys):
    prs, _ = _produce(proj, capsys)
    for s in list(prs.slides)[2:-1]:
        assert any(sh._element.find(".//" + qn("a:fld")) is not None for sh in s.shapes)
        assert not any("견본_엉뚱한도형" in t for t in _texts(s))
        assert any(sh._element.find(".//" + qn("a:tbl")) is not None for sh in s.shapes)  # 유지: [표]


def test_예시_글자와_예시_그림이_산출물에_남지_않는다(proj, capsys):
    _, rep = _produce(proj, capsys)
    path = _out(proj, rep)
    with zipfile.ZipFile(path) as z:
        xml = "".join(z.read(n).decode("utf-8", "replace") for n in z.namelist() if n.startswith("ppt/slides/slide"))
        media = [n for n in z.namelist() if n.startswith("ppt/media/")]
        slide_files = [n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)]
    for s in SAMPLES:
        if s == "견본_표지날짜":
            continue  # 지도가 가리키지 않은 글(위 시험)
        assert s not in xml, s
    assert len(slide_files) == 7  # 고아 슬라이드 없음
    # 예시 그림(빨강)은 없고 산출에 쓴 그림 둘만 있다
    colors = {Image.open(io.BytesIO(zipfile.ZipFile(path).read(n))).convert("RGB").getpixel((1, 1)) for n in media}
    assert (250, 0, 0) not in colors and {(0, 150, 0), (150, 0, 150)} <= colors


def test_그림은_자리에_비율을_유지해_맞춘다(proj, capsys):
    prs, _ = _produce(proj, capsys)
    p1 = [s for s in list(prs.slides)[2].shapes if s.shape_type == 13][0]
    assert abs(p1.width / p1.height - 320 / 200) < 0.01  # a.png 비율
    assert p1.left >= 500000 and p1.top >= 1800000
    assert p1.left + p1.width <= 6500000 and p1.top + p1.height <= 5800000
    p2 = [s for s in list(prs.slides)[3].shapes if s.shape_type == 13][0]
    assert abs(p2.width / p2.height - 0.5) < 0.01  # b.png 는 세로로 긴 그림


def test_그림_파일이_없으면_거부하고_파일을_쓰지_않는다(proj, capsys):
    (proj / "docs" / "img" / "a.png").unlink()
    code, rep = _run(proj, capsys)
    assert code == 1 and rep["status"] == "rejected" and rep["files"] == []
    assert any(p["item"] == "그림" and "본문 · 첫째 화면" in p["sheet"] for p in rep["problems"])
    assert not (proj / "docs" / "out").exists()


def test_그림_경로가_없으면_거부(proj, capsys):
    code, rep = _run(proj, capsys, SPEC.replace("#### 그림\n\ndocs/img/a.png\n\n", ""))
    assert code == 1 and any(p["item"] == "그림" for p in rep["problems"])


def test_찾기가_0개면_항목을_알리고_거부(proj, capsys):
    _write_map(proj, MAP.replace('찾기: "견본_업무명"', '찾기: "바뀐 글자"'))
    code, rep = _run(proj, capsys)
    assert code == 1 and rep["files"] == []
    hit = [p for p in rep["problems"] if p["item"] == "업무명"]
    assert hit and "바뀐 글자" in hit[0]["reason"] and "0개" in hit[0]["reason"] and hit[0]["sheet"] == "본문"


def test_찾기가_2개_이상이면_항목을_알리고_거부(proj, capsys):
    _write_map(proj, MAP.replace('찾기: "견본_업무명"', '찾기: "견본_"'))
    code, rep = _run(proj, capsys)
    assert code == 1
    hit = [p for p in rep["problems"] if p["item"] == "업무명"]
    assert hit and "개라 하나로 정할 수 없다" in hit[0]["reason"]


def test_같은_도형을_두_항목이_가리키면_거부(proj, capsys):
    _write_map(proj, MAP.replace('찾기: "견본_화면제목"', '찾기: "견본_업무명"'))
    code, rep = _run(proj, capsys)
    assert code == 1 and any("두 항목이 가리킬 수 없다" in p["reason"] for p in rep["problems"])


def test_견본_번호가_범위_밖이면_거부(proj, capsys):
    _write_map(proj, MAP.replace("견본: 4", "견본: 9"))
    code, rep = _run(proj, capsys)
    assert code == 1 and any(p["item"] == "견본" and "9" in p["reason"] for p in rep["problems"])


def test_지원하지_않는_키는_읽지_못함(proj, capsys):
    _write_map(proj, MAP.replace("    반복: 화면\n", "    반복: 화면\n    좌표: 1\n"))
    code, rep = _run(proj, capsys)
    assert code == 2 and rep["problems"][0]["item"] == "좌표"


def test_작업본이_지도에_없는_소제목을_쓰면_읽지_못함(proj, capsys):
    code, rep = _run(proj, capsys, SPEC.replace("#### 그림\n\ndocs/img/b.png", "#### 낙서\n\nx\n\n#### 그림\n\ndocs/img/b.png"))
    assert code == 2 and "낙서" in rep["problems"][0]["reason"]


def test_값이_없으면_항목을_알리고_거부(proj, capsys):
    code, rep = _run(proj, capsys, SPEC.replace("| 기능 설명 | 첫째 화면 설명입니다. |", "| 기능 설명 | 미정 |"))
    assert code == 1 and any(p["item"] == "기능 설명" and "첫째 화면" in p["sheet"] for p in rep["problems"])


def test_번호_설명_모양이_틀리면_거부(proj, capsys):
    code, rep = _run(proj, capsys, SPEC.replace("1. 하나\n2. 둘", "그냥 글"))
    assert code == 1 and any(p["item"] == "번호 설명" for p in rep["problems"])


def test_화면을_하나도_안_쓰면_거부(proj, capsys):
    spec = SPEC.split("### 시험매뉴얼 · 본문")[0]
    code, rep = _run(proj, capsys, spec)
    assert code == 1 and any(p["sheet"] == "본문" for p in rep["problems"])


def test_화면_확인_필요는_검사_불능으로_남고_통과가_아니다(proj, capsys):
    code, rep = _run(proj, capsys)
    assert code == 3 and rep["status"] == "warning"
    w = [x for x in rep["warnings"] if x["type"] == "layout-unchecked"]
    assert w and "검사 불능" in w[0]["reason"] and "통과 아님" in w[0]["reason"]
    assert "layout-unchecked" in rep["next"]


def test_check_는_파일을_쓰지_않는다(proj, capsys):
    code, rep = _run(proj, capsys, SPEC, "--check")
    assert code == 3 and rep["dry_run"] is True and not (proj / "docs" / "out").exists()


def test_이미_있는_판은_덮어쓰지_않는다(proj, capsys):
    assert _run(proj, capsys)[0] == 3
    code, rep = _run(proj, capsys)
    assert code == 1 and any("덮어쓰지 않는다" in p["reason"] for p in rep["problems"])


def test_지도가_안_가리키는_견본_쪽은_빼고_알린다(proj, capsys):
    _make_template(proj / "templates" / "시험.pptx", extra_slide=True)
    prs, rep = _produce(proj, capsys)
    assert not any("견본_안쓰는쪽" in t for s in prs.slides for t in _texts(s))
    assert any(w["type"] == "slide-dropped" for w in rep["warnings"])


def test_산출물이_python_pptx로_다시_열리고_xml이_깨지지_않는다(proj, capsys):
    _, rep = _produce(proj, capsys)
    path = _out(proj, rep)
    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None
        for n in z.namelist():
            if n.endswith((".xml", ".rels")):
                etree.fromstring(z.read(n))
        rels = {n: z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("ppt/slides/_rels/")}
        names = set(z.namelist())
    for n, text in rels.items():  # 끊긴 관계 없음
        for t in re.findall(r'Target="\.\./media/([^"]+)"', text):
            assert f"ppt/media/{t}" in names
    Presentation(str(path))


def test_도형_목록_명령은_찾기가_유일한지_보이고_쓰지_않는다(proj, capsys):
    tpl = proj / "templates" / "시험.pptx"
    before = tpl.read_bytes()
    assert export_pptx.main(["--dump", str(tpl)]) == 0
    out = capsys.readouterr().out
    assert "슬라이드 1" in out and "견본_표지LoB" in out and "유일" in out
    assert "중복 2개" in out  # 왼쪽 위치가 다른 같은 글자의 번호 도형 둘
    assert "그림 1번째" in out and "자동 쪽 번호" in out
    assert tpl.read_bytes() == before and not (proj / "docs" / "out").exists()


def test_도형_목록_명령은_못_읽으면_검사_불능(tmp_path, capsys):
    bad = tmp_path / "x.pptx"
    bad.write_bytes(b"not a pptx")
    assert export_pptx.main(["--dump", str(bad)]) == 2
    assert "검사 불능" in capsys.readouterr().out


def test_xlsx_지도는_그대로_시트를_요구한다(proj, capsys):
    _write_map(proj, "양식: 시험매뉴얼\n템플릿: ../시험.xlsx\n산출: {폴더: docs/out, 파일명: x.xlsx}\n쪽: []\n")
    code, rep = _run(proj, capsys)
    assert code == 2 and "`템플릿`·`시트` 가 없다" in rep["problems"][0]["reason"]


def test_xlsx_산출은_pptx_분기를_타지_않는다():
    fmap = {"템플릿": "../a.xlsx", "시트": []}
    assert not export_pptx.is_pptx_map(fmap) and export_pptx.is_pptx_map({"템플릿": "../a.PPTX"})
    assert export.is_srs_map({"쪽": []}) is False  # 시트 없는 지도에서도 죽지 않는다


def test_엔진_코드에_특정_양식의_글자가_없다():
    root = Path(export.__file__).parent
    code = (root / "export_pptx.py").read_text(encoding="utf-8")
    for word in ("사용자 매뉴얼", "사용자매뉴얼", "GL계정", "계정과목", "재무", "Finance", "LoB", "Google Shape"):
        assert word not in code, word
    # 슬라이드 번호·좌표 같은 고정 숫자는 엔진에 없다(지도에서 읽는다)
    assert not re.search(r"\b(?:1793298|7000000|2932580)\b", code)


def _form_columns(proj: Path, text: str):
    from checker.rules import form_columns

    return form_columns._check_spec(text, {"_rules_dir": str(proj / "rules")})


def test_form_columns_맞는_pptx_작업본은_통과한다(proj):
    assert _form_columns(proj, SPEC) == []


def test_form_columns_지도에_없는_쪽은_위반이다(proj):
    vs = _form_columns(proj, SPEC.replace("### 시험매뉴얼 · 목차", "### 시험매뉴얼 · 없는쪽"))
    assert len(vs) == 1 and "양식 밖 항목" in vs[0].message and "없는쪽" in vs[0].message


def test_form_columns_지도에_없는_소제목은_위반이다(proj):
    vs = _form_columns(proj, SPEC.replace("#### 그림\n\ndocs/img/b.png", "#### 낙서\n\nx\n\n#### 그림\n\ndocs/img/b.png"))
    assert len(vs) == 1 and "낙서" in vs[0].message and "양식 밖 항목" in vs[0].message


def test_form_columns_반복_쪽을_소절_이름_없이_쓰면_위반이다(proj):
    vs = _form_columns(proj, SPEC.replace("### 시험매뉴얼 · 본문 · 첫째 화면", "### 시험매뉴얼 · 본문"))
    assert len(vs) == 1 and "반복 쪽" in vs[0].message


def test_form_columns_pptx_지도를_읽지_못하면_검사_불능(proj):
    from checker.model import ConfigError

    _write_map(proj, MAP.replace("    반복: 화면\n", "    반복: 화면\n    좌표: 1\n"))
    with pytest.raises(ConfigError, match="좌표"):
        _form_columns(proj, SPEC)


def _edit_template(proj: Path, fn) -> None:
    path = proj / "templates" / "시험.pptx"
    prs = Presentation(str(path))
    fn(prs)
    prs.save(str(path))


def _shape(slide, text: str):
    return [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.startswith(text)][0]


def test_두_쪽이_같은_견본을_쓰면_거부(proj, capsys):
    _write_map(proj, MAP.replace("견본: 4", "견본: 1"))
    code, rep = _run(proj, capsys)
    assert code == 2 and rep["problems"][0]["item"] == "견본" and "같은 견본 1" in rep["problems"][0]["reason"]
    _write_map(proj, MAP.replace("견본: 4", "견본: 3"))  # 반복 쪽과도 겹치면 안 된다
    code, rep = _run(proj, capsys)
    assert code == 2 and "같은 견본 3" in rep["problems"][0]["reason"]


def test_채우는_중_찾기가_어긋나면_조용히_넘기지_않고_거부(proj, capsys, monkeypatch):
    real = export_pptx._Locator
    calls = {"n": 0}

    class Flaky(real):
        def __init__(self, *a, **k):
            super().__init__(*a, **k)
            calls["n"] += 1
            if calls["n"] > 4:  # 견본 확인(쪽 4개) 뒤 채우기 단계의 첫 호출부터 어긋난다
                self.problems.append(("업무명", "이 쪽에서 도형을 찾지 못했다"))

    monkeypatch.setattr(export_pptx, "_Locator", Flaky)
    code, rep = _run(proj, capsys)
    assert code == 1 and rep["files"] == []
    assert any(p["item"] == "업무명" and "찾지 못했다" in p["reason"] for p in rep["problems"])


def test_layout_unchecked_만_있으면_드롭다운_안내_없이_사람이_확인하라고_안내(proj, capsys):
    code, rep = _run(proj, capsys)
    assert code == 3
    assert rep["next"].startswith("다음: 산출은 됐다.") and "PowerPoint 로 열어" in rep["next"] and "검사 불능" in rep["next"]
    assert "드롭다운" not in rep["next"] and "고쳐 다시 산출" not in rep["next"]


def test_check_도_layout_unchecked_만이면_고치라고_하지_않는다(proj, capsys):
    code, rep = _run(proj, capsys, SPEC, "--check")
    assert code == 3 and rep["dry_run"] is True
    assert "드롭다운" not in rep["next"] and "고친 뒤 다시 확인" not in rep["next"] and "사람이 눈으로 확인" in rep["next"]


def test_xlsx_경고_문구는_그대로다():
    note = export._warning_note(True, True, True, True)
    assert "`type: dropdown` 경고의 값이 템플릿 드롭다운 목록 밖이다" in note and "layout" not in note
    assert "layout-unchecked" in export._warning_note(False, False, layout=True)


def test_서식이_다른_run_이_여럿이면_합쳐진다고_항목_이름으로_경고(proj, capsys):
    def two_runs(prs):
        tf = _shape(list(prs.slides)[2], "견본_업무명").text_frame
        run = tf.paragraphs[0].add_run()
        run.text = " 둘째"
        run.font.bold = True

    _edit_template(proj, two_runs)
    code, rep = _run(proj, capsys)
    assert code == 3
    hit = [w for w in rep["warnings"] if w["type"] == "format-merged"]
    assert len(hit) == 1 and hit[0]["item"] == "업무명" and hit[0]["sheet"] == "본문"  # 이어진 쪽에서 반복되지 않는다


def test_서식이_하나뿐이면_합침_경고가_없다(proj, capsys):
    _, rep = _produce(proj, capsys)
    assert not [w for w in rep["warnings"] if w["type"] == "format-merged"]


def test_왼쪽_위치가_없는_도형은_번호_후보에서_빠진다(proj, capsys):
    def no_xfrm(prs):
        s = _tb(list(prs.slides)[2], 0, 0, 280000, 260000, "견본_번호")
        s._element.find(qn("p:spPr")).remove(s._element.find(".//" + qn("a:xfrm")))
        assert s.left is None

    _edit_template(proj, no_xfrm)
    code, rep = _run(proj, capsys)
    assert code == 3, rep  # TypeError(종료 2)가 아니다
    Presentation(str(_out(proj, rep)))


def test_그림_경로가_루트_밖으로_나가면_거부(proj, capsys):
    (proj.parent / "outside.png").write_bytes(_png())
    code, rep = _run(proj, capsys, SPEC.replace("docs/img/a.png", "../outside.png"))
    assert code == 1 and rep["files"] == []
    assert any(p["item"] == "그림" and "밖" in p["reason"] for p in rep["problems"])


def _add_timing(prs, index: int) -> None:
    slide = list(prs.slides)[index]
    el = etree.fromstring(
        '<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par>'
        '<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"/></p:par></p:tnLst></p:timing>'
    )
    slide._element.append(el)


def test_도형을_지우는_쪽의_애니메이션은_지우고_경고한다(proj, capsys):
    _edit_template(proj, lambda prs: (_add_timing(prs, 0), _add_timing(prs, 2), _add_timing(prs, 3)))
    prs, rep = _produce(proj, capsys)
    slides = list(prs.slides)
    assert all(s._element.find(qn("p:timing")) is None for s in slides[:-1])  # 표지(지우기)·본문(유지)
    assert slides[-1]._element.find(qn("p:timing")) is not None  # 도형을 지우지 않는 끝 쪽은 그대로
    items = {(w["sheet"]) for w in rep["warnings"] if w["type"] == "animation-removed"}
    assert items == {"표지", "본문"}


def test_blank_template_은_pptx_지도를_종료_2로_거절(proj, capsys):
    from checker import blank_template

    src = proj / "templates" / "시험.pptx"
    code = blank_template.main(["--root", str(proj), "--form", "시험매뉴얼", "--src", str(src), "--out", str(proj / "x.xlsx")])
    out = json.loads(capsys.readouterr().out)
    assert code == 2 and "pptx 지도는 빈 양식" in out["problems"][0]["reason"]
