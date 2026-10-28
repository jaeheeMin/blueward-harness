"""엑셀 리더.

시트 이름과, 지정된 자리의 행을 읽는다. 헤더가 몇 행인지는 리더가 추측하지 않는다.
실제 템플릿을 뜯어보니 시트마다 헤더 위치가 다르고, 헤더처럼 보이지만 레이블-값 쌍인
행이 위에 있어서 어떤 heuristic 도 헛짚는다. 그래서 규칙 파일이 시트와 행을 지정한다.
"""
from __future__ import annotations

import atexit
from collections import OrderedDict
from pathlib import Path

from checker.readers import reader

# 같은 파일을 한 프로세스에서 여러 번 읽는 일이 잦다(#515). 개발사양서 템플릿은 styles.xml 이 11.7MB 라
# 한 번 여는 데 6초가 들고, table_headers 규칙마다 다시 열면 규칙 수만큼 곱해진다. 그래서 열린 읽기 전용
# 통합문서를 (경로, mtime_ns, 크기) 단위로 들고 있다가 재사용한다. 파일이 바뀌면 키가 달라져 새로 연다.
# 읽기 전용 통합문서는 zip 핸들을 쥐므로 몇 개만 두고(오래된 것부터 닫음) 프로세스가 끝날 때 모두 닫는다.
# 여는 데 실패한 예외는 캐시에 넣지 않는다 — 다음 호출이 다시 시도한다.
_MAX_OPEN = 4
_OPEN: "OrderedDict[tuple[str, int, int], object]" = OrderedDict()


def _key(path: Path) -> tuple[str, int, int]:
    resolved = Path(path).resolve()
    st = resolved.stat()
    return str(resolved), st.st_mtime_ns, st.st_size


def _workbook(path: Path):
    """읽기 전용 통합문서. 같은 (경로, mtime, 크기) 면 이미 연 것을 돌려준다."""
    import openpyxl

    key = _key(path)
    wb = _OPEN.get(key)
    if wb is not None:
        _OPEN.move_to_end(key)
        return wb
    # 같은 경로의 낡은 판은 버린다(파일이 바뀐 경우).
    for stale in [k for k in _OPEN if k[0] == key[0]]:
        _close(_OPEN.pop(stale))
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    _OPEN[key] = wb
    while len(_OPEN) > _MAX_OPEN:
        _close(_OPEN.popitem(last=False)[1])
    return wb


def _close(wb) -> None:
    try:
        wb.close()
    except Exception:  # noqa: BLE001 - 닫다가 난 오류가 검사 결과를 바꾸면 안 된다
        pass


def clear_cache() -> None:
    """열어 둔 통합문서를 모두 닫는다(프로세스 종료 때와 시험에서 부른다).

    열려 있는 동안 Windows 에서는 그 파일을 바꾸거나 지울 수 없다. 같은 프로세스에서 읽은 xlsx 를
    덮어쓰거나 지우려면 먼저 이것을 부른다(#515).
    """
    while _OPEN:
        _close(_OPEN.popitem(last=False)[1])


atexit.register(clear_cache)


@reader(".xlsx", ".xlsm")
def read_xlsx(path: Path) -> dict:
    wb = _workbook(path)
    return {"sheets": list(wb.sheetnames), "_path": str(path)}


def read_row(path: Path, sheet: str, row: int) -> list[str] | None:
    """한 시트의 한 행을 값 목록으로 읽는다. 시트가 없으면 None.

    꼬리의 빈 칸은 버린다. 엑셀은 쓰다 지운 자리를 빈 셀로 남겨 두는 일이 잦아서,
    그대로 두면 헤더 길이가 파일마다 달라진다.
    """
    wb = _workbook(path)
    if sheet not in wb.sheetnames:
        return None
    ws = wb[sheet]
    values = []
    for r in ws.iter_rows(min_row=row, max_row=row, values_only=True):
        values = ["" if v is None else str(v).strip() for v in r]
        break
    while values and values[-1] == "":
        values.pop()
    return values
